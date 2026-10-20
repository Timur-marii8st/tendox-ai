"""Only synthetic manifests, URLs and DOCX bytes."""

from importlib import import_module
from io import BytesIO

import pytest
from docx import Document

from procurement.access import AccessRole
from procurement.sources import eis_223_documents as source
from procurement.sources.eis_223_archive import parse_notice223_xml
from test_eis_223_archive import UID, URL, xml
from test_eis_document_jobs_api import send


def docx_bytes():
    document = Document()
    document.add_paragraph("Срок поставки: 10 рабочих дней.")
    document.add_paragraph("Оплата после поставки товара.")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_manifest_is_exact_but_does_not_infer_size_or_completeness():
    notice = parse_notice223_xml(xml(), source_member="notice.xml")
    manifest = source.manifest223(notice)
    assert manifest.status == "READY" and manifest.notice.notice_xml_sha256 == notice.xml_sha256
    item = manifest.documents[0]
    assert item.content_uid == UID and item.source_url == URL
    assert item.declared_byte_count is None and manifest.source_manifest_coverage == "UNKNOWN"
    assert item.source_locator.endswith("/t:document[1]/t:url")


@pytest.mark.parametrize("change", ["duplicate", "size", "filename", "private", "count"])
def test_ambiguous_manifest_never_becomes_ready(change):
    raw = xml().decode()
    if change == "duplicate":
        start, end = raw.index("<t:document>"), raw.index("</t:document>") + len("</t:document>")
        raw = raw[:end] + raw[start:end] + raw[end:]
    elif change == "size":
        notice = parse_notice223_xml(raw.encode(), source_member="notice.xml")
        notice = notice.model_copy(update={"attachments": notice.attachments * 33})
        assert "DOCUMENT_SET_RESOURCE_LIMIT" in source.manifest223(notice).reasons
        return
    elif change == "filename":
        raw = raw.replace("ТЗ.docx", "../ТЗ.docx")
    elif change == "private":
        raw = raw.replace("/filestore/public/", "/filestore/lk/")
    else:
        raw = raw.replace("<t:totalDocumentsCount>1", "<t:totalDocumentsCount>2")
    result = source.manifest223(parse_notice223_xml(raw.encode(), source_member="notice.xml"))
    assert result.status == "UNKNOWN"


@pytest.mark.parametrize(
    "url",
    [
        URL.replace("zakupki.gov.ru", "evil.invalid"),
        URL.replace("https:", "http:"),
        URL + "&other=1",
        URL + "#x",
        URL.replace("zakupki.gov.ru", "user@zakupki.gov.ru"),
        URL.replace("zakupki.gov.ru", "zakupki.gov.ru:443"),
        URL.replace(UID, "B" * 32),
    ],
)
def test_unsafe_url_never_reaches_browser(monkeypatch, url):
    monkeypatch.setattr(source, "_host_lease", lambda: pytest.fail("unexpected lease"))
    assert not source.public223_attachment_url(url, UID)
    with pytest.raises(ValueError, match="ATTACHMENT_URL_INVALID"):
        source.download_public223_attachment(url, UID)


def test_read_route_is_exact_tenant_readonly_no_store(monkeypatch):
    from datetime import UTC, datetime

    from procurement.storage.eis_223_documents import Document223Status

    calls = []
    manifest = source.manifest223(parse_notice223_xml(xml(), source_member="notice.xml"))

    def read(storage, company, ref):
        calls.append((company, ref))
        return Document223Status(
            manifest=manifest,
            checked_at=datetime.now(UTC),
            version_status="CURRENT_KNOWN_VERSION",
            blocked_reasons=(),
            documents=(),
            download_status="INCOMPLETE",
        )

    monkeypatch.setattr(import_module("procurement.api.app"), "get_document223_status", read)
    path = "/v1/tender-discovery/eis-223-documents/32600000001"
    url = path + "?version=2&xml_sha256=" + manifest.notice.notice_xml_sha256 + "&company_id=other"
    assert send("GET", url, token="invalid").status_code == 401
    response = send("GET", url, role=AccessRole.AUDITOR)
    assert response.status_code == 200 and response.headers["Cache-Control"] == "no-store"
    assert calls[0][0] == "tenant-a" and calls[0][1] == manifest.notice
    assert send("POST", path, body={}).status_code == 405
    assert send("GET", path + "?version=0&xml_sha256=" + "a" * 64).status_code == 422


@pytest.mark.parametrize("code,expected", [(LookupError("missing"), 404), (ValueError("bad"), 422)])
def test_route_rejects_inaccessible_or_conflicting_source(monkeypatch, code, expected):
    def read(*args):
        raise code

    monkeypatch.setattr(import_module("procurement.api.app"), "get_document223_status", read)
    url = "/v1/tender-discovery/eis-223-documents/32600000001?version=2&xml_sha256=" + "a" * 64
    assert send("GET", url).status_code == expected


def test_queue_api_requires_operator_and_ignores_no_payload_tenant(monkeypatch):
    from datetime import UTC, datetime

    from procurement.storage.eis_223_documents import Document223Status

    manifest = source.manifest223(parse_notice223_xml(xml(), source_member="notice.xml"))
    calls = []

    def plan(storage, company, ref, *, actor_id):
        calls.append((company, ref))
        return Document223Status(
            manifest=manifest,
            checked_at=datetime.now(UTC),
            version_status="CURRENT_KNOWN_VERSION",
            blocked_reasons=(),
            documents=(),
            download_status="INCOMPLETE",
            requested=True,
            queue_status="QUEUED",
        )

    monkeypatch.setattr(import_module("procurement.api.app"), "plan_document223_job", plan)
    path = "/v1/tender-discovery/eis-223-document-jobs"
    body = manifest.notice.model_dump()
    assert send("POST", path, body=body, role=AccessRole.AUDITOR).status_code == 403
    response = send("POST", path, body=body)
    assert response.status_code == 200 and response.headers["Cache-Control"] == "no-store"
    assert calls == [("tenant-a", manifest.notice)]
    assert send("POST", path, body={**body, "company_id": "other"}).status_code == 422
