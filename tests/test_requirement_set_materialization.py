from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document

from procurement.documents import (
    DocumentInput,
    ExecutableRequirementSelection,
    RequirementEntryDisposition,
    RequirementSetMaterializationError,
    RequirementSetMaterializationPreview,
    build_requirement_set_materialization_preview,
    extract_document,
    materialize_reviewed_requirement_entries,
)


def _extraction(*, duplicate: bool = False):
    document = Document()
    for _ in range(2 if duplicate else 1):
        outer = document.add_table(rows=1, cols=1)
        nested = outer.cell(0, 0).add_table(rows=3, cols=3)
        values = (
            ("Тип корпуса", "Rack", "Значение не изменяется"),
            ("Количество ядер", "Больше или равно 16 шт", "Указать значение"),
            ("Высота", "Не более 1", "Указать значение"),
        )
        for row, source in zip(nested.rows, values, strict=True):
            for cell, text in zip(row.cells, source, strict=True):
                cell.text = text
    output = BytesIO()
    document.save(output)
    content = output.getvalue()
    return extract_document(
        DocumentInput(
            document_id="tz-1",
            filename="tz.docx",
            declared_media_type=(
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            ),
            content=content,
        )
    )


def _preview() -> RequirementSetMaterializationPreview:
    return build_requirement_set_materialization_preview(
        company_id="tenant-a",
        source_document_id="tz-1",
        source_document_sha256="a" * 64,
        extraction=_extraction(),
        expected_entry_count=3,
    )


def test_preview_recovers_complete_ordered_table_and_non_executable_hints() -> None:
    preview = _preview()

    assert preview.expected_entry_count == 3
    assert preview.table_locator == "table 1 row 1 cell 1 nested table 1"
    assert [entry.raw_name for entry in preview.entries] == [
        "Тип корпуса",
        "Количество ядер",
        "Высота",
    ]
    assert preview.numeric_suggestion_count == 2
    assert preview.entries[0].numeric_suggestion is None
    assert preview.entries[1].numeric_suggestion is not None
    assert preview.entries[1].numeric_suggestion.value == 16
    assert (
        RequirementSetMaterializationPreview.model_validate_json(preview.model_dump_json())
        == preview
    )


def test_materialization_promotes_only_explicit_numeric_selection() -> None:
    preview = _preview()
    entries = materialize_reviewed_requirement_entries(
        preview,
        (
            ExecutableRequirementSelection(
                ordinal=2,
                property="cpu_cores_each",
                unit="шт",
                mandatory=True,
                review_note="Сверено оператором с обеими ячейками",
            ),
        ),
    )

    assert entries[0].disposition is RequirementEntryDisposition.MANUAL_REVIEW_REQUIRED
    assert entries[1].disposition is RequirementEntryDisposition.EXECUTABLE
    assert entries[1].requirement is not None
    assert entries[1].requirement.property == "cpu_cores_each"
    assert entries[1].requirement.value == 16
    assert entries[2].disposition is RequirementEntryDisposition.MANUAL_REVIEW_REQUIRED


def test_materialization_rejects_non_numeric_and_duplicate_selections() -> None:
    preview = _preview()
    non_numeric = ExecutableRequirementSelection(
        ordinal=1,
        property="chassis_type",
        mandatory=True,
        review_note="attempt",
    )
    with pytest.raises(RequirementSetMaterializationError, match="numeric suggestion"):
        materialize_reviewed_requirement_entries(preview, (non_numeric,))

    numeric = ExecutableRequirementSelection(
        ordinal=2,
        property="cpu_cores_each",
        mandatory=True,
        review_note="checked",
    )
    with pytest.raises(RequirementSetMaterializationError, match="must not repeat"):
        materialize_reviewed_requirement_entries(preview, (numeric, numeric))


def test_preview_rejects_ambiguous_or_incomplete_table_choice() -> None:
    with pytest.raises(RequirementSetMaterializationError, match="exactly one"):
        build_requirement_set_materialization_preview(
            company_id="tenant-a",
            source_document_id="tz-1",
            source_document_sha256="a" * 64,
            extraction=_extraction(duplicate=True),
            expected_entry_count=3,
        )

    with pytest.raises(RequirementSetMaterializationError, match="exactly one"):
        build_requirement_set_materialization_preview(
            company_id="tenant-a",
            source_document_id="tz-1",
            source_document_sha256="a" * 64,
            extraction=_extraction(),
            expected_entry_count=4,
        )
