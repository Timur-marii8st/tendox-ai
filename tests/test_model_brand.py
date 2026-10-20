from __future__ import annotations

import base64
import hashlib
import json
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from procurement.cli import main
from procurement.documents import (
    VISION_SEARCH_TAG_PROMPT_SHA256,
    DocumentInput,
    ExtractionStatus,
    ModelBrandCatalog,
    ModelBrandCatalogEntry,
    ModelBrandResolutionReason,
    ModelBrandSourceEvidence,
    SearchSourceKind,
    VisionSearchTag,
    VisionSearchTagManifest,
    VisionTagReviewStatus,
    build_document_search_entry,
    extract_document,
    load_model_brand_catalog,
    resolve_model_brand,
    search_documents,
)

NOW = datetime(2026, 9, 16, 12, tzinfo=UTC)
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _entry(
    entry_id: str = "general-climate-gw-g140-n1a1",
    *,
    model: str = "GW-G140/N1A1",
    brand: str = "General Climate",
    valid_from: datetime = datetime(2026, 1, 1, tzinfo=UTC),
    valid_to: datetime | None = None,
) -> ModelBrandCatalogEntry:
    return ModelBrandCatalogEntry(
        entry_id=entry_id,
        model=model,
        brand=brand,
        valid_from=valid_from,
        valid_to=valid_to,
        evidence=ModelBrandSourceEvidence(
            source_title="Official installation manual",
            source_url="https://manufacturer.example/manual.pdf",
            source_document_sha256="a" * 64,
            locator="model table, row 140",
            retrieved_at=NOW,
        ),
    )


def _catalog(*entries: ModelBrandCatalogEntry, company_id: str = "tenant-a") -> ModelBrandCatalog:
    return ModelBrandCatalog(
        catalog_id="catalog-20260916-01",
        company_id=company_id,
        created_at=NOW,
        approved_by="operator-7",
        approval_reference="review/model-brand/20260916-01",
        entries=entries or (_entry(),),
    )


def _docx_bytes(text: str, *, image: bool = False) -> bytes:
    document = Document()
    document.add_paragraph(text)
    if image:
        document.add_picture(BytesIO(_TINY_PNG))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _document(text: str, *, image: bool = False) -> DocumentInput:
    return DocumentInput(
        document_id="document-model",
        filename="technical_specification.docx",
        declared_media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        content=_docx_bytes(text, image=image),
    )


def test_exact_active_model_resolves_with_official_evidence() -> None:
    result = resolve_model_brand(
        _catalog(), company_id="tenant-a", model="gw-g140/n1a1", at=NOW
    )

    assert result.status == "PASS"
    assert result.reason is ModelBrandResolutionReason.RESOLVED
    assert result.brand == "General Climate"
    assert result.entries[0].evidence.source_kind == "OFFICIAL_MANUFACTURER"


def test_similar_model_is_not_inferred() -> None:
    result = resolve_model_brand(
        _catalog(), company_id="tenant-a", model="GW-G140W/N1D3", at=NOW
    )

    assert result.status == "UNKNOWN"
    assert result.reason is ModelBrandResolutionReason.NOT_FOUND
    assert result.brand is None


def test_tenant_expiry_and_conflict_remain_unknown() -> None:
    catalog = _catalog(
        _entry(valid_to=datetime(2026, 6, 1, tzinfo=UTC)),
        company_id="tenant-a",
    )
    expired = resolve_model_brand(catalog, company_id="tenant-a", model="GW-G140/N1A1", at=NOW)
    foreign = resolve_model_brand(catalog, company_id="tenant-b", model="GW-G140/N1A1", at=NOW)
    conflicting = resolve_model_brand(
        _catalog(_entry(), _entry("conflict", brand="Another Brand")),
        company_id="tenant-a",
        model="GW-G140/N1A1",
        at=NOW,
    )

    assert expired.reason is ModelBrandResolutionReason.NO_ACTIVE_EVIDENCE
    assert foreign.reason is ModelBrandResolutionReason.TENANT_MISMATCH
    assert conflicting.reason is ModelBrandResolutionReason.CONFLICTING_EVIDENCE
    assert expired.brand is foreign.brand is conflicting.brand is None


def test_catalog_loader_rejects_duplicate_json_keys(tmp_path: Path) -> None:
    path = tmp_path / "catalog.json"
    path.write_text('{"schema_version":"model-brand-catalog-v1","catalog_id":"a","catalog_id":"b"}')

    with pytest.raises(ValueError, match="duplicate key"):
        load_model_brand_catalog(path)


def test_catalog_rejects_evidence_retrieved_after_review() -> None:
    payload = _catalog().model_dump(mode="python")
    payload["created_at"] = datetime(2026, 9, 16, 11, tzinfo=UTC)

    with pytest.raises(ValueError, match="retrieved in the future"):
        ModelBrandCatalog.model_validate(payload)


def test_brand_query_finds_exact_model_in_native_text() -> None:
    document = _document("Наружный блок модели GW-G140/N1A1")
    search_entry = build_document_search_entry(
        document,
        extract_document(document),
        company_id="tenant-a",
    )

    result = search_documents(
        (search_entry,),
        company_id="tenant-a",
        query="General Climate",
        model_brand_catalog=_catalog(),
        at=NOW,
    )[0]

    assert result.status is ExtractionStatus.PASS
    assert [hit.source_kind for hit in result.hits] == [SearchSourceKind.MODEL_BRAND_CATALOG]
    evidence = result.hits[0].model_brand_evidence
    assert evidence is not None
    assert evidence.model_source_kind is SearchSourceKind.NATIVE_TEXT
    assert evidence.model_provenance.locator == "paragraph 1"


def test_catalog_does_not_fuzzy_match_model_in_document() -> None:
    document = _document("Наружный блок модели GW-G140W/N1D3")
    search_entry = build_document_search_entry(
        document,
        extract_document(document),
        company_id="tenant-a",
    )

    assert (
        search_documents(
            (search_entry,),
            company_id="tenant-a",
            query="General Climate",
            model_brand_catalog=_catalog(),
            at=NOW,
        )
        == ()
    )


def test_brand_from_ai_suggested_model_remains_unknown() -> None:
    document = _document("Визуальная спецификация", image=True)
    manifest = VisionSearchTagManifest(
        company_id="tenant-a",
        document_id=document.document_id,
        source_document_sha256=hashlib.sha256(document.content).hexdigest(),
        media_index=0,
        source_media="word/media/image1.png",
        source_media_sha256=hashlib.sha256(_TINY_PNG).hexdigest(),
        producer="vision-review-worker",
        model="approved-vision-model",
        prompt_sha256=VISION_SEARCH_TAG_PROMPT_SHA256,
        source_artifact_sha256="b" * 64,
        produced_at=NOW,
        tags=(
            VisionSearchTag(
                text="GW-G140/N1A1",
                confidence=0.91,
                locator='DOCX media "word/media/image1.png" :: equipment label',
                review_status=VisionTagReviewStatus.AI_SUGGESTED,
            ),
        ),
    )
    search_entry = build_document_search_entry(
        document,
        extract_document(document),
        company_id="tenant-a",
        vision_manifest=manifest,
    )

    result = search_documents(
        (search_entry,),
        company_id="tenant-a",
        query="General Climate",
        model_brand_catalog=_catalog(),
        at=NOW,
    )[0]

    assert result.status is ExtractionStatus.UNKNOWN
    assert result.hits[0].status is ExtractionStatus.UNKNOWN
    assert result.hits[0].vision_review_status is VisionTagReviewStatus.AI_SUGGESTED
    assert result.hits[0].model_brand_evidence is not None


def test_cli_accepts_reviewed_catalog_without_persisting_document(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    document_path = tmp_path / "specification.docx"
    document_path.write_bytes(_docx_bytes("Наружный блок GW-G140/N1A1"))
    catalog_path = tmp_path / "catalog.json"
    catalog_path.write_text(_catalog().model_dump_json(indent=2), encoding="utf-8")

    exit_code = main(
        [
            "search-document",
            str(document_path),
            "--company-id",
            "tenant-a",
            "--query",
            "General Climate",
            "--model-brand-catalog",
            str(catalog_path),
            "--at",
            NOW.isoformat(),
        ]
    )

    assert exit_code == 0
    output = json.loads(capsys.readouterr().out)
    assert output["effective_at"] == NOW.isoformat()
    assert output["hits"][0]["source_kind"] == "MODEL_BRAND_CATALOG"
    assert output["hits"][0]["model_brand_evidence"]["brand"] == "General Climate"
