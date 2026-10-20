"""Синтетический сквозной поиск и отрицательные границы persisted workflow."""

import asyncio
import hashlib
import json
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from typing import cast

import httpx
import pytest
from docx import Document

from procurement.access import AccessRole
from procurement.api import create_app
from procurement.api.contracts import Authenticator, CaseService, Principal
from procurement.domain.discovery_economics import DiscoveryCostInput
from procurement.sources import eis_public
from procurement.sources import eis_workflow as workflow
from procurement.sources.eis_analysis import analyze_document


def _docx() -> bytes:
    document = Document()
    document.add_paragraph("Гарантия не менее 24 месяцев")
    document.add_paragraph("Срок поставки: до 01.12.2027")
    document.add_paragraph("Оплата после приёмки, количество 100 штук")
    target = BytesIO()
    document.save(target)
    return target.getvalue()


def _lead(number: str = "0123456789012345678") -> eis_public.EisPublicLead:
    return eis_public.EisPublicLead(
        purchase_number=number,
        subject="Поставка бумаги",
        law="44-fz",
        method="Электронный аукцион",
        stage="Подача заявок",
        nmck_rub=Decimal("120000"),
        detail_url=f"https://zakupki.gov.ru/epz/order/notice/ea20/view/common-info.html?regNumber={number}",
        documents_url=f"https://zakupki.gov.ru/epz/order/notice/ea20/view/documents.html?regNumber={number}",
        matched_terms=("бумага",),
        source_locator="card[1]",
        source_sha256="a" * 64,
    )


def _costs(goods: str = "60000") -> DiscoveryCostInput:
    return DiscoveryCostInput(
        goods=goods,
        expenses="5000",
        vat_rate="0.2",
        basis="Синтетическая цена полного лота",
        costs_complete_ack=True,
    )


@pytest.fixture
def source(monkeypatch, tmp_path):
    monkeypatch.setenv("EIS_DISCOVERY_DB", str(tmp_path / "runs.sqlite3"))
    content = _docx()
    calls = []

    def search(query, page_number=1, raw_observer=None, page_size=100):
        calls.append(("search", page_number))
        if raw_observer:
            raw_observer(f"<html>search page {page_number}</html>".encode())
        return eis_public.EisPublicSearchResult(
            query=query,
            search_url=eis_public.public_search_url(query, page_number, page_size=page_size),
            status="COMPLETED",
            observed_at=datetime.now(UTC),
            page_number=page_number,
            page_size=page_size,
            source_sha256=hashlib.sha256(f"page {page_number}".encode()).hexdigest(),
            scanned_count=1,
            purchase_numbers=(_lead().purchase_number,),
            unidentified_count=0,
            unique_scanned_count=1,
            has_more_pages=page_number < 2,
            leads=(_lead(),),  # Same notice on two pages must not duplicate work.
        )

    def detail(url, raw_observer=None):
        calls.append(("detail", url))
        if raw_observer:
            raw_observer(b"<html>detail</html>")
        return eis_public.EisPublicDetailResult(
            detail_url=url,
            status="COMPLETED",
            observed_at=datetime.now(UTC),
            exact_deadline=datetime.now(UTC) + timedelta(days=10),
            customer=eis_public.EisPublicCustomer(
                name="ГБУ Пример",
                organization_code="01234567890",
                profile_url=(
                    "https://zakupki.gov.ru/epz/organization/view/info.html"
                    "?organizationCode=01234567890"
                ),
                source_locator="customer-header[1]",
                source_sha256="b" * 64,
            ),
        )

    def documents(url, raw_observer=None):
        calls.append(("documents", url))
        if raw_observer:
            raw_observer(b"<html>documents</html>")
        return eis_public.EisPublicDocumentsResult(
            documents_url=url,
            status="COMPLETED",
            observed_at=datetime.now(UTC),
            documents=(
                eis_public.EisPublicDocument(
                    filename="spec.docx",
                    source_locator="attachment[1]",
                    url="https://zakupki.gov.ru/44fz/filestore/public/1.0/download/spec.docx",
                ),
            ),
        )

    monkeypatch.setattr(eis_public, "search_public_eis", search)
    monkeypatch.setattr(eis_public, "read_public_detail", detail)
    monkeypatch.setattr(eis_public, "read_public_documents", documents)
    monkeypatch.setattr(eis_public, "download_public_attachment", lambda url: content)
    return content, calls


def test_customer_contract_history_is_bounded_scoped_and_does_not_change_estimate(
    source, monkeypatch
):
    _, calls = source
    contracts_link = (
        "https://zakupki.gov.ru/epz/contract/search/results.html"
        "?fz44=on&contractStageList=0&customerIdOrg=-1:ГБУzZz7708410783"
    )

    def profile(url, raw_observer=None):
        calls.append(("profile", url))
        if raw_observer:
            raw_observer(b"<html>official customer profile</html>")
        return eis_public.EisCustomerProfileResult(
            profile_url=url,
            organization_code="01234567890",
            status="COMPLETED",
            observed_at=datetime.now(UTC),
            source_sha256=hashlib.sha256(b"<html>official customer profile</html>").hexdigest(),
            inn="7708410783",
            inn_locator="inn-field[1]",
            contracts_url=contracts_link,
            contracts_locator="customer-contracts-link[2]",
        )

    def contracts(link, inn, name, raw_observer=None, product_query=None):
        calls.append(("contracts", inn, name, product_query))
        assert link == contracts_link and inn == "7708410783" and name == "ГБУ Пример"
        assert product_query == "бумага"
        html = b"<html>one completed contract</html>"
        if raw_observer:
            raw_observer(html)
        digest = hashlib.sha256(html).hexdigest()
        return eis_public.EisCustomerContractsResult(
            search_url=eis_public.completed_contracts_url(link, inn, product_query),
            status="COMPLETED",
            observed_at=datetime.now(UTC),
            source_sha256=digest,
            product_query=product_query,
            scanned_count=1,
            contracts=(
                eis_public.EisCustomerContract(
                    registry_number="2056204250625000119",
                    subject="Бумага А4",
                    price_rub=Decimal("100000.00"),
                    signed_on=datetime(2025, 12, 29, tzinfo=UTC).date(),
                    contract_url=(
                        "https://zakupki.gov.ru/epz/contract/contractCard/common-info.html"
                        "?reestrNumber=2056204250625000119"
                    ),
                    notice_number="0123456789012345678",
                    notice_url=(
                        "https://zakupki.gov.ru/epz/order/notice/view/common-info.html"
                        "?regNumber=0123456789012345678"
                    ),
                    relevance_status="MATCHED_VISIBLE",
                    matched_terms=("бумага",),
                    source_locator="contract-card[1]",
                    source_sha256=digest,
                ),
            ),
        )

    monkeypatch.setattr(eis_public, "read_customer_profile", profile)
    monkeypatch.setattr(eis_public, "read_completed_contracts", contracts)
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    _finish(run.id)
    before = workflow.set_costs("tenant-a", "operator", run.id, _lead().purchase_number, _costs())
    estimate = before.candidates[0].estimate

    queued = workflow.request_customer_history(
        "tenant-a", "operator", run.id, _lead().purchase_number
    )
    assert queued.state == "RUNNING" and queued.pending_count == 1
    after = _finish(run.id)
    candidate = after.candidates[0]
    assert candidate.customer_profile and candidate.customer_profile.inn == "7708410783"
    assert candidate.customer_contracts and len(candidate.customer_contracts.contracts) == 1
    assert candidate.customer_contracts.coverage == "FIRST_PAGE_ONLY"
    assert candidate.customer_contracts.product_query == "бумага"
    assert candidate.customer_contracts.contracts[0].relevance_status == "MATCHED_VISIBLE"
    assert candidate.estimate == estimate
    assert [name for name, *_ in calls].count("profile") == 1
    assert [name for name, *_ in calls].count("contracts") == 1
    assert (
        workflow.read_artifact("tenant-a", run.id, candidate.customer_contracts.source_sha256)[1]
        == b"<html>one completed contract</html>"
    )
    with pytest.raises(workflow.RunNotFound):
        workflow.request_customer_history("tenant-b", "operator", run.id, _lead().purchase_number)
    with workflow._db() as db:
        old = workflow._load(db, "tenant-a", run.id)
        old.candidates[0].customer_contracts = old.candidates[0].customer_contracts.model_copy(
            update={"product_query": None}
        )
        workflow._save(db, old)
    workflow.request_customer_history("tenant-a", "operator", run.id, _lead().purchase_number)
    refreshed = _finish(run.id)
    assert refreshed.candidates[0].customer_contracts.product_query == "бумага"
    assert [name for name, *_ in calls].count("contracts") == 2
    with workflow._db() as db:
        failed = workflow._load(db, "tenant-a", run.id)
        failed.candidates[0].customer_contracts = failed.candidates[
            0
        ].customer_contracts.model_copy(
            update={
                "status": "SOURCE_CHANGED",
                "contracts": (),
                "reason": "CONTRACT_FILTER_UNVERIFIED",
            }
        )
        workflow._save(db, failed)
    retried = workflow.request_customer_history(
        "tenant-a", "operator", run.id, _lead().purchase_number
    )
    assert retried.pending_count == 1
    _finish(run.id)
    assert [name for name, *_ in calls].count("contracts") == 3
    workflow.cancel_run("tenant-a", "operator", run.id)
    with pytest.raises(workflow.RunBusy, match="discovery_run_cancelled"):
        workflow.request_customer_history("tenant-a", "operator", run.id, _lead().purchase_number)


def test_customer_history_does_not_treat_223_customer_as_44(source):
    _, calls = source
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    with workflow._db() as db:
        stored = workflow._load(db, "tenant-a", run.id)
        stored.candidates = [workflow.Candidate(lead=_lead().model_copy(update={"law": "223-fz"}))]
        workflow._save(db, stored)
    with pytest.raises(workflow.RunBusy, match="customer_history_44_only"):
        workflow.request_customer_history("tenant-a", "operator", run.id, _lead().purchase_number)
    assert calls == []


def _ready(run_id: str) -> None:
    with workflow._db() as db:
        run = workflow._load(db, "tenant-a", run_id)
        run.next_retry_at = None
        workflow._save(db, run)


def _finish(run_id: str) -> workflow.DiscoveryRun:
    for _ in range(20):
        _ready(run_id)
        run = workflow.advance_run("tenant-a", "operator", run_id)
        if run.state != "RUNNING":
            return run
    raise AssertionError("bounded workflow did not finish")


def test_connected_discovery_persists_sources_requirements_and_three_scenarios(source):
    content, calls = source
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    result = _finish(run.id)
    assert result.pages_read == 2 and result.scanned_count == 2
    assert len(result.candidates) == 1
    candidate = result.candidates[0]
    assert candidate.documents[0].requirement_candidates
    assert candidate.documents[0].unmapped_count > 0
    assert candidate.application_status == "UNKNOWN"  # No admission from document text.
    digest = hashlib.sha256(content).hexdigest()
    assert workflow.read_artifact("tenant-a", run.id, digest) == ("spec.docx", content)
    assert len([c for c in calls if c[0] == "detail"]) == 1
    result = workflow.set_costs(
        "tenant-a", "operator", run.id, candidate.lead.purchase_number, _costs()
    )
    estimate = result.candidates[0].estimate
    assert estimate and estimate.status == "ESTIMATE"
    assert len(estimate.scenarios) == 3 and estimate.recommendation == "needs_analysis"
    assert estimate.scenarios[0].profit == Decimal("35000.00")
    persisted = workflow.get_run("tenant-a", run.id)
    assert persisted.candidates[0].estimate == estimate
    public = workflow.public_run(persisted)
    assert "company_id" not in public and "tasks" not in public
    assert "blocks" not in public["candidates"][0]["documents"][0]
    with workflow._db() as db:
        assert (
            db.execute("SELECT COUNT(*) FROM eis_audit WHERE company_id='tenant-a'").fetchone()[0]
            >= 7
        )


def test_search_exposes_unparsed_cards_and_source_result_cap(source, monkeypatch):
    def search(query, page_number=1, raw_observer=None, page_size=100):
        del raw_observer
        return eis_public.EisPublicSearchResult(
            query=query,
            search_url=eis_public.public_search_url(query, page_number, page_size=page_size),
            status="COMPLETED",
            observed_at=datetime.now(UTC),
            scanned_count=2,
            unparsed_count=1,
            source_result_count_floor=15000,
            query_exceeds_page_limit=True,
            leads=(_lead(),),
        )

    monkeypatch.setattr(workflow.eis_public, "search_public_eis", search)
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    result = workflow.advance_run("tenant-a", "operator", run.id)
    assert result.pages_read == 1 and result.scanned_count == 2
    assert any("пропущено карточек" in warning for warning in result.warnings)
    assert any("20×100" in warning for warning in result.warnings)


@pytest.mark.parametrize("page_size", [50, 100])
def test_repeat_guard_uses_all_raw_ids_after_reload_and_not_title_shortlist(
    source, monkeypatch, page_size
):
    numbers = ("0000000000000000001", "0000000000000000002")
    called = []

    def search(query, page_number=1, raw_observer=None, page_size=100):
        called.append(page_number)
        return eis_public.EisPublicSearchResult(
            query=query,
            search_url=eis_public.public_search_url(query, page_number, page_size=page_size),
            status="COMPLETED",
            observed_at=datetime.now(UTC),
            source_sha256=hashlib.sha256(f"different HTML {page_number}".encode()).hexdigest(),
            page_number=page_number,
            page_size=page_size,
            scanned_count=2,
            purchase_numbers=numbers if page_number == 1 else numbers[::-1],
            unidentified_count=0,
            unique_scanned_count=2,
            has_more_pages=True,
            leads=(),  # Raw result pages can have no title matches at all.
        )

    monkeypatch.setattr(eis_public, "search_public_eis", search)
    run = workflow.create_run(
        "tenant-a",
        "operator",
        workflow.DiscoveryRequest(
            query="бумага",
            max_pages=20,
            search_page_size=page_size,
        ),
    )
    first = workflow.advance_run("tenant-a", "operator", run.id)
    assert first.unique_scanned_count == 2
    _ready(run.id)
    result = workflow.advance_run("tenant-a", "operator", run.id)
    stored = workflow.get_run("tenant-a", run.id)
    assert called == [1, 2]
    assert result.search_stop_reason == "REPEATED_PAGE"
    assert result.state == "COMPLETED" and result.completion_status == "PARTIAL"
    assert stored.scanned_count == 4 and stored.unique_scanned_count == 2
    assert stored.duplicate_count == 2 and stored.identified_count == 2
    assert len(stored.search_pages) == 2
    assert stored.search_pages[0].fingerprint == stored.search_pages[1].fingerprint
    assert stored.search_pages[0].source_sha256 != stored.search_pages[1].source_sha256
    assert workflow.public_run(stored)["search_stop_reason"] == "REPEATED_PAGE"
    with workflow._db() as db:
        event = db.execute(
            "SELECT detail FROM eis_audit WHERE company_id=? AND run_id=? AND event=?",
            ("tenant-a", run.id, "discovery.search_stopped"),
        ).fetchone()
    assert json.loads(event[0])["page"]["fingerprint"] == stored.search_pages[-1].fingerprint


@pytest.mark.parametrize("incomplete", [False, True])
def test_partial_overlap_and_unidentified_rows_never_inflate_unique(
    source, monkeypatch, incomplete
):
    first_ids = ("0000000000000000001", "0000000000000000002")
    second_ids = ("0000000000000000002", "0000000000000000003")
    third_ids = ("0000000000000000001",)
    called = []

    def search(query, page_number=1, raw_observer=None, page_size=100):
        called.append(page_number)
        ids = (first_ids, second_ids, third_ids)[page_number - 1]
        missing = int(incomplete and page_number == 2)
        return eis_public.EisPublicSearchResult(
            query=query,
            search_url=eis_public.public_search_url(query, page_number),
            status="COMPLETED",
            observed_at=datetime.now(UTC),
            source_sha256="a" * 64,
            page_number=page_number,
            scanned_count=len(ids) + missing,
            purchase_numbers=ids,
            unidentified_count=missing,
            has_more_pages=True,
        )

    monkeypatch.setattr(eis_public, "search_public_eis", search)
    run = workflow.create_run(
        "tenant-a", "operator", workflow.DiscoveryRequest(query="бумага", max_pages=20)
    )
    result = _finish(run.id)
    assert called == ([1, 2] if incomplete else [1, 2, 3])
    assert result.identified_count == 3
    assert result.unique_scanned_count == (None if incomplete else 3)
    assert result.duplicate_count == (1 if incomplete else 2)
    assert result.search_stop_reason == (
        "IDENTIFIERS_INCOMPLETE" if incomplete else "NO_NEW_IDENTIFIERS"
    )


def test_legacy_search_without_page_evidence_stops_before_changing_offsets(source):
    _, calls = source
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    with workflow._db() as db:
        legacy = run.model_dump(mode="json")
        for field in (
            "search_pages",
            "unique_scanned_count",
            "identified_count",
            "duplicate_count",
            "search_stop_reason",
            "unparsed_count",
        ):
            legacy.pop(field, None)
        legacy["request"].pop("search_page_size", None)
        legacy["pages_read"] = 1
        legacy["scanned_count"] = 200
        legacy["tasks"][0]["page"] = 2
        db.execute(
            "UPDATE eis_runs SET payload=? WHERE company_id=? AND id=?",
            (json.dumps(legacy), "tenant-a", run.id),
        )
    result = workflow.advance_run("tenant-a", "operator", run.id)
    assert calls == []
    assert result.search_stop_reason == "LEGACY_STATE_UNVERIFIED"
    assert result.unique_scanned_count is None
    assert result.scanned_count == 200


def test_tenant_isolation_every_operation(source):
    content, _ = source
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    _finish(run.id)
    assert workflow.list_runs("tenant-b") == []
    for operation in (
        lambda: workflow.get_run("tenant-b", run.id),
        lambda: workflow.advance_run("tenant-b", "operator", run.id),
        lambda: workflow.cancel_run("tenant-b", "operator", run.id),
        lambda: workflow.set_costs(
            "tenant-b", "operator", run.id, _lead().purchase_number, _costs()
        ),
        lambda: workflow.read_artifact("tenant-b", run.id, hashlib.sha256(content).hexdigest()),
    ):
        with pytest.raises(workflow.RunNotFound):
            operation()


def test_cooldown_does_not_consume_task_or_contact_source_again(source, monkeypatch):
    _, calls = source
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    future = datetime.now(UTC) + timedelta(seconds=60)
    monkeypatch.setattr(
        eis_public,
        "search_public_eis",
        lambda query, **kwargs: eis_public.EisPublicSearchResult(
            query=query,
            search_url=eis_public.public_search_url(query),
            status="COOLDOWN",
            observed_at=datetime.now(UTC),
            retry_at=future,
        ),
    )
    waiting = workflow.advance_run("tenant-a", "operator", run.id)
    assert waiting.pending_count == 1 and waiting.tasks[0].attempts == 0
    assert waiting.next_retry_at == future
    again = workflow.advance_run("tenant-a", "operator", run.id)
    assert again.pending_count == 1 and not calls


def test_failed_document_is_partial_but_source_and_other_results_retained(source, monkeypatch):
    monkeypatch.setattr(eis_public, "download_public_attachment", lambda url: b"corrupt docx")
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    result = _finish(run.id)
    candidate = result.candidates[0]
    assert candidate.analysis_status == "PARTIAL"
    assert candidate.documents[0].status == "UNKNOWN"
    assert workflow.read_artifact("tenant-a", run.id, hashlib.sha256(b"corrupt docx").hexdigest())


def test_expiration_recomputed_after_completion_blocks_estimate(source):
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    result = _finish(run.id)
    workflow.set_costs("tenant-a", "operator", run.id, _lead().purchase_number, _costs())
    with workflow._db() as db:
        result = workflow._load(db, "tenant-a", run.id)
        result.candidates[0].detail = result.candidates[0].detail.model_copy(
            update={
                "exact_deadline": datetime.now(UTC) - timedelta(seconds=1),
            }
        )
        workflow._save(db, result)
    expired = workflow.get_run("tenant-a", run.id).candidates[0]
    assert expired.application_status == "FAIL"
    assert expired.estimate.status == "BLOCKED" and not expired.estimate.scenarios


def test_active_lease_prevents_duplicate_source_work_and_cancel_stops_it(source):
    _, calls = source
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    with workflow._db() as db:
        db.execute(
            "UPDATE eis_runs SET lease=?,lease_until=? WHERE id=?",
            ("other-worker", (datetime.now(UTC) + timedelta(minutes=1)).isoformat(), run.id),
        )
    workflow.advance_run("tenant-a", "operator", run.id)
    assert not calls
    workflow.cancel_run("tenant-a", "operator", run.id)
    assert workflow.advance_run("tenant-a", "operator", run.id).state == "CANCELLED"
    assert not calls


def test_cancelled_run_cannot_restart_network_analysis(source):
    _, calls = source
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    workflow.advance_run("tenant-a", "operator", run.id)
    workflow.cancel_run("tenant-a", "operator", run.id)

    with pytest.raises(workflow.RunBusy, match="discovery_run_cancelled"):
        workflow.analyze_candidate("tenant-a", "operator", run.id, _lead().purchase_number)
    assert workflow.get_run("tenant-a", run.id).state == "CANCELLED"
    assert calls == [("search", 1)]


def test_unpriced_shortlist_keeps_relevance_order(source):
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    less_relevant = _lead("0123456789012345678")
    more_relevant = _lead("0123456789012345679").model_copy(
        update={"matched_terms": ("бумага", "офисная")}
    )
    run.candidates = [
        workflow.Candidate(lead=less_relevant),
        workflow.Candidate(lead=more_relevant),
    ]

    result = workflow._refresh(run)
    assert result.candidates[0].lead.purchase_number == more_relevant.purchase_number


def test_document_mapping_uses_text_after_preview_boundary():
    document = Document()
    document.add_paragraph("Описание " * 700)
    document.add_paragraph("Гарантия не менее 24 месяцев")
    target = BytesIO()
    document.save(target)
    result = analyze_document(target.getvalue(), "spec.docx", "https://zakupki.gov.ru/file")
    assert sum(len(b.text) for b in result.blocks) > 4000
    assert result.requirement_candidates[0].property == "warranty"
    assert result.requirement_candidates[0].provenance.document_id == result.sha256


class _Auth:
    async def authenticate(self, credential: str) -> Principal | None:
        if credential == "auditor":
            return Principal(
                company_id="tenant-a", actor_id="reviewer", access_role=AccessRole.AUDITOR
            )
        if credential not in {"tenant-a", "tenant-b"}:
            return None
        return Principal(
            company_id=credential, actor_id="operator", access_role=AccessRole.OPERATOR
        )


def test_workflow_api_auth_scope_and_validation(source):
    app = create_app(cast(CaseService, object()), cast(Authenticator, _Auth()))

    async def scenario():
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://test"
        ) as client:
            assert (
                await client.post("/v1/tender-discovery/runs", json={"query": "бумага"})
            ).status_code == 401
            headers = {"Authorization": "Bearer tenant-a"}
            created = await client.post(
                "/v1/tender-discovery/runs", headers=headers, json={"query": "бумага"}
            )
            assert created.status_code == 200
            run_id = created.json()["id"]
            assert "company_id" not in created.json()
            assert created.headers["cache-control"] == "no-store"
            path = f"/v1/tender-discovery/runs/{run_id}"
            assert (
                await client.get(path, headers={"Authorization": "Bearer tenant-b"})
            ).status_code == 404
            advanced = await client.post(path + "/advance", headers=headers)
            assert advanced.status_code == 200
            assert advanced.json()["unique_scanned_count"] == 1
            assert advanced.json()["search_pages"][0]["page_size"] == 100
            assert "company_id" not in advanced.json()
            assert (
                await client.post(
                    "/v1/tender-discovery/runs",
                    headers=headers,
                    json={"query": "бумага", "search_page_size": 200},
                )
            ).status_code == 422
            history_path = path + "/customer-history/0123456789012345678"
            assert (
                await client.post(history_path, headers={"Authorization": "Bearer tenant-b"})
            ).status_code == 404
            assert (await client.post(history_path, headers=headers)).status_code == 200
            bad = await client.post(
                path + "/costs/0123456789012345678",
                headers=headers,
                json={
                    **_costs().model_dump(mode="json"),
                    "goods": 1.2,
                },
            )
            assert bad.status_code == 422
            assert (await client.post(path + "/cancel", headers=headers)).json()[
                "state"
            ] == "CANCELLED"

    asyncio.run(scenario())


def test_blocked_search_is_failed_not_empty_success(source, monkeypatch):
    monkeypatch.setattr(
        eis_public,
        "search_public_eis",
        lambda query, **kwargs: eis_public.EisPublicSearchResult(
            query=query,
            search_url=eis_public.public_search_url(query),
            status="BLOCKED",
            reason="BLOCKED_PAGE",
            observed_at=datetime.now(UTC),
        ),
    )
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    result = workflow.advance_run("tenant-a", "operator", run.id)
    assert result.state == "COMPLETED" and result.completion_status == "FAILED"
    assert result.warnings and not result.candidates


def test_unsupported_file_does_not_consume_supported_document_budget(source, monkeypatch):
    def documents(url, raw_observer=None):
        return eis_public.EisPublicDocumentsResult(
            documents_url=url,
            status="COMPLETED",
            observed_at=datetime.now(UTC),
            documents=(
                eis_public.EisPublicDocument(
                    filename="old.doc",
                    supported=False,
                    source_locator="attachment[1]",
                    url="https://zakupki.gov.ru/epz/old.doc",
                ),
                eis_public.EisPublicDocument(
                    filename="spec.docx",
                    supported=True,
                    source_locator="attachment[2]",
                    url="https://zakupki.gov.ru/epz/spec.docx",
                ),
            ),
        )

    monkeypatch.setattr(eis_public, "read_public_documents", documents)
    run = workflow.create_run(
        "tenant-a", "operator", workflow.DiscoveryRequest(query="бумага", max_documents=1)
    )
    candidate = _finish(run.id).candidates[0]
    assert candidate.documents[0].filename == "spec.docx"
    assert any("old.doc" in w for w in candidate.warnings)
    assert candidate.analysis_status == "PARTIAL"


def test_detail_cancellation_blocks_economics_even_when_search_was_open(source):
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    _finish(run.id)
    with workflow._db() as db:
        current = workflow._load(db, "tenant-a", run.id)
        current.candidates[0].detail = current.candidates[0].detail.model_copy(
            update={
                "fields": (
                    eis_public.EisPublicDetailField(
                        name="stage",
                        label="Этап закупки",
                        value="Закупка отменена",
                        source_locator="stage",
                        evidence_snippet="Закупка отменена",
                    ),
                ),
            }
        )
        workflow._save(db, current)
    result = workflow.set_costs("tenant-a", "operator", run.id, _lead().purchase_number, _costs())
    assert result.candidates[0].estimate.status == "BLOCKED"


def test_analysis_limit_is_visible_and_never_becomes_approved_requirement():
    result = analyze_document(
        _docx(), "spec.docx", "https://zakupki.gov.ru/epz/doc", max_characters=1
    )
    assert result.status == "UNKNOWN" and "ANALYSIS_LIMIT_FULL_SOURCE_RETAINED" in result.reasons
    assert not result.blocks and result.document_role == "UNKNOWN"
    full = analyze_document(_docx(), "spec.docx", "https://zakupki.gov.ru/epz/doc")
    assert all(not item.reviewed for item in full.requirement_candidates)
    assert all("mandatory" not in item.model_dump() for item in full.requirement_candidates)


def test_artifact_budget_and_cancelled_run_reject_new_writes(source, monkeypatch):
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))
    monkeypatch.setattr(workflow, "MAX_RUN_BYTES", 2)
    with pytest.raises(ValueError, match="RUN_ARTIFACT_LIMIT"):
        workflow._artifact(run, b"large", "test.bin", "https://zakupki.gov.ru/epz/file")
    workflow.cancel_run("tenant-a", "operator", run.id)
    with pytest.raises(ValueError, match="RUN_STOPPED"):
        workflow._artifact(run, b"x", "test.bin", "https://zakupki.gov.ru/epz/file")


def test_auditor_can_read_but_cannot_mutate_discovery(source):
    app = create_app(cast(CaseService, object()), cast(Authenticator, _Auth()))
    run = workflow.create_run("tenant-a", "operator", workflow.DiscoveryRequest(query="бумага"))

    async def scenario():
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app),
            base_url="http://test",
            headers={"Authorization": "Bearer auditor"},
        ) as client:
            path = f"/v1/tender-discovery/runs/{run.id}"
            assert (await client.get(path)).status_code == 200
            assert (
                await client.post("/v1/tender-discovery/runs", json={"query": "бумага"})
            ).status_code == 403
            for endpoint in [
                "advance",
                "cancel",
                "analyze/0123456789012345678",
                "customer-history/0123456789012345678",
            ]:
                assert (await client.post(path + "/" + endpoint)).status_code == 403
            assert (
                await client.post(
                    path + "/costs/0123456789012345678", json=_costs().model_dump(mode="json")
                )
            ).status_code == 403

    asyncio.run(scenario())
