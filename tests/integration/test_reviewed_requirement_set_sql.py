"""PostgreSQL invariants for reviewed requirement sets and case bindings."""

from __future__ import annotations

import hashlib
import os
import uuid
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from alembic import command
from alembic.config import Config
from docx import Document
from sqlalchemy import create_engine, delete, func, select, update
from sqlalchemy.exc import DBAPIError

from case_support import paper_case_at
from procurement.documents import (
    CaseHardGateDecisionInput,
    CaseHardGateEvidenceRef,
    CompletenessStatus,
    ConstituentDecisionRef,
    DocumentInput,
    DocumentRole,
    DocumentSetCrossAssessment,
    DocumentSetInventoryAssessment,
    DocumentSetMemberAssessment,
    DocumentSetMemberAttestation,
    RequirementEntryDisposition,
    RequirementSetCoverageAssessment,
    RequirementSetNormalizationAssessment,
    ReviewedRequirementEntry,
    ServerOfferBomLineInput,
    ServerOfferBomScopeAssessment,
    ServerOfferCommercialAssessment,
    ServerOfferCommercialTermsInput,
    ServerOfferEvidenceKind,
    ServerOfferEvidenceRef,
    ServerOfferOperatingCostsInput,
    ServerOfferPartKind,
    ServerOfferRequirementDecision,
    ServerOfferScenarioInput,
    extract_document,
)
from procurement.domain.cases import GateKind, ProcurementCase, case_input_digest, evaluate_case
from procurement.domain.compliance import Verdict
from procurement.domain.decisions import RecommendationDecision
from procurement.domain.economics import ScenarioKind
from procurement.domain.provenance import Extractor, Provenance
from procurement.domain.requirements import Operator, Requirement
from procurement.domain.sourcing import (
    LeadSourceKind,
    ServerLead,
    SourcingRunRecord,
    SourcingRunStatus,
    shortlist_server_leads,
)
from procurement.storage import IdempotencyConflict, PostgresStorage
from procurement.storage.models import (
    case_requirement_set_bindings,
    reviewed_requirement_set_decisions,
    reviewed_requirement_set_proposals,
    server_offer_economics_packages,
    server_offer_proposals,
    server_offer_review_decisions,
)


@pytest.fixture(scope="session")
def database_url() -> str:
    value = os.environ.get("DATABASE_URL")
    if not value:
        pytest.skip("DATABASE_URL is not set; PostgreSQL integration tests are skipped")
    return value


@pytest.fixture(scope="session", autouse=True)
def migrated_database(database_url: str) -> None:
    admin_url = os.environ.get("DATABASE_ADMIN_URL")
    if admin_url:
        config = Config("alembic.ini")
        config.set_main_option("sqlalchemy.url", admin_url.replace("%", "%%"))
        command.upgrade(config, "head")


@pytest.fixture(scope="session")
def storage(database_url: str) -> PostgresStorage:
    engine = create_engine(database_url, pool_pre_ping=True)
    yield PostgresStorage(engine)
    engine.dispose()


def _case(tenant: str) -> ProcurementCase:
    payload = paper_case_at(datetime.now(UTC)).model_dump(mode="python")
    payload["company_id"] = tenant
    payload["case_id"] = f"case-{uuid.uuid4()}"
    payload["profile"]["company_id"] = tenant
    if payload["selected_quote"] is not None:
        payload["selected_quote"]["company_id"] = tenant
    return ProcurementCase.model_validate(payload)


def _source(storage: PostgresStorage, tenant: str):
    document_id = f"tz-{uuid.uuid4()}"
    source = Document()
    source.add_paragraph("Ядер каждого процессора не менее 16")
    output = BytesIO()
    source.save(output)
    content = output.getvalue()
    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    extraction = extract_document(
        DocumentInput(
            document_id=document_id,
            filename="tz.docx",
            declared_media_type=media_type,
            content=content,
        )
    )
    record = storage.store_document(
        tenant,
        "operator-a",
        "tz.docx",
        media_type,
        DocumentRole.TENDER_REQUIREMENTS,
        content,
        extraction,
        f"document-{uuid.uuid4()}",
        hashlib.sha256(content).hexdigest(),
        document_id=document_id,
    )
    provenance = Provenance(
        document_id=document_id,
        locator="paragraph 1",
        extractor=Extractor.MANUAL,
    )
    entry = ReviewedRequirementEntry(
        ordinal=1,
        raw_name="Ядер каждого CPU",
        raw_value="Не менее 16",
        provenance=provenance,
        disposition=RequirementEntryDisposition.EXECUTABLE,
        requirement=Requirement(
            property="cpu_cores_each",
            operator=Operator.GE,
            value=Decimal(16),
            mandatory=True,
            raw_text="Ядер каждого процессора не менее 16",
            provenance=provenance,
        ),
        review_note="Сверено с исходной строкой",
    )
    return record, entry


def _document_binding(storage: PostgresStorage, tenant: str, document_id: str):
    support_id = f"support-{uuid.uuid4()}"
    support_source = Document()
    support_source.add_paragraph("Сопроводительный документ закупки")
    support_output = BytesIO()
    support_source.save(support_output)
    support_content = support_output.getvalue()
    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    support_document = storage.store_document(
        tenant,
        "operator-a",
        "support.docx",
        media_type,
        DocumentRole.TENDER_REQUIREMENTS,
        support_content,
        extract_document(
            DocumentInput(
                document_id=support_id,
                filename="support.docx",
                declared_media_type=media_type,
                content=support_content,
            )
        ),
        "hard-gate-support-document",
        hashlib.sha256(support_content).hexdigest(),
        document_id=support_id,
    )
    reviewed = []
    for ordinal, source_document_id in enumerate(
        (document_id, support_document.document_id), start=1
    ):
        draft = storage.store_completeness_review(
            tenant,
            "operator-a",
            source_document_id,
            f"hard-gate-completeness-review-{ordinal}",
            str(ordinal + 4) * 64,
            review_id=f"review-{uuid.uuid4()}",
        )
        reviewed.append(
            storage.decide_completeness_review(
                tenant,
                "document-reviewer",
                draft.snapshot.review_id,
                draft.snapshot.evidence_sha256,
                CompletenessStatus.UNKNOWN,
                "entire selected document",
                "formal completeness remains unresolved",
                (),
                f"hard-gate-completeness-decision-{ordinal}",
                str(ordinal + 6) * 64,
            )
        )
    manifest = storage.store_document_set_evidence_manifest(
        tenant,
        "operator-a",
        "exact tender package",
        tuple((item.snapshot.review_id, item.snapshot.evidence_sha256) for item in reviewed),
        "hard-gate-manifest",
        "9" * 64,
        manifest_id=f"manifest-{uuid.uuid4()}",
    )
    evidence = storage.get_document_set_review_evidence(tenant, manifest.manifest_id)
    assert evidence is not None
    decision = storage.store_document_set_review_decision(
        tenant,
        "package-reviewer",
        manifest.manifest_id,
        manifest.evidence_sha256,
        tuple(
            DocumentSetMemberAttestation(
                reference=ConstituentDecisionRef(
                    member_ordinal=constituent.member.ordinal,
                    review_id=constituent.member.completeness_review_id,
                    review_evidence_sha256=(
                        constituent.member.completeness_review_evidence_sha256
                    ),
                    decision_sha256=constituent.decision_sha256 or "",
                ),
                assessment=DocumentSetMemberAssessment.UNRESOLVED,
                rationale="formal completeness remains unresolved",
            )
            for constituent in evidence.constituents
        ),
        DocumentSetInventoryAssessment.UNRESOLVED,
        "package inventory remains unresolved",
        DocumentSetCrossAssessment.UNRESOLVED,
        "cross-document completeness remains unresolved",
        (),
        CompletenessStatus.UNKNOWN,
        "document gate remains fail-closed",
        "hard-gate-document-set-decision",
        "b" * 64,
    )
    return storage.store_case_document_evidence_binding(
        tenant,
        "operator-a",
        manifest.manifest_id,
        manifest.evidence_sha256,
        decision.decision_sha256,
        "hard-gate-document-binding",
        "c" * 64,
        binding_id=f"document-binding-{uuid.uuid4()}",
    )


def test_reviewed_requirement_set_is_exact_immutable_and_tenant_scoped(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    other = f"tenant-{uuid.uuid4()}"
    case = _case(tenant)
    storage.save_case(tenant, case)
    document, entry = _source(storage, tenant)
    document_binding = _document_binding(storage, tenant, document.document_id)
    set_id = f"set-{uuid.uuid4()}"
    prepared = storage.store_reviewed_requirement_set_proposal(
        tenant,
        "operator-a",
        set_id=set_id,
        source_document_id=document.document_id,
        source_document_sha256=document.sha256,
        declared_scope="entire requirement table",
        expected_entry_count=1,
        entries=(entry,),
        idempotency_key="requirement-set-prepare",
        request_sha256="a" * 64,
    )
    decided = storage.decide_reviewed_requirement_set(
        tenant,
        "reviewer-a",
        set_id=set_id,
        proposal_sha256=prepared.proposal.proposal_sha256,
        coverage_assessment=RequirementSetCoverageAssessment.CONFIRMED,
        normalization_assessment=RequirementSetNormalizationAssessment.REVIEWED,
        status=Verdict.PASS,
        reason="all rows reviewed",
        idempotency_key="requirement-set-decision",
        request_sha256="b" * 64,
    )
    assert decided.decision is not None
    binding = storage.bind_reviewed_requirement_set_to_case(
        tenant,
        "operator-b",
        binding_id=f"binding-{uuid.uuid4()}",
        case_id=case.case_id,
        case_version=case.version,
        set_id=set_id,
        proposal_sha256=prepared.proposal.proposal_sha256,
        decision_sha256=decided.decision.decision_sha256,
        idempotency_key="requirement-set-binding",
        request_sha256="c" * 64,
    )
    loaded = storage.get_bound_requirement_set(tenant, case.case_id, case.version)
    assert loaded is not None and loaded[0] == binding and loaded[1] == decided
    assert storage.get_reviewed_requirement_set(other, set_id) is None
    assert storage.get_case_requirement_set_binding(other, binding.binding_id) is None

    observed_at = datetime.now(UTC)
    source_content = b"<html><title>HPE DL360 Gen12</title></html>"
    source_sha256 = hashlib.sha256(source_content).hexdigest()
    lead = ServerLead(
        lead_id="lead-1",
        supplier="HPE",
        manufacturer="HPE",
        source_kind=LeadSourceKind.MANUFACTURER_CATALOG,
        product_name="HPE DL360 Gen12",
        source_url="https://example.com/catalog",
        product_url="https://example.com/server",
        source_sha256=source_sha256,
        observed_at=observed_at,
    )
    run = SourcingRunRecord(
        company_id=tenant,
        run_id=f"run-{uuid.uuid4()}",
        case_id=case.case_id,
        case_version=case.version,
        query="DL360 Gen12",
        source_name="HPE",
        source_url="https://example.com/catalog",
        status=SourcingRunStatus.COMPLETED,
        source_sha256=source_sha256,
        source_byte_size=len(source_content),
        observed_at=observed_at,
        shortlist=shortlist_server_leads((lead,), prepared.proposal.executable_requirements),
        requirement_set_id=binding.set_id,
        requirement_set_proposal_sha256=binding.proposal_sha256,
        requirement_set_decision_sha256=binding.decision_sha256,
        requirement_set_binding_sha256=binding.binding_sha256,
        technical_scope_reason="exact set applied; offer still needs review",
        created_at=observed_at,
    )
    storage.store_sourcing_run(
        tenant,
        "operator-b",
        run,
        source_content,
        idempotency_key="offer-sourcing",
        request_sha256="e" * 64,
    )
    quote_source = Document()
    quote_source.add_paragraph("КП Q-1: HPE P-1, 1 шт, 1000 RUB без НДС")
    quote_output = BytesIO()
    quote_source.save(quote_output)
    quote_content = quote_output.getvalue()
    quote_id = f"quote-{uuid.uuid4()}"
    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    quote_document = storage.store_document(
        tenant,
        "operator-b",
        "quote.docx",
        media_type,
        DocumentRole.SUPPLIER_QUOTE,
        quote_content,
        extract_document(
            DocumentInput(
                document_id=quote_id,
                filename="quote.docx",
                declared_media_type=media_type,
                content=quote_content,
            )
        ),
        "offer-quote-document",
        hashlib.sha256(quote_content).hexdigest(),
        document_id=quote_id,
    )
    offer = storage.store_server_offer_proposal(
        tenant,
        "operator-b",
        offer_id=f"offer-{uuid.uuid4()}",
        case_id=case.case_id,
        case_version=case.version,
        requirement_set_binding_id=binding.binding_id,
        requirement_set_binding_sha256=binding.binding_sha256,
        sourcing_run_id=run.run_id,
        lead_id=lead.lead_id,
        source_document_id=quote_document.document_id,
        source_document_sha256=quote_document.sha256,
        bom_lines=(
            ServerOfferBomLineInput(
                ordinal=1,
                kind=ServerOfferPartKind.CHASSIS,
                manufacturer="HPE",
                sku="P-1",
                description="Exact configured chassis",
                quantity="1",
                unit="шт",
                line_total_ex_vat="1000.00",
                source_locator="paragraph 1",
            ),
        ),
        commercial_terms=ServerOfferCommercialTermsInput(
            quote_number="Q-1",
            supplier_id="supplier-1",
            supplier_name="Поставщик",
            amount_ex_vat="1000.00",
            nds_rate="0.20",
            issued_at=observed_at,
            valid_until=observed_at + timedelta(days=13),
            covers_full_scope=True,
            availability_terms="В наличии",
            delivery_terms="10 дней",
            payment_terms="После поставки",
            source_locator="paragraph 1",
        ),
        idempotency_key="server-offer-proposal",
        request_sha256="f" * 64,
    )
    assert storage.get_server_offer_proposal(tenant, offer.offer_id) == offer
    assert storage.get_server_offer_proposal(other, offer.offer_id) is None
    review = storage.get_server_offer_review(tenant, offer.offer_id)
    assert review is not None and review.decision is None
    offer_decision = storage.decide_server_offer_proposal(
        tenant,
        "reviewer-b",
        offer_id=offer.offer_id,
        proposal_sha256=offer.proposal_sha256,
        requirement_decisions=(
            ServerOfferRequirementDecision(
                requirement_ordinal=1,
                status=Verdict.PASS,
                reason="BOM row confirms the configured chassis",
                evidence=(
                    ServerOfferEvidenceRef(
                        kind=ServerOfferEvidenceKind.BOM_LINE,
                        bom_line_ordinal=1,
                    ),
                ),
            ),
        ),
        bom_scope_assessment=ServerOfferBomScopeAssessment.CONFIRMED_COMPLETE,
        commercial_terms_assessment=ServerOfferCommercialAssessment.CONFIRMED,
        reason="Every requirement row and the complete quote were reviewed",
        idempotency_key="server-offer-review",
        request_sha256="1" * 64,
    )
    assert offer_decision.status is Verdict.PASS
    loaded_review = storage.get_server_offer_review(tenant, offer.offer_id)
    assert loaded_review is not None and loaded_review.decision == offer_decision
    assert storage.get_server_offer_review(other, offer.offer_id) is None
    operating_costs = ServerOfferOperatingCostsInput(
        packaging="0", logistics="0", loading_unloading="0", storage="0",
        installation="0", payroll_travel="0", subcontracting="0", insurance="0",
        finance="0", bank_guarantee="0", platform_fees="0", taxes="0",
        certification="0", warranty="0", risk_reserve="0", penalty_reserve="0",
        other="0",
    )
    economics = storage.store_server_offer_economics(
        tenant,
        "operator-c",
        economics_id=f"economics-{uuid.uuid4()}",
        case_id=case.case_id,
        case_version=case.version,
        offer_id=offer.offer_id,
        offer_proposal_sha256=offer.proposal_sha256,
        offer_review_decision_sha256=offer_decision.decision_sha256,
        scenarios=tuple(
            ServerOfferScenarioInput(
                kind=kind,
                contract_revenue="1000000",
                costs=operating_costs,
                assumptions=(f"explicit {kind.value} assumptions",),
            )
            for kind in ScenarioKind
        ),
        idempotency_key="server-offer-economics",
        request_sha256="2" * 64,
    )
    assert economics.recommendation_status is Verdict.UNKNOWN
    loaded_economics = storage.get_server_offer_economics(
        tenant, case.case_id, case.version, offer.offer_id
    )
    assert loaded_economics is not None and loaded_economics.package == economics
    assert storage.get_server_offer_economics(
        other, case.case_id, case.version, offer.offer_id
    ) is None
    derived = storage.materialize_case_from_server_offer_economics(
        tenant,
        "operator-c",
        case_id=case.case_id,
        case_version=case.version,
        offer_id=offer.offer_id,
        economics_id=economics.economics_id,
        package_sha256=economics.package_sha256,
        idempotency_key="server-offer-case-materialization",
        request_sha256="3" * 64,
    )
    assert derived.version == case.version + 1
    assert derived.selected_quote is not None
    assert derived.selected_quote.quote_id == offer.offer_id
    assert derived.scenarios == economics.scenario_set
    assert derived.assessment is not None
    assert [(check.kind.value, check.status) for check in derived.assessment.checks] == [
        ("technical", Verdict.PASS)
    ]
    assert evaluate_case(
        derived, at=derived.assessment.assessed_at
    ).decision is RecommendationDecision.NEEDS_ANALYSIS
    assert storage.get_case(tenant, case.case_id, derived.version) == derived
    assert storage.get_case(other, case.case_id, derived.version) is None
    assert storage.materialize_case_from_server_offer_economics(
        tenant,
        "operator-c",
        case_id=case.case_id,
        case_version=case.version,
        offer_id=offer.offer_id,
        economics_id=economics.economics_id,
        package_sha256=economics.package_sha256,
        idempotency_key="server-offer-case-materialization",
        request_sha256="3" * 64,
    ) == derived
    gate_evidence = CaseHardGateEvidenceRef(
        document_id=document.document_id,
        document_sha256=document.sha256,
        locator="paragraph 1",
    )
    assessed = storage.materialize_case_hard_gate_assessment(
        tenant,
        "operator-d",
        case_id=case.case_id,
        case_version=derived.version,
        source_case_input_sha256=case_input_digest(derived),
        document_binding_id=document_binding.binding_id,
        manifest_evidence_sha256=document_binding.manifest_evidence_sha256,
        document_decision_sha256=document_binding.decision_sha256,
        decisions=tuple(
            CaseHardGateDecisionInput(
                kind=kind,
                status=Verdict.PASS,
                reason=f"exact evidence reviewed for {kind.value}",
                evidence=(gate_evidence,),
            )
            for kind in GateKind
            if kind not in {GateKind.TECHNICAL, GateKind.DOCUMENT}
        ),
        idempotency_key="case-hard-gate-assessment",
        request_sha256="4" * 64,
    )
    assert assessed.version == derived.version + 1
    assert assessed.assessment is not None
    assert tuple(check.kind for check in assessed.assessment.checks) == tuple(GateKind)
    assert assessed.assessment.checks[2] == derived.assessment.checks[0]
    assert storage.get_case(tenant, case.case_id, assessed.version) == assessed
    assert storage.get_case(other, case.case_id, assessed.version) is None
    assert storage.get_latest_evaluation(tenant, case.case_id, assessed.version) is None
    assert storage.materialize_case_hard_gate_assessment(
        tenant,
        "operator-d",
        case_id=case.case_id,
        case_version=derived.version,
        source_case_input_sha256=case_input_digest(derived),
        document_binding_id=document_binding.binding_id,
        manifest_evidence_sha256=document_binding.manifest_evidence_sha256,
        document_decision_sha256=document_binding.decision_sha256,
        decisions=tuple(
            CaseHardGateDecisionInput(
                kind=kind,
                status=Verdict.PASS,
                reason=f"exact evidence reviewed for {kind.value}",
                evidence=(gate_evidence,),
            )
            for kind in GateKind
            if kind not in {GateKind.TECHNICAL, GateKind.DOCUMENT}
        ),
        idempotency_key="case-hard-gate-assessment",
        request_sha256="4" * 64,
    ) == assessed

    with pytest.raises(IdempotencyConflict):
        storage.store_reviewed_requirement_set_proposal(
            tenant,
            "operator-a",
            set_id=f"set-{uuid.uuid4()}",
            source_document_id=document.document_id,
            source_document_sha256=document.sha256,
            declared_scope="changed",
            expected_entry_count=1,
            entries=(entry,),
            idempotency_key="requirement-set-prepare",
            request_sha256="d" * 64,
        )

    with storage.transaction(other) as connection:
        assert (
            connection.execute(
                select(func.count()).select_from(reviewed_requirement_set_proposals)
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                select(func.count()).select_from(reviewed_requirement_set_decisions)
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                select(func.count()).select_from(case_requirement_set_bindings)
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                select(func.count()).select_from(server_offer_proposals)
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                select(func.count()).select_from(server_offer_review_decisions)
            ).scalar_one()
            == 0
        )
        assert connection.execute(
            select(func.count()).select_from(server_offer_economics_packages)
        ).scalar_one() == 0

    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            update(reviewed_requirement_set_decisions)
            .where(reviewed_requirement_set_decisions.c.set_id == set_id)
            .values(status="FAIL")
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            delete(server_offer_economics_packages).where(
                server_offer_economics_packages.c.economics_id == economics.economics_id
            )
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            delete(case_requirement_set_bindings).where(
                case_requirement_set_bindings.c.binding_id == binding.binding_id
            )
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            delete(server_offer_proposals).where(
                server_offer_proposals.c.offer_id == offer.offer_id
            )
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            update(server_offer_review_decisions)
            .where(server_offer_review_decisions.c.offer_id == offer.offer_id)
            .values(status="FAIL")
        )
