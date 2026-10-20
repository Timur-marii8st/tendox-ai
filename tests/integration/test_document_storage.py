"""PostgreSQL checks for immutable tenant-scoped source documents."""

from __future__ import annotations

import hashlib
import json
import os
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime, timedelta
from io import BytesIO
from threading import Barrier, Event, Lock
from zipfile import ZIP_DEFLATED, ZipFile

import httpx
import pytest
from alembic import command
from alembic.config import Config
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from docx import Document
from docx.shared import Inches
from PIL import Image
from pypdf import PdfWriter
from sqlalchemy import create_engine, func, insert, select, text
from sqlalchemy.exc import DBAPIError

from procurement.access import AccessRole
from procurement.documents import (
    AiRequirementAnalysisReport,
    AiRequirementCandidateDecision,
    CandidateDisposition,
    CandidateKind,
    CompletenessStatus,
    ConstituentDecisionRef,
    DocumentInput,
    DocumentRole,
    DocumentSetCrossAssessment,
    DocumentSetEvidenceManifest,
    DocumentSetInventoryAssessment,
    DocumentSetMemberAssessment,
    DocumentSetMemberAttestation,
    ExtractionResult,
    ModelBrandCatalog,
    ModelBrandCatalogEntry,
    ModelBrandReviewProposal,
    ModelBrandReviewStatus,
    ModelBrandSourceEvidence,
    OcrEngineInput,
    analyze_requirements_with_model,
    build_document_search_entry,
    build_ocr_derivative,
    bundle_member_source_id,
    canonical_json_sha256,
    docx_media_source_id,
    extract_document,
    extract_document_bundle,
    extract_docx_media,
    validate_ai_requirement_report,
)
from procurement.documents.ai_requirement_review import (
    AiRequirementTransferBatchRunSnapshot,
    AiRequirementTransferRunSnapshot,
)
from procurement.documents.ai_requirements import (
    DEFAULT_OPENROUTER_MODEL,
    OpenRouterJsonResponse,
    prepare_ai_requirements_transfer,
)
from procurement.documents.ai_transfer import (
    build_persisted_ai_transfer_approval,
    build_persisted_ai_transfer_request,
)
from procurement.documents.ai_transfer_execution import (
    build_ai_transfer_dispatch_claim,
    build_ai_transfer_execution_admission,
)
from procurement.documents.ai_transfer_transport import (
    AiTransferTransportAttempt,
    AiTransferTransportOutcome,
    AiTransferTransportReason,
    ExactAiTransferHttpClient,
)
from procurement.documents.ai_transfer_worker import (
    AiTransferBatchCoordinator,
    AiTransferPreparedDispatch,
)
from procurement.storage import (
    AiRequirementReviewConflict,
    AiRequirementRunSourceMismatch,
    AiTransferConflict,
    AiTransferNotFound,
    AiTransferSourceMismatch,
    BundleMemberOcrConflict,
    CompletenessReviewConflict,
    CompletenessReviewSourceMismatch,
    DocumentBundleConflict,
    DocumentConflict,
    DocumentSearchIndexConflict,
    DocumentSetEvidenceConflict,
    DocumentSetEvidenceSourceMismatch,
    DocumentSetReviewConflict,
    IdempotencyConflict,
    ModelBrandCatalogConflict,
    ModelBrandReviewConflict,
    ModelBrandSourceArtifactConflict,
    ModelBrandSourceEvidenceMismatch,
    PostgresStorage,
    StorageError,
)
from procurement.storage.models import (
    MAX_DOCUMENT_BYTES,
    ai_transfer_approvals,
    ai_transfer_chunks,
    ai_transfer_dispatch_claims,
    ai_transfer_dispatch_receipts,
    ai_transfer_execution_admissions,
    ai_transfer_raw_response_blobs,
    ai_transfer_raw_responses,
    ai_transfer_requests,
    ai_transfer_requirement_batch_results,
    ai_transfer_requirement_chunk_rejections,
    ai_transfer_requirement_chunk_results,
    audit_events,
    document_ai_requirement_review_decisions,
    document_ai_requirement_runs,
    document_bundle_member_ocr_derivatives,
    document_bundles,
    document_completeness_review_decisions,
    document_completeness_reviews,
    document_docx_media_ocr_derivatives,
    document_search_indexes,
    document_set_completeness_decisions,
    document_set_completeness_member_attestations,
    document_set_evidence_manifests,
    document_set_evidence_members,
    documents,
    model_brand_catalogs,
    model_brand_review_decisions,
    model_brand_review_proposals,
    model_brand_source_artifacts,
    write_receipts,
)

SYNTHETIC_RAW_RESPONSE = b'{"synthetic":"provider response must never persist"}'


class _EphemeralTenantKekProvider:
    def __init__(self) -> None:
        self._key = bytes(range(32))

    def active_key_id(self, *, company_id: str) -> str:
        return f"tenant/{company_id}/test-kek/v1"

    def wrap_dek(
        self,
        *,
        company_id: str,
        key_id: str,
        nonce: bytes,
        dek: bytes,
        aad: bytes,
    ) -> bytes:
        del company_id, key_id
        return AESGCM(self._key).encrypt(nonce, dek, aad)

    def unwrap_dek(
        self,
        *,
        company_id: str,
        key_id: str,
        nonce: bytes,
        wrapped_dek: bytes,
        aad: bytes,
    ) -> bytes:
        del company_id, key_id
        return AESGCM(self._key).decrypt(nonce, wrapped_dek, aad)


def _store_transfer_source(storage: PostgresStorage, tenant: str, *, multi_chunk: bool = False):
    document_id = f"transfer-source-{uuid.uuid4()}"
    doc = Document()
    if multi_chunk:
        for number in range(1, 4):
            doc.add_paragraph(f"Раздел {number}. " + ("Насос мощностью 5 кВт. " * 700))
    else:
        doc.add_paragraph("Насос мощностью не менее 5 кВт.")
    output = BytesIO()
    doc.save(output)
    content = output.getvalue()
    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    extraction = extract_document(
        DocumentInput(
            document_id=document_id,
            filename="source.docx",
            declared_media_type=media_type,
            content=content,
        )
    )
    return storage.store_document(
        tenant,
        "preparer",
        "source.docx",
        media_type,
        DocumentRole.TENDER_REQUIREMENTS,
        content,
        extraction,
        "transfer-source",
        "a" * 64,
        document_id=document_id,
    )


def _claimed_ai_transfer(
    storage: PostgresStorage,
    tenant: str,
    *,
    worker_id: str = "receipt-worker",
    multi_chunk: bool = False,
):
    source = _store_transfer_source(storage, tenant, multi_chunk=multi_chunk)
    request = storage.store_ai_transfer_request(
        tenant,
        "preparer",
        source.document_id,
        DEFAULT_OPENROUTER_MODEL,
        f"receipt-prepare-{uuid.uuid4()}",
        "a" * 64,
    )
    storage.approve_ai_transfer_request(
        tenant,
        "approver",
        request.transfer_id,
        request.payload_sha256,
        f"receipt-approve-{uuid.uuid4()}",
        "b" * 64,
    )
    admission = storage.admit_ai_transfer_execution(
        tenant,
        "preparer",
        request.transfer_id,
        request.payload_sha256,
        f"receipt-admit-{uuid.uuid4()}",
        "c" * 64,
    )
    claim = storage.claim_first_ai_transfer_dispatch(tenant, worker_id, admission.execution_id)
    return admission, claim


def _successful_transport_attempt(claim) -> AiTransferTransportAttempt:
    raw_response = SYNTHETIC_RAW_RESPONSE
    assistant_content = b'{"items":[]}'
    started_at = claim.claimed_at
    finished_at = max(datetime.now(UTC), started_at)
    return AiTransferTransportAttempt(
        request_sha256=claim.request_sha256,
        outcome=AiTransferTransportOutcome.SUCCESS,
        reason=AiTransferTransportReason.VALIDATED_RESPONSE,
        http_status=200,
        capture_complete=True,
        captured_byte_count=len(raw_response),
        response_body_sha256=hashlib.sha256(raw_response).hexdigest(),
        provider_response_id="synthetic-provider-response",
        provider_model=DEFAULT_OPENROUTER_MODEL,
        assistant_content_sha256=hashlib.sha256(assistant_content).hexdigest(),
        assistant_content_byte_count=len(assistant_content),
        started_at=started_at,
        finished_at=finished_at,
    )


def test_next_ai_transfer_claim_is_ordered_committed_and_single_winner(
    raw_storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    admission, first_claim = _claimed_ai_transfer(raw_storage, tenant, multi_chunk=True)
    assert admission.chunk_count > 1
    first_attempt = _successful_transport_attempt(first_claim)
    raw_storage.persist_captured_ai_transfer_result(
        tenant,
        first_claim.claimed_by,
        admission.execution_id,
        first_attempt,
        SYNTHETIC_RAW_RESPONSE,
    )

    def claim_next():
        try:
            return raw_storage.claim_next_ai_transfer_dispatch(
                tenant, "next-worker", admission.execution_id
            )
        except AiTransferConflict:
            return None

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(lambda _: claim_next(), range(2)))
    claims = [claim for claim in results if claim is not None]

    assert len(claims) == 1
    second_claim = claims[0]
    assert second_claim.ordinal == 2
    assert second_claim.chunk_id == "chunk-0002"
    assert raw_storage.get_ai_transfer_execution(tenant, admission.execution_id) == (
        admission,
        second_claim,
        None,
    )
    with pytest.raises(AiTransferConflict, match="not terminal"):
        raw_storage.claim_next_ai_transfer_dispatch(tenant, "next-worker", admission.execution_id)
    with raw_storage.transaction(tenant) as connection:
        ordinals = connection.execute(
            select(ai_transfer_dispatch_claims.c.ordinal)
            .where(ai_transfer_dispatch_claims.c.execution_id == admission.execution_id)
            .order_by(ai_transfer_dispatch_claims.c.ordinal)
        ).scalars()
        assert tuple(ordinals) == (1, 2)


def test_next_ai_transfer_claim_requires_successful_prior_receipt(
    raw_storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    admission, first_claim = _claimed_ai_transfer(raw_storage, tenant, multi_chunk=True)
    rejected_body = b'{"error":"synthetic rejection"}'
    failed = AiTransferTransportAttempt(
        request_sha256=first_claim.request_sha256,
        outcome=AiTransferTransportOutcome.FAILED,
        reason=AiTransferTransportReason.HTTP_STATUS_REJECTED,
        http_status=400,
        capture_complete=True,
        captured_byte_count=len(rejected_body),
        response_body_sha256=hashlib.sha256(rejected_body).hexdigest(),
        provider_response_id=None,
        provider_model=None,
        assistant_content_sha256=None,
        assistant_content_byte_count=None,
        started_at=first_claim.claimed_at,
        finished_at=max(datetime.now(UTC), first_claim.claimed_at),
    )
    raw_storage.persist_captured_ai_transfer_result(
        tenant,
        first_claim.claimed_by,
        admission.execution_id,
        failed,
        rejected_body,
    )

    with pytest.raises(AiTransferConflict, match="next dispatch claim is invalid"):
        raw_storage.claim_next_ai_transfer_dispatch(tenant, "next-worker", admission.execution_id)


def test_batch_coordinator_commits_then_sends_one_exact_next_request(
    raw_storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    admission, first_claim = _claimed_ai_transfer(raw_storage, tenant, multi_chunk=True)
    first_attempt, first_raw_response = _valid_requirement_capture(first_claim)
    raw_storage.persist_captured_ai_transfer_result(
        tenant,
        first_claim.claimed_by,
        admission.execution_id,
        first_attempt,
        first_raw_response,
    )
    with raw_storage.transaction(tenant) as connection:
        second_chunk = connection.execute(
            select(ai_transfer_chunks.c.chunk).where(
                ai_transfer_chunks.c.transfer_id == admission.transfer_id,
                ai_transfer_chunks.c.ordinal == 2,
            )
        ).scalar_one()
    source_block_ids = second_chunk["source_block_ids"]
    assistant_content = json.dumps(
        {
            "requirements": [],
            "unresolved_block_ids": source_block_ids,
            "document_notes": [],
            "items": [],
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )
    response_body = json.dumps(
        {
            "id": "next-generation",
            "model": DEFAULT_OPENROUTER_MODEL,
            "choices": [
                {
                    "index": 0,
                    "finish_reason": "stop",
                    "message": {"role": "assistant", "content": assistant_content},
                }
            ],
        },
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    network_calls = 0

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal network_calls
        network_calls += 1
        with raw_storage.transaction(tenant) as connection:
            committed = connection.execute(
                select(func.count())
                .select_from(ai_transfer_dispatch_claims)
                .where(
                    ai_transfer_dispatch_claims.c.execution_id == admission.execution_id,
                    ai_transfer_dispatch_claims.c.ordinal == 2,
                )
            ).scalar_one()
        assert committed == 1
        return httpx.Response(200, stream=httpx.ByteStream(response_body), request=request)

    coordinated = AiTransferBatchCoordinator(
        storage=raw_storage,
        client=ExactAiTransferHttpClient(transport=httpx.MockTransport(handler)),
    ).run_once(
        company_id=tenant,
        worker_id="next-worker",
        execution_id=admission.execution_id,
        api_key="integration-secret",
    )

    assert network_calls == 1
    assert coordinated.action == "NEXT_CHUNK"
    assert coordinated.claim_ordinal == 2
    assert coordinated.worker_outcome is not None
    outcome = coordinated.worker_outcome
    assert outcome.receipt.outcome == "SUCCESS"
    assert outcome.requirement_batch_pending is True
    assert outcome.materialization_id is None
    assert outcome.raw_response_id is not None
    assert raw_storage.get_ai_transfer_execution(tenant, admission.execution_id) == (
        admission,
        outcome.claim,
        outcome.receipt,
    )


def test_concurrent_batch_coordinators_make_only_one_new_network_attempt(
    raw_storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    admission, first_claim = _claimed_ai_transfer(raw_storage, tenant, multi_chunk=True)
    first_attempt, first_raw_response = _valid_requirement_capture(first_claim)
    raw_storage.persist_captured_ai_transfer_result(
        tenant,
        first_claim.claimed_by,
        admission.execution_id,
        first_attempt,
        first_raw_response,
    )
    with raw_storage.transaction(tenant) as connection:
        second_chunk = connection.execute(
            select(ai_transfer_chunks.c.chunk).where(
                ai_transfer_chunks.c.transfer_id == admission.transfer_id,
                ai_transfer_chunks.c.ordinal == 2,
            )
        ).scalar_one()
    assistant_content = json.dumps(
        {
            "requirements": [],
            "unresolved_block_ids": second_chunk["source_block_ids"],
            "document_notes": [],
            "items": [],
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )
    response_body = json.dumps(
        {
            "id": "concurrent-next-generation",
            "model": DEFAULT_OPENROUTER_MODEL,
            "choices": [
                {
                    "index": 0,
                    "finish_reason": "stop",
                    "message": {"role": "assistant", "content": assistant_content},
                }
            ],
        },
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    initial_reads = Barrier(2)
    network_started = Event()
    losing_call_finished = Event()
    calls_lock = Lock()
    network_calls = 0

    class SynchronizedStorage:
        def get_ai_transfer_execution(self, company_id: str, execution_id: str):
            current = raw_storage.get_ai_transfer_execution(company_id, execution_id)
            initial_reads.wait(timeout=10)
            return current

        def __getattr__(self, name: str):
            return getattr(raw_storage, name)

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal network_calls
        with calls_lock:
            network_calls += 1
        network_started.set()
        assert losing_call_finished.wait(timeout=10)
        return httpx.Response(200, stream=httpx.ByteStream(response_body), request=request)

    def coordinate(worker_id: str):
        coordinator = AiTransferBatchCoordinator(
            storage=SynchronizedStorage(),
            client=ExactAiTransferHttpClient(transport=httpx.MockTransport(handler)),
        )
        try:
            return coordinator.run_once(
                company_id=tenant,
                worker_id=worker_id,
                execution_id=admission.execution_id,
                api_key="integration-secret",
            )
        except AiTransferConflict as error:
            assert network_started.wait(timeout=10)
            losing_call_finished.set()
            return error

    with ThreadPoolExecutor(max_workers=2) as executor:
        outcomes = list(executor.map(coordinate, ("concurrent-worker-a", "concurrent-worker-b")))

    assert network_calls == 1
    assert sum(isinstance(value, AiTransferConflict) for value in outcomes) == 1
    coordinated = next(value for value in outcomes if not isinstance(value, AiTransferConflict))
    assert coordinated.action == "NEXT_CHUNK"
    assert coordinated.claim_ordinal == 2
    with raw_storage.transaction(tenant) as connection:
        assert (
            connection.execute(
                select(func.count())
                .select_from(ai_transfer_dispatch_claims)
                .where(
                    ai_transfer_dispatch_claims.c.execution_id == admission.execution_id,
                    ai_transfer_dispatch_claims.c.ordinal == 2,
                )
            ).scalar_one()
            == 1
        )


def test_complete_transfer_batch_materializes_once_into_v3_review_run(
    raw_storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    admission, claim = _claimed_ai_transfer(raw_storage, tenant, multi_chunk=True)
    assert admission.chunk_count > 1
    latest_results = ()

    for ordinal in range(1, admission.chunk_count + 1):
        prepared = raw_storage.prepare_claimed_ai_transfer_dispatch(
            tenant,
            claim.claimed_by,
            admission.execution_id,
            claim,
        )
        assert isinstance(prepared, AiTransferPreparedDispatch)
        content = json.dumps(
            {
                "requirements": [],
                "unresolved_block_ids": list(prepared.request.source_block_ids),
                "document_notes": [],
                "items": [],
            },
            ensure_ascii=False,
            separators=(",", ":"),
        )
        raw_response = json.dumps(
            {
                "id": f"batch-generation-{ordinal}",
                "model": DEFAULT_OPENROUTER_MODEL,
                "choices": [
                    {
                        "index": 0,
                        "finish_reason": "stop",
                        "message": {"role": "assistant", "content": content},
                    }
                ],
            },
            ensure_ascii=False,
            separators=(",", ":"),
        ).encode("utf-8")
        started_at = max(datetime.now(UTC), claim.claimed_at)
        attempt = AiTransferTransportAttempt(
            request_sha256=claim.request_sha256,
            outcome=AiTransferTransportOutcome.SUCCESS,
            reason=AiTransferTransportReason.VALIDATED_RESPONSE,
            http_status=200,
            capture_complete=True,
            captured_byte_count=len(raw_response),
            response_body_sha256=hashlib.sha256(raw_response).hexdigest(),
            provider_response_id=f"batch-generation-{ordinal}",
            provider_model=DEFAULT_OPENROUTER_MODEL,
            assistant_content_sha256=hashlib.sha256(content.encode("utf-8")).hexdigest(),
            assistant_content_byte_count=len(content.encode("utf-8")),
            started_at=started_at,
            finished_at=started_at,
        )
        raw_storage.persist_captured_ai_transfer_result(
            tenant,
            claim.claimed_by,
            admission.execution_id,
            attempt,
            raw_response,
        )
        latest_results, record, rejection = raw_storage.materialize_ai_transfer_requirement_batch(
            tenant,
            "batch-materializer",
            admission.execution_id,
        )
        assert tuple(result.ordinal for result in latest_results) == tuple(range(1, ordinal + 1))
        assert rejection is None
        if ordinal < admission.chunk_count:
            assert record is None
            claim = raw_storage.claim_next_ai_transfer_dispatch(
                tenant,
                f"next-worker-{ordinal + 1}",
                admission.execution_id,
            )
        else:
            assert record is not None

    assert record is not None
    assert isinstance(record.snapshot, AiRequirementTransferBatchRunSnapshot)
    assert record.snapshot.completeness_status is CompletenessStatus.UNKNOWN
    assert record.snapshot.result.chunk_results == latest_results
    serialized = record.snapshot.model_dump_json()
    assert "response_content" not in serialized
    assert "ciphertext" not in serialized
    replay = raw_storage.materialize_ai_transfer_requirement_batch(
        tenant,
        "recovery-materializer",
        admission.execution_id,
    )
    assert replay == (latest_results, record, None)
    network_calls = 0

    def forbidden_network(_: httpx.Request) -> httpx.Response:
        nonlocal network_calls
        network_calls += 1
        raise AssertionError("terminal recovery must not access the network")

    recovered = AiTransferBatchCoordinator(
        storage=raw_storage,
        client=ExactAiTransferHttpClient(transport=httpx.MockTransport(forbidden_network)),
    ).run_once(
        company_id=tenant,
        worker_id="coordinator-recovery",
        execution_id=admission.execution_id,
        api_key="must-not-be-used",
    )
    assert network_calls == 0
    assert recovered.action == "RECOVERY"
    assert recovered.status == "MATERIALIZED"
    assert recovered.network_attempted is False
    assert recovered.run_id == record.snapshot.run_id
    with raw_storage.transaction(tenant) as connection:
        assert (
            connection.execute(
                select(func.count()).select_from(ai_transfer_requirement_chunk_results)
            ).scalar_one()
            == admission.chunk_count
        )
        assert (
            connection.execute(
                select(func.count()).select_from(ai_transfer_requirement_batch_results)
            ).scalar_one()
            == 1
        )
    other = f"tenant-{uuid.uuid4()}"
    with raw_storage.transaction(other) as connection:
        assert (
            connection.execute(
                select(func.count()).select_from(ai_transfer_requirement_chunk_results)
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                select(func.count()).select_from(ai_transfer_requirement_batch_results)
            ).scalar_one()
            == 0
        )
    with pytest.raises(DBAPIError), raw_storage.transaction(tenant) as connection:
        connection.execute(
            ai_transfer_requirement_chunk_results.update().values(materialized_by="forged-actor")
        )
    with pytest.raises(DBAPIError), raw_storage.transaction(tenant) as connection:
        connection.execute(ai_transfer_requirement_batch_results.delete())


def test_coordinator_recovers_single_capture_after_materializer_crash_without_network(
    raw_storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    admission, claim = _claimed_ai_transfer(raw_storage, tenant)
    attempt, raw_response = _valid_requirement_capture(claim)
    receipt, raw_metadata = raw_storage.persist_captured_ai_transfer_result(
        tenant,
        claim.claimed_by,
        admission.execution_id,
        attempt,
        raw_response,
    )
    assert receipt.outcome == "SUCCESS"
    assert raw_metadata.execution_id == admission.execution_id
    network_calls = 0

    def forbidden_network(_: httpx.Request) -> httpx.Response:
        nonlocal network_calls
        network_calls += 1
        raise AssertionError("post-capture recovery must not access the network")

    coordinator = AiTransferBatchCoordinator(
        storage=raw_storage,
        client=ExactAiTransferHttpClient(transport=httpx.MockTransport(forbidden_network)),
    )
    recovered = coordinator.run_once(
        company_id=tenant,
        worker_id="recovery-worker-a",
        execution_id=admission.execution_id,
        api_key="must-not-be-used",
    )
    replay = coordinator.run_once(
        company_id=tenant,
        worker_id="recovery-worker-b",
        execution_id=admission.execution_id,
        api_key="must-not-be-used",
    )

    assert network_calls == 0
    assert recovered.action == "RECOVERY"
    assert recovered.status == "MATERIALIZED"
    assert recovered.network_attempted is False
    assert recovered.materialization_id == replay.materialization_id
    assert recovered.run_id == replay.run_id


def test_coordinator_recovers_semantic_rejection_after_capture_without_network(
    raw_storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    admission, claim = _claimed_ai_transfer(raw_storage, tenant)
    attempt, raw_response = _valid_requirement_capture(
        claim,
        wire={
            "requirements": [],
            "unresolved_block_ids": [],
            "document_notes": [],
            "items": [
                {
                    "item_name": {
                        "value": "Вымышленный товар",
                        "source_block_ids": ["b000001"],
                    },
                    "okpd2_codes": [],
                    "ktru_codes": [],
                    "unit": None,
                    "quantity": None,
                    "uncertainty": None,
                }
            ],
        },
    )
    raw_storage.persist_captured_ai_transfer_result(
        tenant,
        claim.claimed_by,
        admission.execution_id,
        attempt,
        raw_response,
    )
    network_calls = 0

    def forbidden_network(_: httpx.Request) -> httpx.Response:
        nonlocal network_calls
        network_calls += 1
        raise AssertionError("semantic recovery must not access the network")

    coordinator = AiTransferBatchCoordinator(
        storage=raw_storage,
        client=ExactAiTransferHttpClient(transport=httpx.MockTransport(forbidden_network)),
    )
    rejected = coordinator.run_once(
        company_id=tenant,
        worker_id="rejection-recovery-a",
        execution_id=admission.execution_id,
        api_key="must-not-be-used",
    )
    replay = coordinator.run_once(
        company_id=tenant,
        worker_id="rejection-recovery-b",
        execution_id=admission.execution_id,
        api_key="must-not-be-used",
    )

    assert network_calls == 0
    assert rejected.action == "RECOVERY"
    assert rejected.status == "REJECTED"
    assert rejected.network_attempted is False
    assert rejected.run_id is None
    assert rejected.materialization_id == replay.materialization_id


def test_invalid_multi_chunk_candidate_records_terminal_rejection_without_network_replay(
    raw_storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    admission, claim = _claimed_ai_transfer(raw_storage, tenant, multi_chunk=True)
    assert admission.chunk_count > 1
    attempt, raw_response = _valid_requirement_capture(
        claim,
        wire={
            "requirements": [],
            "unresolved_block_ids": [],
            "document_notes": [],
            "items": [
                {
                    "item_name": {
                        "value": "Вымышленный товар",
                        "source_block_ids": ["b000001"],
                    },
                    "okpd2_codes": [],
                    "ktru_codes": [],
                    "unit": None,
                    "quantity": None,
                    "uncertainty": None,
                }
            ],
        },
    )
    raw_storage.persist_captured_ai_transfer_result(
        tenant,
        claim.claimed_by,
        admission.execution_id,
        attempt,
        raw_response,
    )

    first = raw_storage.materialize_ai_transfer_requirement_batch(
        tenant,
        "batch-materializer",
        admission.execution_id,
    )
    replay = raw_storage.materialize_ai_transfer_requirement_batch(
        tenant,
        "recovery-materializer",
        admission.execution_id,
    )

    results, record, rejection = first
    assert results == ()
    assert record is None
    assert rejection is not None
    assert rejection.reason == "INVALID_CANDIDATE_EVIDENCE"
    assert rejection.ordinal == 1
    assert replay == first
    execution = raw_storage.get_ai_transfer_execution(tenant, admission.execution_id)
    assert execution is not None and execution[2] is not None
    assert execution[2].outcome == "SUCCESS"
    with pytest.raises(AiTransferConflict):
        raw_storage.claim_next_ai_transfer_dispatch(
            tenant,
            "must-not-dispatch",
            admission.execution_id,
        )
    with raw_storage.transaction(tenant) as connection:
        assert (
            connection.execute(
                select(func.count()).select_from(ai_transfer_requirement_chunk_rejections)
            ).scalar_one()
            == 1
        )
        assert (
            connection.execute(
                select(func.count()).select_from(ai_transfer_requirement_chunk_results)
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                select(func.count()).select_from(ai_transfer_requirement_batch_results)
            ).scalar_one()
            == 0
        )
    with pytest.raises(DBAPIError), raw_storage.transaction(tenant) as connection:
        connection.execute(ai_transfer_requirement_chunk_rejections.delete())


def _valid_requirement_capture(
    claim,
    *,
    wire: dict[str, object] | None = None,
) -> tuple[AiTransferTransportAttempt, bytes]:
    content = json.dumps(
        wire
        or {
            "requirements": [
                {
                    "item_name": "Насос",
                    "property": "мощность",
                    "operator": ">=",
                    "value": "5",
                    "unit": "кВт",
                    "mandatory": True,
                    "source_block_ids": ["b000001"],
                    "uncertainty": None,
                }
            ],
            "unresolved_block_ids": [],
            "document_notes": [],
            "items": [],
        },
        ensure_ascii=False,
        separators=(",", ":"),
    )
    raw_response = json.dumps(
        {
            "id": "materialization-generation-1",
            "model": DEFAULT_OPENROUTER_MODEL,
            "choices": [
                {
                    "index": 0,
                    "finish_reason": "stop",
                    "message": {"role": "assistant", "content": content},
                }
            ],
        },
        ensure_ascii=False,
        separators=(",", ":"),
    ).encode("utf-8")
    content_bytes = content.encode("utf-8")
    started_at = claim.claimed_at
    finished_at = max(datetime.now(UTC), started_at)
    return (
        AiTransferTransportAttempt(
            request_sha256=claim.request_sha256,
            outcome=AiTransferTransportOutcome.SUCCESS,
            reason=AiTransferTransportReason.VALIDATED_RESPONSE,
            http_status=200,
            capture_complete=True,
            captured_byte_count=len(raw_response),
            response_body_sha256=hashlib.sha256(raw_response).hexdigest(),
            provider_response_id="materialization-generation-1",
            provider_model=DEFAULT_OPENROUTER_MODEL,
            assistant_content_sha256=hashlib.sha256(content_bytes).hexdigest(),
            assistant_content_byte_count=len(content_bytes),
            started_at=started_at,
            finished_at=finished_at,
        ),
        raw_response,
    )


def _forged_receipt_values(receipt, claim) -> tuple[dict[str, object], dict[str, object]]:
    """Return matching SQL projections and receipt JSON for a different claim."""
    receipt_json = receipt.model_dump(mode="json")
    recorded_at = datetime.now(UTC)
    started_at = max(claim.claimed_at, recorded_at - timedelta(seconds=2))
    finished_at = max(started_at, recorded_at - timedelta(seconds=1))
    receipt_json.update(
        company_id=claim.company_id,
        transfer_id=claim.transfer_id,
        execution_id=claim.execution_id,
        ordinal=claim.ordinal,
        claim_id=claim.claim_id,
        chunk_id=claim.chunk_id,
        claim_payload_sha256=claim.payload_sha256,
        receipt_id=str(uuid.uuid4()),
        request_sha256=claim.request_sha256,
        recorded_by=claim.claimed_by,
        started_at=started_at.isoformat(),
        finished_at=finished_at.isoformat(),
        recorded_at=recorded_at.isoformat(),
        payload_sha256="f" * 64,
    )
    values = {
        column.name: receipt_json[column.name]
        for column in ai_transfer_dispatch_receipts.columns
        if column.name != "receipt"
    }
    values.update(
        started_at=started_at,
        finished_at=finished_at,
        recorded_at=recorded_at,
        receipt=receipt_json,
    )
    return values, receipt_json


def test_ai_transfer_storage_atomic_independent_replay_and_tenant_scope(
    storage: PostgresStorage,
) -> None:
    tenant, other = f"tenant-{uuid.uuid4()}", f"tenant-{uuid.uuid4()}"
    source = _store_transfer_source(storage, tenant)

    def prepare():
        return storage.store_ai_transfer_request(
            tenant, "preparer", source.document_id, DEFAULT_OPENROUTER_MODEL, "prepare", "b" * 64
        )

    with ThreadPoolExecutor(max_workers=2) as executor:
        first, replay = list(executor.map(lambda _: prepare(), range(2)))
    assert first == replay
    assert storage.get_ai_transfer_request(other, first.transfer_id) is None
    with pytest.raises(AiTransferConflict):
        storage.store_ai_transfer_request(
            tenant,
            "different-actor",
            source.document_id,
            DEFAULT_OPENROUTER_MODEL,
            "prepare",
            "b" * 64,
        )
    with pytest.raises(IdempotencyConflict):
        storage.store_ai_transfer_request(
            tenant, "preparer", source.document_id, DEFAULT_OPENROUTER_MODEL, "prepare", "c" * 64
        )
    with pytest.raises(AiTransferConflict):
        storage.approve_ai_transfer_request(
            tenant, "preparer", first.transfer_id, first.payload_sha256, "self", "d" * 64
        )
    with pytest.raises(AiTransferConflict):
        storage.approve_ai_transfer_request(
            tenant, "approver", first.transfer_id, "0" * 64, "stale", "e" * 64
        )

    def approve():
        return storage.approve_ai_transfer_request(
            tenant, "approver", first.transfer_id, first.payload_sha256, "approve", "f" * 64
        )

    with ThreadPoolExecutor(max_workers=2) as executor:
        approval, approval_replay = list(executor.map(lambda _: approve(), range(2)))
    assert approval == approval_replay
    assert approval.expires_at - approval.approved_at == timedelta(minutes=15)
    assert prepare() == first
    assert storage.get_ai_transfer_request(tenant, first.transfer_id) == (first, approval)
    with pytest.raises(AiTransferConflict):
        storage.approve_ai_transfer_request(
            tenant, "approver", first.transfer_id, first.payload_sha256, "second", "1" * 64
        )

    def admit():
        return storage.admit_ai_transfer_execution(
            tenant, "preparer", first.transfer_id, first.payload_sha256, "admit", "3" * 64
        )

    with ThreadPoolExecutor(max_workers=2) as executor:
        admission, admission_replay = list(executor.map(lambda _: admit(), range(2)))
    assert admission == admission_replay
    assert admission.status == "ADMITTED"
    assert storage.get_ai_transfer_execution(other, admission.execution_id) is None
    with pytest.raises(AiTransferConflict):
        storage.admit_ai_transfer_execution(
            tenant,
            "different-actor",
            first.transfer_id,
            first.payload_sha256,
            "admit-other",
            "4" * 64,
        )
    with pytest.raises(AiTransferConflict):
        storage.admit_ai_transfer_execution(
            tenant, "preparer", first.transfer_id, first.payload_sha256, "admit-again", "5" * 64
        )

    def claim_once():
        try:
            return storage.claim_first_ai_transfer_dispatch(
                tenant, "worker-1", admission.execution_id
            )
        except AiTransferConflict:
            return None

    with ThreadPoolExecutor(max_workers=2) as executor:
        claim_results = list(executor.map(lambda _: claim_once(), range(2)))
    claims = [item for item in claim_results if item is not None]
    assert len(claims) == 1
    claim = claims[0]
    assert claim.ordinal == 1
    assert claim.delivery_state == "DELIVERY_UNKNOWN"
    assert storage.get_ai_transfer_execution(tenant, admission.execution_id) == (
        admission,
        claim,
        None,
    )
    with pytest.raises(AiTransferConflict):
        storage.claim_first_ai_transfer_dispatch(tenant, "worker-1", admission.execution_id)

    with storage.transaction(tenant) as connection:
        assert (
            connection.execute(select(func.count()).select_from(ai_transfer_requests)).scalar() == 1
        )
        assert (
            connection.execute(select(func.count()).select_from(ai_transfer_approvals)).scalar()
            == 1
        )
        assert (
            connection.execute(
                select(func.count()).select_from(ai_transfer_execution_admissions)
            ).scalar()
            == 1
        )
        assert (
            connection.execute(
                select(func.count()).select_from(ai_transfer_dispatch_claims)
            ).scalar()
            == 1
        )
        events = connection.execute(
            select(audit_events.c.event_type, audit_events.c.basis).where(
                audit_events.c.object_type == "ai_transfer_request"
            )
        ).all()
        assert {row.event_type for row in events} == {
            "document.ai_transfer_prepared",
            "document.ai_transfer_approved",
        }
        execution_events = connection.execute(
            select(audit_events.c.event_type, audit_events.c.basis).where(
                audit_events.c.object_type == "ai_transfer_execution"
            )
        ).all()
        assert {row.event_type for row in execution_events} == {
            "document.ai_transfer_execution_admitted",
            "document.ai_transfer_dispatch_claimed",
        }
        assert "request_body_json" not in json.dumps([row.basis for row in events])
        assert "request_body_json" not in json.dumps([row.basis for row in execution_events])
        assert (
            connection.execute(
                select(func.count())
                .select_from(write_receipts)
                .where(
                    write_receipts.c.operation.in_(
                        [
                            "store_ai_transfer_request",
                            "approve_ai_transfer_request",
                            "admit_ai_transfer_execution",
                        ]
                    )
                )
            ).scalar()
            == 3
        )
    with storage.transaction(other) as connection:
        assert (
            connection.execute(select(func.count()).select_from(ai_transfer_requests)).scalar() == 0
        )
    for table in (
        ai_transfer_requests,
        ai_transfer_chunks,
        ai_transfer_approvals,
        ai_transfer_execution_admissions,
        ai_transfer_dispatch_claims,
        ai_transfer_dispatch_receipts,
    ):
        with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
            connection.execute(table.delete())


def test_ai_transfer_sql_rejects_incomplete_projection_and_invalid_approvals(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    source = _store_transfer_source(storage, tenant)
    manifest = prepare_ai_requirements_transfer(
        source.extraction,
        company_id=tenant,
        source_document_id=source.document_id,
        source_sha256=source.sha256,
    )
    for mutation in ("missing", "wrong_body", "wrong_order", "projection"):
        prepared = build_persisted_ai_transfer_request(
            manifest,
            transfer_id=str(uuid.uuid4()),
            prepared_by="preparer",
            prepared_at=datetime.now(UTC),
        )
        payload = prepared.model_dump(mode="json")
        if mutation == "projection":
            payload["chunks"][0]["source_block_ids"] = ["b999999"]
        with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
            connection.execute(
                insert(ai_transfer_requests).values(
                    company_id=tenant,
                    transfer_id=prepared.transfer_id,
                    source_document_id=source.document_id,
                    source_sha256=source.sha256,
                    payload_sha256=prepared.payload_sha256,
                    chunk_count=len(prepared.chunks),
                    request=payload,
                    prepared_by="preparer",
                    created_at=prepared.prepared_at,
                )
            )
            if mutation != "missing":
                for ordinal, chunk in enumerate(prepared.chunks, 1):
                    connection.execute(
                        insert(ai_transfer_chunks).values(
                            company_id=tenant,
                            transfer_id=prepared.transfer_id,
                            payload_sha256=prepared.payload_sha256,
                            ordinal=ordinal + (1 if mutation == "wrong_order" else 0),
                            chunk_id=chunk.chunk_id,
                            request_body_json=(
                                "{}" if mutation == "wrong_body" else chunk.request_body_json
                            ),
                            request_sha256=chunk.request_sha256,
                            chunk=chunk.model_dump(mode="json"),
                        )
                    )
            connection.execute(text("SET CONSTRAINTS ALL IMMEDIATE"))
    prepared = storage.store_ai_transfer_request(
        tenant, "preparer", source.document_id, DEFAULT_OPENROUTER_MODEL, "valid", "2" * 64
    )
    valid = build_persisted_ai_transfer_approval(
        prepared, approved_by="approver", approved_at=datetime.now(UTC)
    )
    for mutation in ("self", "predated", "expiry", "digest"):
        payload = valid.model_dump(mode="json")
        actor = "preparer" if mutation == "self" else valid.approved_by
        approved_at = (
            prepared.prepared_at - timedelta(seconds=1)
            if mutation == "predated"
            else valid.approved_at
        )
        expires_at = approved_at + timedelta(minutes=16 if mutation == "expiry" else 15)
        digest = "0" * 64 if mutation == "digest" else prepared.payload_sha256
        payload.update(
            approved_by=actor,
            approved_at=approved_at.isoformat(),
            expires_at=expires_at.isoformat(),
            request_payload_sha256=digest,
        )
        with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
            connection.execute(
                insert(ai_transfer_approvals).values(
                    company_id=tenant,
                    transfer_id=prepared.transfer_id,
                    request_payload_sha256=digest,
                    approval=payload,
                    approved_by=actor,
                    approved_at=approved_at,
                    expires_at=expires_at,
                )
            )
    assert storage.get_ai_transfer_request(tenant, prepared.transfer_id) == (prepared, None)


def test_ai_transfer_execution_sql_rejects_forged_time_expiry_and_claims(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    source = _store_transfer_source(storage, tenant)
    request = storage.store_ai_transfer_request(
        tenant, "preparer", source.document_id, DEFAULT_OPENROUTER_MODEL, "prepare", "6" * 64
    )
    approval = storage.approve_ai_transfer_request(
        tenant, "approver", request.transfer_id, request.payload_sha256, "approve", "7" * 64
    )
    valid_admission = build_ai_transfer_execution_admission(
        request,
        approval,
        company_id=tenant,
        transfer_id=request.transfer_id,
        execution_id=str(uuid.uuid4()),
        admitted_by="preparer",
        admitted_at=datetime.now(UTC),
        existing_admission=None,
    )

    for mutation in ("wrong_actor", "backdated", "projection"):
        payload = valid_admission.model_dump(mode="json")
        execution_id = str(uuid.uuid4())
        admitted_by = "intruder" if mutation == "wrong_actor" else "preparer"
        admitted_at = (
            valid_admission.admitted_at - timedelta(seconds=3)
            if mutation == "backdated"
            else valid_admission.admitted_at
        )
        payload["admitted_by"] = admitted_by
        payload["admitted_at"] = admitted_at.isoformat()
        payload["execution_id"] = execution_id
        if mutation == "projection":
            payload["chunk_count"] = valid_admission.chunk_count + 1
        with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
            connection.execute(
                insert(ai_transfer_execution_admissions).values(
                    company_id=tenant,
                    execution_id=execution_id,
                    transfer_id=request.transfer_id,
                    request_payload_sha256=request.payload_sha256,
                    approved_by=approval.approved_by,
                    approved_at=approval.approved_at,
                    approval_expires_at=approval.expires_at,
                    chunk_count=payload["chunk_count"],
                    payload_sha256=valid_admission.payload_sha256,
                    admission=payload,
                    admitted_by=admitted_by,
                    admitted_at=admitted_at,
                )
            )

    admission = storage.admit_ai_transfer_execution(
        tenant, "preparer", request.transfer_id, request.payload_sha256, "admit", "8" * 64
    )
    valid_claim = build_ai_transfer_dispatch_claim(
        request,
        approval,
        admission,
        claim_id=str(uuid.uuid4()),
        claimed_by="worker",
        claimed_at=datetime.now(UTC),
        existing_claims=(),
        existing_receipts=(),
    )
    for mutation in ("ordinal", "backdated", "admission_digest", "projection"):
        payload = valid_claim.model_dump(mode="json")
        claim_id = str(uuid.uuid4())
        ordinal = 2 if mutation == "ordinal" else 1
        claimed_at = (
            valid_claim.claimed_at - timedelta(seconds=3)
            if mutation == "backdated"
            else valid_claim.claimed_at
        )
        admission_digest = (
            "0" * 64 if mutation == "admission_digest" else valid_claim.admission_payload_sha256
        )
        payload["ordinal"] = ordinal
        payload["claim_id"] = claim_id
        payload["claimed_at"] = claimed_at.isoformat()
        payload["admission_payload_sha256"] = admission_digest
        if mutation == "projection":
            payload["attempt"] = 2
        with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
            connection.execute(
                insert(ai_transfer_dispatch_claims).values(
                    company_id=tenant,
                    claim_id=claim_id,
                    execution_id=admission.execution_id,
                    transfer_id=request.transfer_id,
                    request_payload_sha256=request.payload_sha256,
                    admission_payload_sha256=admission_digest,
                    ordinal=ordinal,
                    chunk_id=valid_claim.chunk_id,
                    request_sha256=valid_claim.request_sha256,
                    claim=payload,
                    attempt=payload["attempt"],
                    claim_status=valid_claim.claim_status,
                    delivery_state=valid_claim.delivery_state,
                    claimed_by=valid_claim.claimed_by,
                    claimed_at=claimed_at,
                    payload_sha256=valid_claim.payload_sha256,
                )
            )

    past_manifest = prepare_ai_requirements_transfer(
        source.extraction,
        company_id=tenant,
        source_document_id=source.document_id,
        source_sha256=source.sha256,
    )
    past_request = build_persisted_ai_transfer_request(
        past_manifest,
        transfer_id=str(uuid.uuid4()),
        prepared_by="preparer",
        prepared_at=datetime.now(UTC) - timedelta(minutes=20),
    )
    with storage.transaction(tenant) as connection:
        connection.execute(
            insert(ai_transfer_requests).values(
                company_id=tenant,
                transfer_id=past_request.transfer_id,
                source_document_id=source.document_id,
                source_sha256=source.sha256,
                payload_sha256=past_request.payload_sha256,
                chunk_count=len(past_request.chunks),
                request=past_request.model_dump(mode="json"),
                prepared_by=past_request.prepared_by,
                created_at=past_request.prepared_at,
            )
        )
        connection.execute(
            insert(ai_transfer_chunks),
            [
                {
                    "company_id": tenant,
                    "transfer_id": past_request.transfer_id,
                    "payload_sha256": past_request.payload_sha256,
                    "ordinal": ordinal,
                    "chunk_id": chunk.chunk_id,
                    "request_body_json": chunk.request_body_json,
                    "request_sha256": chunk.request_sha256,
                    "chunk": chunk.model_dump(mode="json"),
                }
                for ordinal, chunk in enumerate(past_request.chunks, 1)
            ],
        )
        connection.execute(text("SET CONSTRAINTS ALL IMMEDIATE"))
    past_approval = build_persisted_ai_transfer_approval(
        past_request,
        approved_by="approver",
        approved_at=past_request.prepared_at + timedelta(minutes=1),
    )
    with storage.transaction(tenant) as connection:
        connection.execute(
            insert(ai_transfer_approvals).values(
                company_id=tenant,
                transfer_id=past_request.transfer_id,
                request_payload_sha256=past_request.payload_sha256,
                approval=past_approval.model_dump(mode="json"),
                approved_by=past_approval.approved_by,
                approved_at=past_approval.approved_at,
                expires_at=past_approval.expires_at,
            )
        )
    historical_admission = build_ai_transfer_execution_admission(
        past_request,
        past_approval,
        company_id=tenant,
        transfer_id=past_request.transfer_id,
        execution_id=str(uuid.uuid4()),
        admitted_by="preparer",
        admitted_at=past_approval.approved_at + timedelta(minutes=1),
        existing_admission=None,
    )
    with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
        connection.execute(
            insert(ai_transfer_execution_admissions).values(
                company_id=tenant,
                execution_id=historical_admission.execution_id,
                transfer_id=past_request.transfer_id,
                request_payload_sha256=past_request.payload_sha256,
                approved_by=past_approval.approved_by,
                approved_at=past_approval.approved_at,
                approval_expires_at=past_approval.expires_at,
                chunk_count=historical_admission.chunk_count,
                payload_sha256=historical_admission.payload_sha256,
                admission=historical_admission.model_dump(mode="json"),
                admitted_by=historical_admission.admitted_by,
                admitted_at=historical_admission.admitted_at,
            )
        )


def test_ai_transfer_dispatch_receipt_is_atomic_safe_and_tenant_scoped(
    raw_storage: PostgresStorage,
) -> None:
    storage = raw_storage
    tenant, other = f"tenant-{uuid.uuid4()}", f"tenant-{uuid.uuid4()}"
    admission, claim = _claimed_ai_transfer(storage, tenant)
    attempt = _successful_transport_attempt(claim)
    prepared = storage.prepare_claimed_ai_transfer_dispatch(
        tenant,
        claim.claimed_by,
        admission.execution_id,
        claim,
    )
    assert isinstance(prepared, AiTransferPreparedDispatch)
    assert prepared.claim == claim
    assert prepared.request.request_sha256 == claim.request_sha256
    assert prepared.request.source_block_ids
    assert prepared.checked_at < prepared.approval_expires_at

    with pytest.raises(AiTransferConflict):
        storage.record_ai_transfer_dispatch_receipt(
            tenant, "other-worker", admission.execution_id, attempt
        )
    with pytest.raises(AiTransferConflict):
        storage.record_ai_transfer_dispatch_receipt(
            tenant,
            claim.claimed_by,
            admission.execution_id,
            attempt.model_copy(update={"request_sha256": "0" * 64}),
        )

    def record():
        try:
            return storage.persist_captured_ai_transfer_result(
                tenant,
                claim.claimed_by,
                admission.execution_id,
                attempt,
                SYNTHETIC_RAW_RESPONSE,
            )
        except AiTransferConflict:
            return None

    with ThreadPoolExecutor(max_workers=2) as executor:
        receipt_results = list(executor.map(lambda _: record(), range(2)))
    results = [item for item in receipt_results if item is not None]
    assert len(results) == 1
    receipt, raw_metadata = results[0]
    assert receipt.outcome == "SUCCESS"
    assert receipt.request_sha256 == claim.request_sha256
    assert storage.get_ai_transfer_execution(tenant, admission.execution_id) == (
        admission,
        claim,
        receipt,
    )
    assert storage.get_ai_transfer_execution(other, admission.execution_id) is None
    assert (
        storage.get_ai_transfer_raw_metadata_for_execution(tenant, admission.execution_id)
        == raw_metadata
    )
    assert storage.get_ai_transfer_raw_metadata_for_execution(other, admission.execution_id) is None
    with pytest.raises(AiTransferConflict):
        storage.persist_captured_ai_transfer_result(
            tenant,
            claim.claimed_by,
            admission.execution_id,
            attempt,
            SYNTHETIC_RAW_RESPONSE,
        )

    plaintext, read_metadata = storage.read_ai_transfer_raw_response(
        tenant,
        "reviewer",
        AccessRole.AI_REQUIREMENTS_REVIEWER,
        raw_metadata.raw_response_id,
    )
    assert plaintext == SYNTHETIC_RAW_RESPONSE
    assert read_metadata == raw_metadata
    with pytest.raises(AiTransferSourceMismatch):
        storage.read_ai_transfer_raw_response(
            tenant,
            "operator",
            AccessRole.OPERATOR,
            raw_metadata.raw_response_id,
        )
    with pytest.raises(AiTransferNotFound):
        storage.read_ai_transfer_raw_response(
            other,
            "auditor",
            AccessRole.AUDITOR,
            raw_metadata.raw_response_id,
        )

    with storage.transaction(tenant) as connection:
        row = connection.execute(
            select(
                ai_transfer_dispatch_receipts.c.receipt,
                ai_transfer_dispatch_receipts.c.response_body_sha256,
                ai_transfer_dispatch_receipts.c.assistant_content_sha256,
            )
        ).one()
        assert row.response_body_sha256 == attempt.response_body_sha256
        assert row.assistant_content_sha256 == attempt.assistant_content_sha256
        stored_text = json.dumps(row.receipt)
        assert "provider response must never persist" not in stored_text
        assert "sk-or-v1-not-a-real-key" not in stored_text
        event = connection.execute(
            select(audit_events.c.event_type, audit_events.c.basis).where(
                audit_events.c.object_type == "ai_transfer_execution",
                audit_events.c.event_type == "document.ai_transfer_dispatch_receipted",
            )
        ).one()
        assert event.event_type == "document.ai_transfer_dispatch_receipted"
        audit_text = json.dumps(event.basis)
        assert "provider response must never persist" not in audit_text
        assert "sk-or-v1-not-a-real-key" not in audit_text
        raw_row = connection.execute(select(ai_transfer_raw_responses)).mappings().one()
        blob_row = connection.execute(select(ai_transfer_raw_response_blobs)).mappings().one()
        assert raw_row["metadata_sha256"] == raw_metadata.metadata_sha256
        assert SYNTHETIC_RAW_RESPONSE not in bytes(blob_row["ciphertext"])
        assert SYNTHETIC_RAW_RESPONSE not in bytes(blob_row["wrapped_dek"])
        access_event = connection.execute(
            select(audit_events.c.event_type).where(
                audit_events.c.object_type == "ai_transfer_raw_response",
                audit_events.c.event_type == "document.ai_transfer_raw_response_accessed",
            )
        ).one()
        assert access_event.event_type == "document.ai_transfer_raw_response_accessed"
    with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
        connection.execute(
            ai_transfer_raw_responses.update()
            .where(ai_transfer_raw_responses.c.raw_response_id == raw_metadata.raw_response_id)
            .values(key_id="forged-key")
        )
    with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
        connection.execute(
            ai_transfer_raw_response_blobs.delete().where(
                ai_transfer_raw_response_blobs.c.raw_response_id == raw_metadata.raw_response_id
            )
        )
    with storage.transaction(other) as connection:
        assert (
            connection.execute(
                select(func.count()).select_from(ai_transfer_dispatch_receipts)
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                select(func.count()).select_from(ai_transfer_raw_responses)
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                select(func.count()).select_from(ai_transfer_raw_response_blobs)
            ).scalar_one()
            == 0
        )


def test_ai_transfer_success_materializes_once_into_raw_free_review_run(
    raw_storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    admission, claim = _claimed_ai_transfer(raw_storage, tenant)
    attempt, raw_response = _valid_requirement_capture(claim)
    raw_storage.persist_captured_ai_transfer_result(
        tenant,
        claim.claimed_by,
        admission.execution_id,
        attempt,
        raw_response,
    )

    def materialize():
        return raw_storage.materialize_ai_transfer_requirement_run(
            tenant,
            "materializer-1",
            admission.execution_id,
        )

    with ThreadPoolExecutor(max_workers=2) as executor:
        first, replay = list(executor.map(lambda _: materialize(), range(2)))

    assert first == replay
    materialization, record = first
    assert materialization.status == "MATERIALIZED"
    assert materialization.reason == "MATERIALIZED"
    assert record is not None
    assert isinstance(record.snapshot, AiRequirementTransferRunSnapshot)
    assert record.snapshot.prepared_by == "preparer"
    assert record.snapshot.materialized_by == "materializer-1"
    assert record.snapshot.completeness_status is CompletenessStatus.UNKNOWN
    assert len(record.snapshot.result.candidates) == 1
    serialized = record.snapshot.model_dump_json()
    assert "response_content" not in serialized
    assert raw_response.decode("utf-8") not in serialized
    assert raw_storage.get_ai_requirement_run(tenant, record.snapshot.run_id) == record
    with pytest.raises(AiTransferNotFound):
        raw_storage.materialize_ai_transfer_requirement_run(
            "foreign-tenant",
            "materializer-1",
            admission.execution_id,
        )


def test_ai_transfer_invalid_candidate_is_rejected_without_mutating_transport(
    raw_storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    admission, claim = _claimed_ai_transfer(raw_storage, tenant)
    attempt, raw_response = _valid_requirement_capture(
        claim,
        wire={
            "requirements": [],
            "unresolved_block_ids": [],
            "document_notes": [],
            "items": [
                {
                    "item_name": {
                        "value": "Вымышленный товар",
                        "source_block_ids": ["b000001"],
                    },
                    "okpd2_codes": [],
                    "ktru_codes": [],
                    "unit": None,
                    "quantity": None,
                    "uncertainty": None,
                }
            ],
        },
    )
    raw_storage.persist_captured_ai_transfer_result(
        tenant,
        claim.claimed_by,
        admission.execution_id,
        attempt,
        raw_response,
    )

    materialization, record = raw_storage.materialize_ai_transfer_requirement_run(
        tenant,
        "materializer-1",
        admission.execution_id,
    )
    replay = raw_storage.materialize_ai_transfer_requirement_run(
        tenant,
        "different-recovery-worker",
        admission.execution_id,
    )

    assert record is None
    assert replay == (materialization, None)
    assert materialization.status == "REJECTED"
    assert materialization.reason == "INVALID_CANDIDATE_EVIDENCE"
    execution = raw_storage.get_ai_transfer_execution(tenant, admission.execution_id)
    assert execution is not None
    assert execution[2] is not None
    assert execution[2].outcome == "SUCCESS"


def test_ai_transfer_dispatch_receipt_sql_is_append_only_and_rejects_forgery(
    raw_storage: PostgresStorage,
) -> None:
    storage = raw_storage
    tenant = f"tenant-{uuid.uuid4()}"
    admission, claim = _claimed_ai_transfer(storage, tenant)
    receipt, _ = storage.persist_captured_ai_transfer_result(
        tenant,
        claim.claimed_by,
        admission.execution_id,
        _successful_transport_attempt(claim),
        SYNTHETIC_RAW_RESPONSE,
    )

    with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
        connection.execute(
            ai_transfer_dispatch_receipts.update()
            .where(ai_transfer_dispatch_receipts.c.receipt_id == receipt.receipt_id)
            .values(outcome="FAILED")
        )
    with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
        connection.execute(
            ai_transfer_dispatch_receipts.delete().where(
                ai_transfer_dispatch_receipts.c.receipt_id == receipt.receipt_id
            )
        )

    semantic_admission, semantic_claim = _claimed_ai_transfer(storage, tenant)
    values, forged_json = _forged_receipt_values(receipt, semantic_claim)
    forged_json.update(
        outcome="SUCCESS",
        reason="VALIDATED_RESPONSE",
        capture_complete=False,
        response_body_sha256=None,
    )
    values.update(
        outcome="SUCCESS",
        reason="VALIDATED_RESPONSE",
        capture_complete=False,
        response_body_sha256=None,
        receipt=forged_json,
    )
    with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
        connection.execute(insert(ai_transfer_dispatch_receipts).values(**values))
    assert storage.get_ai_transfer_execution(tenant, semantic_admission.execution_id) == (
        semantic_admission,
        semantic_claim,
        None,
    )

    projection_admission, projection_claim = _claimed_ai_transfer(storage, tenant)
    values, forged_json = _forged_receipt_values(receipt, projection_claim)
    forged_json["outcome"] = "FAILED"
    values["receipt"] = forged_json
    with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
        connection.execute(insert(ai_transfer_dispatch_receipts).values(**values))
    assert storage.get_ai_transfer_execution(tenant, projection_admission.execution_id) == (
        projection_admission,
        projection_claim,
        None,
    )


class _AiReportClient:
    def __init__(self, content: str) -> None:
        self.content = content

    def complete_json(self, **kwargs: object) -> OpenRouterJsonResponse:
        model = kwargs["model"]
        assert isinstance(model, str)
        return OpenRouterJsonResponse(
            request_sha256="a" * 64,
            response_id="integration-response",
            model=model,
            content=self.content,
        )


def _ai_report(
    extraction: ExtractionResult,
    *,
    company_id: str,
    source_sha256: str,
    complete: bool = True,
):
    content = (
        json.dumps(
            {
                "items": [],
                "requirements": [
                    {
                        "item_name": "Насос",
                        "property": "power",
                        "operator": ">=",
                        "value": "5",
                        "unit": "кВт",
                        "mandatory": True,
                        "source_block_ids": ["b000001"],
                        "uncertainty": None,
                    }
                ],
                "unresolved_block_ids": [],
                "document_notes": [],
            }
        )
        if complete
        else "{}"
    )
    return analyze_requirements_with_model(
        extraction,
        company_id=company_id,
        source_sha256=source_sha256,
        authorization_reference="integration import",
        client=_AiReportClient(content),
        created_at=datetime.now(UTC),
    )


def test_ai_requirement_run_review_is_immutable_atomic_and_tenant_scoped(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    other = f"tenant-{uuid.uuid4()}"
    document_id = f"document-{uuid.uuid4()}"
    source = Document()
    source.add_paragraph("Насос; мощность не менее 5 кВт")
    output = BytesIO()
    source.save(output)
    content = output.getvalue()
    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    extraction = extract_document(
        DocumentInput(
            document_id=document_id,
            filename="requirements.docx",
            declared_media_type=media_type,
            content=content,
        )
    )
    document = storage.store_document(
        tenant,
        "operator-1",
        "requirements.docx",
        media_type,
        DocumentRole.TENDER_REQUIREMENTS,
        content,
        extraction,
        "ai-run-source",
        "a" * 64,
        document_id=document_id,
    )
    report_v2 = _ai_report(
        extraction,
        company_id=tenant,
        source_sha256=document.sha256,
    )
    run_id = f"ai-run-{uuid.uuid4()}"
    created = storage.store_ai_requirement_run(
        tenant,
        "operator-1",
        document_id,
        report_v2,
        "ai-run-create",
        "b" * 64,
        run_id=run_id,
    )
    replay = storage.store_ai_requirement_run(
        tenant,
        "operator-1",
        document_id,
        report_v2,
        "ai-run-create",
        "b" * 64,
    )
    assert replay == created
    for null_field in ("company_id", "schema_version"):
        rejected_run_id = f"ai-run-null-{null_field}-{uuid.uuid4()}"
        rejected_evidence = "7" * 64
        rejected_snapshot = created.snapshot.model_dump(mode="json")
        rejected_snapshot["run_id"] = rejected_run_id
        rejected_snapshot["evidence_sha256"] = rejected_evidence
        rejected_snapshot[null_field] = None
        with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
            connection.execute(
                insert(document_ai_requirement_runs).values(
                    company_id=tenant,
                    run_id=rejected_run_id,
                    source_document_id=document_id,
                    source_sha256=document.sha256,
                    evidence_sha256=rejected_evidence,
                    snapshot=rejected_snapshot,
                    created_at=datetime.now(UTC),
                )
            )
    assert storage.get_ai_requirement_run(other, run_id) is None
    assert storage.list_ai_requirement_runs(other) == ()
    with pytest.raises(AiRequirementRunSourceMismatch):
        storage.store_ai_requirement_run(
            tenant,
            "operator-1",
            document_id,
            report_v2.model_copy(update={"company_id": other}),
            "ai-run-source-mismatch",
            "8" * 64,
        )
    with pytest.raises(AiRequirementReviewConflict, match="actor"):
        storage.store_ai_requirement_run(
            tenant,
            "another-operator",
            document_id,
            report_v2,
            "ai-run-create",
            "b" * 64,
        )
    candidate = created.snapshot.report.candidates[0]
    candidate_decisions = (
        AiRequirementCandidateDecision(
            candidate_kind=CandidateKind.REQUIREMENT,
            candidate_id=candidate.candidate_id,
            disposition=CandidateDisposition.ACCEPTED,
            reason="exact evidence checked",
        ),
    )
    with pytest.raises(AiRequirementReviewConflict, match="own AI run"):
        storage.decide_ai_requirement_review(
            tenant,
            "operator-1",
            run_id,
            created.snapshot.evidence_sha256,
            candidate_decisions,
            "ai-run-self",
            "c" * 64,
        )
    with pytest.raises(AiRequirementReviewConflict, match="stale"):
        storage.decide_ai_requirement_review(
            tenant,
            "reviewer-1",
            run_id,
            "0" * 64,
            candidate_decisions,
            "ai-run-stale",
            "d" * 64,
        )
    decided = storage.decide_ai_requirement_review(
        tenant,
        "reviewer-1",
        run_id,
        created.snapshot.evidence_sha256,
        candidate_decisions,
        "ai-run-decision",
        "e" * 64,
    )
    exact_replay = storage.decide_ai_requirement_review(
        tenant,
        "reviewer-1",
        run_id,
        created.snapshot.evidence_sha256,
        candidate_decisions,
        "ai-run-decision",
        "e" * 64,
    )
    assert exact_replay == decided
    assert decided.decision is not None
    with pytest.raises(AiRequirementReviewConflict, match="already"):
        storage.decide_ai_requirement_review(
            tenant,
            "reviewer-2",
            run_id,
            created.snapshot.evidence_sha256,
            candidate_decisions,
            "ai-run-second-decision",
            "f" * 64,
        )
    create_replay_after_decision = storage.store_ai_requirement_run(
        tenant,
        "operator-1",
        document_id,
        report_v2,
        "ai-run-create",
        "b" * 64,
    )
    assert create_replay_after_decision.decision is None

    v1_payload = report_v2.model_dump(mode="json")
    v1_payload["schema_version"] = "ai-requirement-analysis-v1"
    v1_payload["prompt_version"] = "ai-requirements-v1"
    v1_payload.pop("item_candidates")
    for chunk in v1_payload["chunks"]:
        chunk.pop("item_candidates")
    v1_payload["evidence_sha256"] = canonical_json_sha256(
        {key: value for key, value in v1_payload.items() if key != "evidence_sha256"}
    )
    report_v1 = validate_ai_requirement_report(v1_payload)
    assert isinstance(report_v1, AiRequirementAnalysisReport)
    v1_run = storage.store_ai_requirement_run(
        tenant,
        "operator-1",
        document_id,
        report_v1,
        "ai-run-v1",
        "1" * 64,
        run_id=f"ai-run-v1-{uuid.uuid4()}",
    )
    assert type(v1_run.snapshot.report) is AiRequirementAnalysisReport
    with pytest.raises(AiRequirementReviewConflict, match="receipt"):
        storage.decide_ai_requirement_review(
            tenant,
            "reviewer-1",
            v1_run.snapshot.run_id,
            v1_run.snapshot.evidence_sha256,
            candidate_decisions,
            "ai-run-decision",
            "e" * 64,
        )

    incomplete = _ai_report(
        extraction,
        company_id=tenant,
        source_sha256=document.sha256,
        complete=False,
    )
    incomplete_run = storage.store_ai_requirement_run(
        tenant,
        "operator-1",
        document_id,
        incomplete,
        "ai-run-incomplete",
        "2" * 64,
        run_id=f"ai-run-incomplete-{uuid.uuid4()}",
    )
    with pytest.raises(AiRequirementReviewConflict, match="invalid"):
        storage.decide_ai_requirement_review(
            tenant,
            "reviewer-1",
            incomplete_run.snapshot.run_id,
            incomplete_run.snapshot.evidence_sha256,
            (),
            "ai-run-incomplete-decision",
            "3" * 64,
        )

    events = storage.list_audit_events(
        tenant,
        object_type="document_ai_requirement_run",
        object_id=run_id,
    )
    assert [event.event_type for event in events] == [
        "document.ai_requirement_run_imported",
        "document.ai_requirement_review_decided",
    ]
    assert "source_text" not in json.dumps([event.basis for event in events], ensure_ascii=False)
    with storage.transaction(other) as connection:
        assert (
            connection.execute(
                select(func.count()).select_from(document_ai_requirement_runs)
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                select(func.count()).select_from(document_ai_requirement_review_decisions)
            ).scalar_one()
            == 0
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_ai_requirement_runs.update()
            .where(document_ai_requirement_runs.c.run_id == run_id)
            .values(evidence_sha256="9" * 64)
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_ai_requirement_review_decisions.delete().where(
                document_ai_requirement_review_decisions.c.run_id == run_id
            )
        )


def test_document_completeness_review_is_atomic_immutable_and_tenant_scoped(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    other = f"tenant-{uuid.uuid4()}"
    document_id = f"document-{uuid.uuid4()}"
    source = Document()
    source.add_paragraph("Гарантия не менее 24 мес.")
    output = BytesIO()
    source.save(output)
    content = output.getvalue()
    media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    extraction = extract_document(
        DocumentInput(
            document_id=document_id,
            filename="requirements.docx",
            declared_media_type=media_type,
            content=content,
        )
    )
    document = storage.store_document(
        tenant,
        "operator-1",
        "requirements.docx",
        media_type,
        DocumentRole.TENDER_REQUIREMENTS,
        content,
        extraction,
        "completeness-source",
        "1" * 64,
        document_id=document_id,
    )
    review_id = f"review-{uuid.uuid4()}"
    created = storage.store_completeness_review(
        tenant,
        "operator-1",
        document.document_id,
        "completeness-create",
        "2" * 64,
        review_id=review_id,
    )
    replay = storage.store_completeness_review(
        tenant,
        "operator-1",
        document.document_id,
        "completeness-create",
        "2" * 64,
    )
    assert replay == created
    assert created.snapshot.mapping.completeness.status is CompletenessStatus.UNKNOWN
    assert storage.get_completeness_review(other, review_id) is None
    assert storage.list_completeness_reviews(other) == ()
    with pytest.raises(IdempotencyConflict):
        storage.store_completeness_review(
            tenant,
            "operator-1",
            document.document_id,
            "completeness-create",
            "3" * 64,
        )
    with pytest.raises(CompletenessReviewConflict, match="own snapshot"):
        storage.decide_completeness_review(
            tenant,
            "operator-1",
            review_id,
            created.snapshot.evidence_sha256,
            CompletenessStatus.PASS,
            "all requirements",
            "reviewed",
            (),
            "completeness-self-review",
            "4" * 64,
        )
    with pytest.raises(CompletenessReviewConflict, match="stale"):
        storage.decide_completeness_review(
            tenant,
            "reviewer-1",
            review_id,
            "0" * 64,
            CompletenessStatus.PASS,
            "all requirements",
            "reviewed",
            (),
            "completeness-stale",
            "5" * 64,
        )
    changed_extraction = created.snapshot.extraction.model_copy(update={"characters_extracted": 0})
    changed_snapshot = created.snapshot.model_copy(update={"extraction": changed_extraction})
    with (
        storage.transaction(tenant) as connection,
        pytest.raises(CompletenessReviewSourceMismatch, match="evidence changed"),
    ):
        storage._verify_completeness_source_bytes(connection, tenant, changed_snapshot)
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_completeness_review_decisions.insert().values(
                company_id=tenant,
                review_id=review_id,
                evidence_sha256="0" * 64,
                status="UNKNOWN",
                decision={
                    "schema_version": "document-completeness-decision-v1",
                    "company_id": tenant,
                    "review_id": review_id,
                    "evidence_sha256": "0" * 64,
                    "status": "UNKNOWN",
                },
                decided_by="reviewer-1",
                created_at=datetime.now(UTC),
            )
        )
    decided = storage.decide_completeness_review(
        tenant,
        "reviewer-1",
        review_id,
        created.snapshot.evidence_sha256,
        CompletenessStatus.PASS,
        "all requirements",
        "reviewed",
        (),
        "completeness-decision",
        "6" * 64,
    )
    exact_replay = storage.decide_completeness_review(
        tenant,
        "reviewer-1",
        review_id,
        created.snapshot.evidence_sha256,
        CompletenessStatus.PASS,
        "all requirements",
        "reviewed",
        (),
        "completeness-decision",
        "6" * 64,
    )
    assert exact_replay == decided
    assert decided.decision is not None
    assert decided.decision.status is CompletenessStatus.PASS
    create_replay_after_decision = storage.store_completeness_review(
        tenant,
        "operator-1",
        document.document_id,
        "completeness-create",
        "2" * 64,
    )
    assert create_replay_after_decision == created
    assert create_replay_after_decision.decision is None
    events = storage.list_audit_events(
        tenant,
        object_type="document_completeness_review",
        object_id=review_id,
    )
    assert [event.event_type for event in events] == [
        "document.completeness_review_created",
        "document.completeness_review_decided",
    ]
    assert events[-1].basis["scope"] == "all requirements"
    assert len(events[-1].basis["decision_sha256"]) == 64
    with pytest.raises(CompletenessReviewConflict):
        storage.decide_completeness_review(
            tenant,
            "reviewer-1",
            review_id,
            created.snapshot.evidence_sha256,
            CompletenessStatus.UNKNOWN,
            "all requirements",
            "changed",
            (),
            "completeness-second-decision",
            "7" * 64,
        )
    with storage.transaction(other) as connection:
        assert (
            connection.execute(
                select(func.count()).select_from(document_completeness_reviews)
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                select(func.count()).select_from(document_completeness_review_decisions)
            ).scalar_one()
            == 0
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_completeness_reviews.update()
            .where(document_completeness_reviews.c.review_id == review_id)
            .values(evidence_sha256="f" * 64)
        )


def test_document_set_evidence_manifest_is_atomic_ordered_and_tenant_scoped(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    other = f"tenant-{uuid.uuid4()}"
    reviews = []
    for number in (1, 2):
        document_id = f"document-{uuid.uuid4()}"
        source = Document()
        source.add_paragraph(f"Мощность не менее {number + 4} кВт")
        output = BytesIO()
        source.save(output)
        content = output.getvalue()
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        extraction = extract_document(
            DocumentInput(
                document_id=document_id,
                filename=f"requirements-{number}.docx",
                declared_media_type=media_type,
                content=content,
            )
        )
        document = storage.store_document(
            tenant,
            "operator-1",
            f"requirements-{number}.docx",
            media_type,
            DocumentRole.TENDER_REQUIREMENTS,
            content,
            extraction,
            f"document-set-source-{number}",
            str(number) * 64,
            document_id=document_id,
        )
        reviews.append(
            storage.store_completeness_review(
                tenant,
                "operator-1",
                document.document_id,
                f"document-set-review-{number}",
                str(number + 2) * 64,
                review_id=f"review-{uuid.uuid4()}",
            )
        )
    refs = tuple((record.snapshot.review_id, record.snapshot.evidence_sha256) for record in reviews)
    manifest_id = f"manifest-{uuid.uuid4()}"
    created = storage.store_document_set_evidence_manifest(
        tenant,
        "operator-1",
        "Полный выбранный комплект ТЗ",
        refs,
        "document-set-create",
        "7" * 64,
        manifest_id=manifest_id,
    )
    replay = storage.store_document_set_evidence_manifest(
        tenant,
        "operator-1",
        "Полный выбранный комплект ТЗ",
        refs,
        "document-set-create",
        "7" * 64,
    )
    assert replay == created
    assert created.completeness_status is CompletenessStatus.UNKNOWN
    assert [member.completeness_review_id for member in created.members] == [
        review_id for review_id, _ in refs
    ]
    assert storage.get_document_set_evidence_manifest(tenant, manifest_id) == created
    assert storage.get_document_set_evidence_manifest(other, manifest_id) is None
    incomplete_manifest_id = f"manifest-{uuid.uuid4()}"
    incomplete_payload = created.model_dump(mode="json")
    incomplete_payload["manifest_id"] = incomplete_manifest_id
    with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
        connection.execute(
            insert(document_set_evidence_manifests).values(
                company_id=tenant,
                manifest_id=incomplete_manifest_id,
                scope=created.scope,
                member_count=len(created.members),
                completeness_status=created.completeness_status.value,
                reason=created.reason,
                evidence_sha256=created.evidence_sha256,
                manifest=incomplete_payload,
                created_by=created.created_by,
                created_at=created.created_at,
            )
        )
    assert storage.get_document_set_evidence_manifest(tenant, incomplete_manifest_id) is None
    for forged_field, forged_value in (
        ("source_document_id", f"forged-{uuid.uuid4()}"),
        ("mapper_version", "forged-mapper"),
    ):
        forged_payload = created.model_dump(mode="json", exclude={"evidence_sha256"})
        forged_payload["manifest_id"] = f"manifest-{uuid.uuid4()}"
        forged_payload["members"][0][forged_field] = forged_value
        forged_payload["evidence_sha256"] = canonical_json_sha256(forged_payload)
        forged = DocumentSetEvidenceManifest.model_validate_json(json.dumps(forged_payload))
        with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
            connection.execute(
                insert(document_set_evidence_manifests).values(
                    company_id=tenant,
                    manifest_id=forged.manifest_id,
                    scope=forged.scope,
                    member_count=len(forged.members),
                    completeness_status=forged.completeness_status.value,
                    reason=forged.reason,
                    evidence_sha256=forged.evidence_sha256,
                    manifest=forged.model_dump(mode="json"),
                    created_by=forged.created_by,
                    created_at=forged.created_at,
                )
            )
            connection.execute(
                insert(document_set_evidence_members),
                [
                    {
                        "company_id": tenant,
                        "manifest_id": forged.manifest_id,
                        "ordinal": member.ordinal,
                        "source_document_id": member.source_document_id,
                        "source_document_sha256": member.source_document_sha256,
                        "completeness_review_id": member.completeness_review_id,
                        "completeness_review_evidence_sha256": (
                            member.completeness_review_evidence_sha256
                        ),
                        "mapper_version": member.mapper_version,
                        "mapper_min_confidence": member.mapper_min_confidence,
                        "member": member.model_dump(mode="json"),
                    }
                    for member in forged.members
                ],
            )
    with pytest.raises(IdempotencyConflict):
        storage.store_document_set_evidence_manifest(
            tenant,
            "operator-2",
            "Полный выбранный комплект ТЗ",
            refs,
            "document-set-create",
            "8" * 64,
        )
    with pytest.raises(DocumentSetEvidenceSourceMismatch, match="stale"):
        storage.store_document_set_evidence_manifest(
            tenant,
            "operator-1",
            "Другой комплект",
            ((refs[0][0], "0" * 64), refs[1]),
            "document-set-stale",
            "9" * 64,
        )
    duplicate_review = storage.store_completeness_review(
        tenant,
        "operator-1",
        reviews[0].snapshot.source_document_id,
        "document-set-duplicate-review",
        "a" * 64,
        review_id=f"review-{uuid.uuid4()}",
    )
    with pytest.raises(DocumentSetEvidenceSourceMismatch, match="cannot form"):
        storage.store_document_set_evidence_manifest(
            tenant,
            "operator-1",
            "Duplicate source",
            (
                refs[0],
                (
                    duplicate_review.snapshot.review_id,
                    duplicate_review.snapshot.evidence_sha256,
                ),
            ),
            "document-set-duplicate-source",
            "b" * 64,
        )
    with pytest.raises(DocumentSetEvidenceConflict):
        storage.store_document_set_evidence_manifest(
            tenant,
            "operator-1",
            "Different immutable identity",
            refs,
            "document-set-identity-conflict",
            "c" * 64,
            manifest_id=manifest_id,
        )
    events = storage.list_audit_events(
        tenant,
        object_type="document_set_evidence_manifest",
        object_id=manifest_id,
    )
    assert [event.event_type for event in events] == ["document_set.evidence_manifest_created"]
    assert events[0].basis["member_count"] == 2
    with storage.transaction(other) as connection:
        assert (
            connection.execute(
                select(func.count()).select_from(document_set_evidence_manifests)
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                select(func.count()).select_from(document_set_evidence_members)
            ).scalar_one()
            == 0
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_set_evidence_manifests.update()
            .where(document_set_evidence_manifests.c.manifest_id == manifest_id)
            .values(scope="tampered")
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_set_evidence_members.delete().where(
                document_set_evidence_members.c.manifest_id == manifest_id
            )
        )


def test_document_set_completeness_decision_is_exact_atomic_and_tenant_scoped(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    other = f"tenant-{uuid.uuid4()}"
    records = []
    for number in (1, 2):
        document_id = f"document-{uuid.uuid4()}"
        source = Document()
        source.add_paragraph(f"Гарантия не менее {number + 23} месяцев")
        output = BytesIO()
        source.save(output)
        content = output.getvalue()
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        extraction = extract_document(
            DocumentInput(
                document_id=document_id,
                filename=f"package-{number}.docx",
                declared_media_type=media_type,
                content=content,
            )
        )
        document = storage.store_document(
            tenant,
            "snapshot-author",
            f"package-{number}.docx",
            media_type,
            DocumentRole.TENDER_REQUIREMENTS,
            content,
            extraction,
            f"set-review-source-{number}",
            str(number) * 64,
            document_id=document_id,
        )
        draft = storage.store_completeness_review(
            tenant,
            "snapshot-author",
            document.document_id,
            f"set-review-snapshot-{number}",
            str(number + 2) * 64,
            review_id=f"review-{uuid.uuid4()}",
        )
        records.append(
            storage.decide_completeness_review(
                tenant,
                "single-reviewer",
                draft.snapshot.review_id,
                draft.snapshot.evidence_sha256,
                CompletenessStatus.PASS,
                "entire document",
                "reviewed in full",
                (),
                f"set-review-decision-{number}",
                str(number + 4) * 64,
            )
        )

    refs = tuple((record.snapshot.review_id, record.snapshot.evidence_sha256) for record in records)
    manifest_id = f"manifest-{uuid.uuid4()}"
    manifest = storage.store_document_set_evidence_manifest(
        tenant,
        "manifest-author",
        "exact declared package",
        refs,
        "set-terminal-manifest",
        "7" * 64,
        manifest_id=manifest_id,
    )
    evidence = storage.get_document_set_review_evidence(tenant, manifest_id)
    assert evidence is not None
    assert storage.get_document_set_review_evidence(other, manifest_id) is None
    attestations = tuple(
        DocumentSetMemberAttestation(
            reference=ConstituentDecisionRef(
                member_ordinal=item.member.ordinal,
                review_id=item.member.completeness_review_id,
                review_evidence_sha256=item.member.completeness_review_evidence_sha256,
                decision_sha256=item.decision_sha256 or "",
            ),
            assessment=DocumentSetMemberAssessment.REVIEWED_IN_SCOPE,
            rationale="reviewed together in the declared package scope",
        )
        for item in evidence.constituents
    )
    decision = storage.store_document_set_review_decision(
        tenant,
        "package-reviewer",
        manifest_id,
        manifest.evidence_sha256,
        attestations,
        DocumentSetInventoryAssessment.CONFIRMED_FOR_DECLARED_SCOPE,
        "all declared package documents are present",
        DocumentSetCrossAssessment.CHECKED_NO_OPEN_ISSUES,
        "all exact constituent decisions were compared",
        (),
        CompletenessStatus.PASS,
        "complete for the exact declared scope",
        "set-terminal-decision",
        "8" * 64,
    )
    replay = storage.store_document_set_review_decision(
        tenant,
        "package-reviewer",
        manifest_id,
        manifest.evidence_sha256,
        attestations,
        DocumentSetInventoryAssessment.CONFIRMED_FOR_DECLARED_SCOPE,
        "all declared package documents are present",
        DocumentSetCrossAssessment.CHECKED_NO_OPEN_ISSUES,
        "all exact constituent decisions were compared",
        (),
        CompletenessStatus.PASS,
        "complete for the exact declared scope",
        "set-terminal-decision",
        "8" * 64,
    )
    assert replay == decision
    assert storage.get_document_set_review_decision(tenant, manifest_id) == decision
    assert storage.get_document_set_review_decision(other, manifest_id) is None
    with pytest.raises(DocumentSetReviewConflict, match="terminal decision"):
        storage.store_document_set_review_decision(
            tenant,
            "package-reviewer",
            manifest_id,
            manifest.evidence_sha256,
            attestations,
            DocumentSetInventoryAssessment.CONFIRMED_FOR_DECLARED_SCOPE,
            "all declared package documents are present",
            DocumentSetCrossAssessment.CHECKED_NO_OPEN_ISSUES,
            "all exact constituent decisions were compared",
            (),
            CompletenessStatus.PASS,
            "second decision",
            "set-terminal-second",
            "9" * 64,
        )
    events = storage.list_audit_events(
        tenant,
        object_type="document_set_completeness_decision",
        object_id=manifest_id,
    )
    assert [event.event_type for event in events] == ["document_set.completeness_decided"]
    assert events[0].basis["decision_sha256"] == decision.decision_sha256

    incomplete_manifest = storage.store_document_set_evidence_manifest(
        tenant,
        "manifest-author",
        "exact declared package",
        refs,
        "set-incomplete-manifest",
        "a" * 64,
        manifest_id=f"manifest-{uuid.uuid4()}",
    )
    now = datetime.now(UTC)
    with pytest.raises(DBAPIError), storage.transaction(tenant) as connection:
        connection.execute(
            insert(document_set_completeness_decisions).values(
                company_id=tenant,
                manifest_id=incomplete_manifest.manifest_id,
                manifest_evidence_sha256=incomplete_manifest.evidence_sha256,
                status="PASS",
                decision_sha256="b" * 64,
                decision={
                    "schema_version": "document-set-completeness-decision-v1",
                    "company_id": tenant,
                    "manifest_id": incomplete_manifest.manifest_id,
                    "manifest_evidence_sha256": incomplete_manifest.evidence_sha256,
                    "scope": incomplete_manifest.scope,
                    "members": [attestations[0].model_dump(mode="json")],
                    "status": "PASS",
                    "decision_sha256": "b" * 64,
                    "decided_by": "direct-sql-reviewer",
                    "decided_at": now.isoformat(),
                },
                decided_by="direct-sql-reviewer",
                created_at=now,
            )
        )
    with storage.transaction(other) as connection:
        assert (
            connection.execute(
                select(func.count()).select_from(document_set_completeness_decisions)
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                select(func.count()).select_from(document_set_completeness_member_attestations)
            ).scalar_one()
            == 0
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_set_completeness_decisions.update()
            .where(document_set_completeness_decisions.c.manifest_id == manifest_id)
            .values(status="UNKNOWN")
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_set_completeness_member_attestations.delete().where(
                document_set_completeness_member_attestations.c.manifest_id == manifest_id
            )
        )


@pytest.fixture(scope="session")
def database_url() -> str:
    value = os.environ.get("DATABASE_URL")
    if not value:
        pytest.skip("DATABASE_URL is not set; PostgreSQL integration tests are skipped")
    return value


@pytest.fixture(scope="session", autouse=True)
def migrated_database(database_url: str) -> None:
    admin_url = os.environ.get("DATABASE_ADMIN_URL")
    if admin_url:
        config = Config("alembic.ini")
        config.set_main_option("sqlalchemy.url", admin_url.replace("%", "%%"))
        command.upgrade(config, "head")


@pytest.fixture(scope="session")
def storage(database_url: str) -> PostgresStorage:
    engine = create_engine(database_url, pool_pre_ping=True)
    yield PostgresStorage(engine)
    engine.dispose()


@pytest.fixture(scope="session")
def raw_storage(database_url: str) -> PostgresStorage:
    engine = create_engine(database_url, pool_pre_ping=True)
    yield PostgresStorage(engine, raw_key_provider=_EphemeralTenantKekProvider())
    engine.dispose()


def _store(
    storage: PostgresStorage,
    tenant: str,
    *,
    key: str,
    digest: str,
    document_id: str | None = None,
):
    return storage.store_document(
        tenant,
        "document-parser",
        "technical-requirements.pdf",
        "application/pdf",
        DocumentRole.TENDER_REQUIREMENTS,
        b"%PDF-1.7 synthetic source",
        {
            "status": "UNKNOWN",
            "reasons": ["OCR_REQUIRED"],
            "source_units_total": 1,
            "source_units_extracted": 0,
            "characters_extracted": 0,
        },
        key,
        digest,
        document_id=document_id,
    )


def _bundle_content() -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("unknown.ods", b"unsupported synthetic spreadsheet")
    return output.getvalue()


def _store_bundle(
    storage: PostgresStorage,
    tenant: str,
    *,
    key: str,
    digest: str,
    bundle_id: str,
):
    content = _bundle_content()
    result = extract_document_bundle(
        DocumentInput(
            document_id=bundle_id,
            filename="tender.zip",
            declared_media_type="application/zip",
            content=content,
        )
    )
    return storage.store_document_bundle(
        tenant,
        "bundle-parser",
        "tender.zip",
        "application/zip",
        DocumentRole.TENDER_REQUIREMENTS,
        content,
        result,
        key,
        digest,
        bundle_id=bundle_id,
    )


def _blank_pdf() -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.write(output)
    return output.getvalue()


def _model_brand_catalog(
    tenant: str,
    actor: str,
    catalog_id: str,
    *,
    artifact_id: str,
    source_sha256: str,
    retrieved_at: datetime,
) -> ModelBrandCatalog:
    now = datetime.now(UTC)
    return ModelBrandCatalog(
        catalog_id=catalog_id,
        company_id=tenant,
        created_at=now,
        approved_by=actor,
        approval_reference=f"review/{catalog_id}",
        entries=(
            ModelBrandCatalogEntry(
                entry_id="general-climate-gw-g140-n1a1",
                model="GW-G140/N1A1",
                brand="General Climate",
                valid_from=datetime(2026, 1, 1, tzinfo=UTC),
                evidence=ModelBrandSourceEvidence(
                    source_artifact_id=artifact_id,
                    source_title="Official installation manual",
                    source_url="https://manufacturer.example/manual.pdf",
                    source_document_sha256=source_sha256,
                    locator="model table, row 140",
                    retrieved_at=retrieved_at,
                ),
            ),
        ),
    )


def _store_docx_with_image(storage: PostgresStorage, tenant: str):
    image_output = BytesIO()
    Image.new("RGB", (80, 40), color="white").save(image_output, format="PNG")
    output = BytesIO()
    document = Document()
    document.add_picture(BytesIO(image_output.getvalue()), width=Inches(1))
    document.save(output)
    content = output.getvalue()
    document_id = f"document-{uuid.uuid4()}"
    source = DocumentInput(
        document_id=document_id,
        filename="requirements.docx",
        declared_media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        content=content,
    )
    stored = storage.store_document(
        tenant,
        "document-parser",
        source.filename,
        source.declared_media_type,
        DocumentRole.TENDER_REQUIREMENTS,
        content,
        extract_document(source),
        f"source-{document_id}",
        hashlib.sha256(document_id.encode()).hexdigest(),
        document_id=document_id,
    )
    media = extract_docx_media(source).media[0]
    derivative = build_ocr_derivative(
        OcrEngineInput(
            source_document_id=docx_media_source_id(document_id, media.media_index),
            source_sha256=media.sha256,
            page_count=1,
        ),
        ExtractionResult(
            status="UNKNOWN",
            reasons=("EMBEDDED_MEDIA_IGNORED",),
            source_units_total=1,
            source_units_extracted=0,
            characters_extracted=0,
        ),
        (),
        engine="test-image-ocr/1",
        threshold=0.75,
    )
    return stored, media, derivative


def _store_pdf_bundle(
    storage: PostgresStorage,
    tenant: str,
    *,
    bundle_id: str,
):
    pdf = _blank_pdf()
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("scans/spec.pdf", pdf)
    content = output.getvalue()
    result = extract_document_bundle(
        DocumentInput(
            document_id=bundle_id,
            filename="scans.zip",
            declared_media_type="application/zip",
            content=content,
        )
    )
    stored = storage.store_document_bundle(
        tenant,
        "bundle-parser",
        "scans.zip",
        "application/zip",
        DocumentRole.TENDER_REQUIREMENTS,
        content,
        result,
        f"source-{bundle_id}",
        hashlib.sha256(bundle_id.encode()).hexdigest(),
        bundle_id=bundle_id,
    )
    member = stored.result.members[0]
    derivative = build_ocr_derivative(
        OcrEngineInput(
            source_document_id=bundle_member_source_id(bundle_id, 0),
            source_sha256=member.sha256,
            page_count=member.extraction.source_units_total,
        ),
        member.extraction,
        (),
        engine="test-ocr/1",
        threshold=0.75,
    )
    return stored, member, derivative


def test_document_roundtrip_is_tenant_scoped_and_audited(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    other = f"tenant-{uuid.uuid4()}"

    stored = _store(storage, tenant, key="document-once", digest="1" * 64)

    assert storage.get_document(tenant, stored.document_id) == stored
    assert storage.get_document(other, stored.document_id) is None
    assert stored.byte_size == len(stored.content)
    assert stored.sha256 == hashlib.sha256(stored.content).hexdigest()
    assert stored.document_role is DocumentRole.TENDER_REQUIREMENTS
    events = storage.list_audit_events(
        tenant,
        object_type="source_document",
        object_id=stored.document_id,
    )
    assert len(events) == 1
    assert events[0].event_type == "document.stored"
    assert events[0].actor_id == "document-parser"
    assert events[0].basis["sha256"] == stored.sha256
    assert events[0].basis["document_role"] == "TENDER_REQUIREMENTS"
    assert (
        storage.list_audit_events(
            other,
            object_type="source_document",
            object_id=stored.document_id,
        )
        == []
    )


def test_document_idempotency_replays_and_rejects_digest_or_identity_reuse(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    document_id = f"document-{uuid.uuid4()}"
    first = _store(
        storage,
        tenant,
        key="same-upload",
        digest="2" * 64,
        document_id=document_id,
    )

    replay = _store(storage, tenant, key="same-upload", digest="2" * 64)
    assert replay == first
    with pytest.raises(IdempotencyConflict):
        _store(storage, tenant, key="same-upload", digest="3" * 64)
    with pytest.raises(DocumentConflict):
        _store(
            storage,
            tenant,
            key="another-upload",
            digest="4" * 64,
            document_id=document_id,
        )


def test_same_document_request_is_serialized_under_concurrency(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"

    def save() -> object:
        return _store(storage, tenant, key="concurrent-document", digest="5" * 64)

    with ThreadPoolExecutor(max_workers=2) as executor:
        first, second = executor.map(lambda _: save(), range(2))

    assert first == second
    with storage.transaction(tenant) as connection:
        count = connection.execute(select(func.count()).select_from(documents)).scalar_one()
    assert count == 1


def test_document_rows_are_append_only_and_hidden_by_rls(storage: PostgresStorage) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    other = f"tenant-{uuid.uuid4()}"
    stored = _store(storage, tenant, key="immutable-document", digest="6" * 64)

    with storage.transaction(other) as connection:
        assert connection.execute(select(func.count()).select_from(documents)).scalar_one() == 0
        assert (
            connection.execute(
                select(documents.c.document_id).where(documents.c.document_id == stored.document_id)
            ).scalar_one_or_none()
            is None
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            documents.update()
            .where(documents.c.document_id == stored.document_id)
            .values(original_filename="tampered.pdf")
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(documents.delete().where(documents.c.document_id == stored.document_id))


def test_document_size_and_json_envelope_are_validated(storage: PostgresStorage) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    common = {
        "company_id": tenant,
        "actor_id": "parser",
        "original_filename": "source.pdf",
        "media_type": "application/pdf",
        "document_role": DocumentRole.UNKNOWN,
        "extraction": {"status": "UNKNOWN", "segments": []},
        "idempotency_key": "invalid-size",
        "request_sha256": "7" * 64,
    }
    with pytest.raises(ValueError, match="must not be empty"):
        storage.store_document(content=b"", **common)
    with pytest.raises(ValueError, match="exceeds"):
        storage.store_document(content=b"x" * (MAX_DOCUMENT_BYTES + 1), **common)
    with pytest.raises(ValueError, match="typed extraction schema"):
        storage.store_document(
            content=b"valid",
            **(common | {"extraction": {"invalid": {1, 2}}}),
        )


def test_database_rejects_mismatched_declared_size(storage: PostgresStorage) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            documents.insert().values(
                company_id=tenant,
                document_id=f"invalid-{uuid.uuid4()}",
                original_filename="source.pdf",
                media_type="application/pdf",
                document_role=DocumentRole.UNKNOWN.value,
                byte_size=99,
                sha256="8" * 64,
                content=b"short",
                extraction={"status": "UNKNOWN"},
            )
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            documents.insert().values(
                company_id=tenant,
                document_id=f"invalid-role-{uuid.uuid4()}",
                original_filename="source.pdf",
                media_type="application/pdf",
                document_role="UNTRUSTED_GUESS",
                byte_size=5,
                sha256=hashlib.sha256(b"short").hexdigest(),
                content=b"short",
                extraction={"status": "UNKNOWN"},
            )
        )


def test_document_bundle_roundtrip_is_tenant_scoped_audited_and_idempotent(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    other = f"tenant-{uuid.uuid4()}"
    first_id = f"bundle-{uuid.uuid4()}"
    stored = _store_bundle(
        storage,
        tenant,
        key="bundle-once",
        digest="a" * 64,
        bundle_id=first_id,
    )

    assert storage.get_document_bundle(tenant, first_id) == stored
    assert storage.get_document_bundle(other, first_id) is None
    assert stored.result.source_document_id == stored.bundle_id
    assert stored.result.source_sha256 == stored.sha256
    assert stored.result.members_total == 1
    replay = _store_bundle(
        storage,
        tenant,
        key="bundle-once",
        digest="a" * 64,
        bundle_id=f"bundle-{uuid.uuid4()}",
    )
    assert replay == stored
    with pytest.raises(IdempotencyConflict):
        _store_bundle(
            storage,
            tenant,
            key="bundle-once",
            digest="b" * 64,
            bundle_id=f"bundle-{uuid.uuid4()}",
        )

    events = storage.list_audit_events(
        tenant,
        object_type="document_bundle",
        object_id=first_id,
    )
    assert len(events) == 1
    assert events[0].event_type == "document_bundle.stored"
    assert events[0].actor_id == "bundle-parser"
    assert events[0].basis["sha256"] == stored.sha256
    assert events[0].basis["status"] == "UNKNOWN"
    assert events[0].basis["unsupported_members"] == 1


def test_document_bundle_identity_is_immutable_and_hidden_by_rls(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    other = f"tenant-{uuid.uuid4()}"
    bundle_id = f"bundle-{uuid.uuid4()}"
    _store_bundle(
        storage,
        tenant,
        key="immutable-bundle",
        digest="c" * 64,
        bundle_id=bundle_id,
    )

    with storage.transaction(other) as connection:
        count = connection.execute(select(func.count()).select_from(document_bundles)).scalar_one()
        assert count == 0
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_bundles.update()
            .where(document_bundles.c.bundle_id == bundle_id)
            .values(original_filename="tampered.zip")
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_bundles.delete().where(document_bundles.c.bundle_id == bundle_id)
        )


def test_document_bundle_rejects_identity_reuse_and_mismatched_result(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    bundle_id = f"bundle-{uuid.uuid4()}"
    _store_bundle(
        storage,
        tenant,
        key="bundle-original",
        digest="d" * 64,
        bundle_id=bundle_id,
    )
    with pytest.raises(DocumentBundleConflict):
        _store_bundle(
            storage,
            tenant,
            key="bundle-duplicate",
            digest="e" * 64,
            bundle_id=bundle_id,
        )

    content = _bundle_content()
    wrong_result = extract_document_bundle(
        DocumentInput(
            document_id="different-bundle",
            filename="tender.zip",
            declared_media_type="application/zip",
            content=content,
        )
    )
    with pytest.raises(ValueError, match="source id"):
        storage.store_document_bundle(
            tenant,
            "bundle-parser",
            "tender.zip",
            "application/zip",
            DocumentRole.UNKNOWN,
            content,
            wrong_result,
            "bundle-mismatch",
            "f" * 64,
            bundle_id=f"bundle-{uuid.uuid4()}",
        )

    direct_id = f"bundle-{uuid.uuid4()}"
    direct_sha256 = hashlib.sha256(content).hexdigest()
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_bundles.insert().values(
                company_id=tenant,
                bundle_id=direct_id,
                original_filename="direct.zip",
                media_type="application/zip",
                document_role=DocumentRole.UNKNOWN.value,
                byte_size=len(content),
                sha256=direct_sha256,
                content=content,
                result={
                    "source_document_id": direct_id,
                    "source_sha256": direct_sha256,
                    "status": "UNKNOWN",
                },
            )
        )


def test_bundle_member_ocr_roundtrip_is_tenant_scoped_audited_and_idempotent(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    other = f"tenant-{uuid.uuid4()}"
    bundle_id = f"bundle-{uuid.uuid4()}"
    _, member, derivative = _store_pdf_bundle(storage, tenant, bundle_id=bundle_id)
    ocr_id = f"ocr-{uuid.uuid4()}"

    stored = storage.store_bundle_member_ocr_derivative(
        tenant,
        "ocr-worker",
        bundle_id,
        0,
        member.archive_path,
        member.sha256,
        engine="test-ocr/1",
        threshold=0.75,
        status=derivative.status.value,
        result=derivative,
        mean_confidence=derivative.mean_confidence,
        idempotency_key="bundle-member-ocr",
        request_sha256="1" * 64,
        ocr_id=ocr_id,
    )

    assert storage.get_bundle_member_ocr_derivative(tenant, ocr_id) == stored
    assert storage.get_bundle_member_ocr_derivative(other, ocr_id) is None
    assert storage.get_latest_ocr_for_bundle_member(tenant, bundle_id, 0) == stored
    assert stored.result.source_sha256 == member.sha256
    replay = storage.store_bundle_member_ocr_derivative(
        tenant,
        "ocr-worker",
        bundle_id,
        0,
        member.archive_path,
        member.sha256,
        engine="test-ocr/1",
        threshold=0.75,
        status=derivative.status.value,
        result=derivative,
        mean_confidence=derivative.mean_confidence,
        idempotency_key="bundle-member-ocr",
        request_sha256="1" * 64,
        ocr_id=f"ocr-{uuid.uuid4()}",
    )
    assert replay == stored
    with pytest.raises(IdempotencyConflict):
        storage.store_bundle_member_ocr_derivative(
            tenant,
            "ocr-worker",
            bundle_id,
            0,
            member.archive_path,
            member.sha256,
            engine="test-ocr/1",
            threshold=0.75,
            status=derivative.status.value,
            result=derivative,
            mean_confidence=derivative.mean_confidence,
            idempotency_key="bundle-member-ocr",
            request_sha256="2" * 64,
        )

    events = storage.list_audit_events(
        tenant,
        object_type="document_bundle",
        object_id=bundle_id,
    )
    ocr_events = [event for event in events if event.event_type.endswith("member_ocr_stored")]
    assert len(ocr_events) == 1
    assert ocr_events[0].actor_id == "ocr-worker"
    assert ocr_events[0].basis["source_member_sha256"] == member.sha256


def test_bundle_member_ocr_is_append_only_and_hidden_by_rls(storage: PostgresStorage) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    other = f"tenant-{uuid.uuid4()}"
    bundle_id = f"bundle-{uuid.uuid4()}"
    _, member, derivative = _store_pdf_bundle(storage, tenant, bundle_id=bundle_id)
    stored = storage.store_bundle_member_ocr_derivative(
        tenant,
        "ocr-worker",
        bundle_id,
        0,
        member.archive_path,
        member.sha256,
        engine="test-ocr/1",
        threshold=0.75,
        status=derivative.status.value,
        result=derivative,
        mean_confidence=derivative.mean_confidence,
        idempotency_key="immutable-member-ocr",
        request_sha256="3" * 64,
    )

    with storage.transaction(other) as connection:
        count = connection.execute(
            select(func.count()).select_from(document_bundle_member_ocr_derivatives)
        ).scalar_one()
        assert count == 0
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_bundle_member_ocr_derivatives.update()
            .where(document_bundle_member_ocr_derivatives.c.ocr_id == stored.ocr_id)
            .values(status="PASS")
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_bundle_member_ocr_derivatives.delete().where(
                document_bundle_member_ocr_derivatives.c.ocr_id == stored.ocr_id
            )
        )


def test_bundle_member_ocr_rejects_wrong_member_evidence_and_identity_reuse(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    bundle_id = f"bundle-{uuid.uuid4()}"
    _, member, derivative = _store_pdf_bundle(storage, tenant, bundle_id=bundle_id)
    ocr_id = f"ocr-{uuid.uuid4()}"
    storage.store_bundle_member_ocr_derivative(
        tenant,
        "ocr-worker",
        bundle_id,
        0,
        member.archive_path,
        member.sha256,
        engine="test-ocr/1",
        threshold=0.75,
        status=derivative.status.value,
        result=derivative,
        mean_confidence=derivative.mean_confidence,
        idempotency_key="original-member-ocr",
        request_sha256="4" * 64,
        ocr_id=ocr_id,
    )

    with pytest.raises(BundleMemberOcrConflict):
        storage.store_bundle_member_ocr_derivative(
            tenant,
            "ocr-worker",
            bundle_id,
            0,
            member.archive_path,
            member.sha256,
            engine="test-ocr/1",
            threshold=0.75,
            status=derivative.status.value,
            result=derivative,
            mean_confidence=derivative.mean_confidence,
            idempotency_key="duplicate-member-ocr",
            request_sha256="5" * 64,
            ocr_id=ocr_id,
        )
    with pytest.raises(StorageError, match="verified OCR source"):
        storage.store_bundle_member_ocr_derivative(
            tenant,
            "ocr-worker",
            bundle_id,
            0,
            "different.pdf",
            member.sha256,
            engine="test-ocr/1",
            threshold=0.75,
            status=derivative.status.value,
            result=derivative,
            mean_confidence=derivative.mean_confidence,
            idempotency_key="wrong-member-ocr",
            request_sha256="6" * 64,
        )

    contradictory = derivative.model_dump(mode="json") | {"status": "PASS"}
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_bundle_member_ocr_derivatives.insert().values(
                company_id=tenant,
                ocr_id=f"ocr-{uuid.uuid4()}",
                source_bundle_id=bundle_id,
                member_index=0,
                archive_path=member.archive_path,
                source_member_sha256=member.sha256,
                engine="test-ocr/1",
                threshold=0.75,
                status="UNKNOWN",
                result=contradictory,
                mean_confidence=derivative.mean_confidence,
            )
        )


def test_docx_media_ocr_is_tenant_scoped_audited_idempotent_and_append_only(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    other = f"tenant-{uuid.uuid4()}"
    document, media, derivative = _store_docx_with_image(storage, tenant)
    ocr_id = f"ocr-{uuid.uuid4()}"

    stored = storage.store_docx_media_ocr_derivative(
        tenant,
        "ocr-worker",
        document.document_id,
        media.media_index,
        media.archive_path,
        document.sha256,
        media.sha256,
        media.media_type,
        engine="test-image-ocr/1",
        threshold=0.75,
        status=derivative.status.value,
        result=derivative,
        mean_confidence=derivative.mean_confidence,
        idempotency_key="docx-media-ocr",
        request_sha256="7" * 64,
        ocr_id=ocr_id,
    )

    assert storage.get_docx_media_ocr_derivative(tenant, ocr_id) == stored
    assert storage.get_docx_media_ocr_derivative(other, ocr_id) is None
    assert (
        storage.get_latest_ocr_for_docx_media(
            tenant,
            document.document_id,
            media.media_index,
        )
        == stored
    )
    replay = storage.store_docx_media_ocr_derivative(
        tenant,
        "ocr-worker",
        document.document_id,
        media.media_index,
        media.archive_path,
        document.sha256,
        media.sha256,
        media.media_type,
        engine="test-image-ocr/1",
        threshold=0.75,
        status=derivative.status.value,
        result=derivative,
        mean_confidence=derivative.mean_confidence,
        idempotency_key="docx-media-ocr",
        request_sha256="7" * 64,
    )
    assert replay == stored
    with pytest.raises(IdempotencyConflict):
        storage.store_docx_media_ocr_derivative(
            tenant,
            "ocr-worker",
            document.document_id,
            media.media_index,
            media.archive_path,
            document.sha256,
            media.sha256,
            media.media_type,
            engine="test-image-ocr/1",
            threshold=0.75,
            status=derivative.status.value,
            result=derivative,
            mean_confidence=derivative.mean_confidence,
            idempotency_key="docx-media-ocr",
            request_sha256="8" * 64,
        )

    events = storage.list_audit_events(
        tenant,
        object_type="document",
        object_id=document.document_id,
    )
    assert [event.event_type for event in events] == ["document.docx_media_ocr_stored"]
    assert events[0].basis["source_media_sha256"] == media.sha256
    with storage.transaction(other) as connection:
        count = connection.execute(
            select(func.count()).select_from(document_docx_media_ocr_derivatives)
        ).scalar_one()
        assert count == 0
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_docx_media_ocr_derivatives.update()
            .where(document_docx_media_ocr_derivatives.c.ocr_id == ocr_id)
            .values(status="PASS")
        )


def test_document_search_index_is_verified_versioned_audited_and_tenant_scoped(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    other = f"tenant-{uuid.uuid4()}"
    document = _store(
        storage,
        tenant,
        key="search-source",
        digest="a" * 64,
    )
    source = DocumentInput(
        document_id=document.document_id,
        filename=document.original_filename,
        declared_media_type=document.media_type,
        content=document.content,
    )
    first_entry = build_document_search_entry(
        source,
        document.extraction,
        company_id=tenant,
        title="Первый поисковый заголовок",
    )
    first_id = f"search-{uuid.uuid4()}"

    first = storage.store_document_search_index(
        tenant,
        "search-worker",
        document.document_id,
        first_entry,
        "search-index-first",
        "b" * 64,
        search_index_id=first_id,
    )

    assert first.entry == first_entry
    assert storage.list_latest_document_search_indexes(tenant) == (first,)
    assert storage.list_latest_document_search_indexes(other) == ()
    replay = storage.store_document_search_index(
        tenant,
        "search-worker",
        document.document_id,
        first_entry,
        "search-index-first",
        "b" * 64,
    )
    assert replay == first
    with pytest.raises(IdempotencyConflict):
        storage.store_document_search_index(
            tenant,
            "search-worker",
            document.document_id,
            first_entry,
            "search-index-first",
            "c" * 64,
        )

    second_entry = first_entry.model_copy(update={"title": "Второй поисковый заголовок"})
    second = storage.store_document_search_index(
        tenant,
        "search-worker",
        document.document_id,
        second_entry,
        "search-index-second",
        "d" * 64,
        search_index_id=f"search-{uuid.uuid4()}",
    )
    assert storage.list_latest_document_search_indexes(tenant) == (second,)

    tampered = first_entry.model_copy(update={"source_sha256": "0" * 64})
    with pytest.raises(ValueError, match="immutable source"):
        storage.store_document_search_index(
            tenant,
            "search-worker",
            document.document_id,
            tampered,
            "search-index-tampered",
            "e" * 64,
        )

    events = storage.list_audit_events(
        tenant,
        object_type="document_search_index",
    )
    assert [event.event_type for event in events] == [
        "document.search_index_stored",
        "document.search_index_stored",
    ]
    assert events[0].basis["source_document_sha256"] == document.sha256
    with storage.transaction(other) as connection:
        count = connection.execute(
            select(func.count()).select_from(document_search_indexes)
        ).scalar_one()
        assert count == 0
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            document_search_indexes.update()
            .where(document_search_indexes.c.search_index_id == first_id)
            .values(source_document_sha256="f" * 64)
        )


def test_document_search_index_identity_is_immutable(storage: PostgresStorage) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    document = _store(storage, tenant, key="search-conflict-source", digest="1" * 64)
    entry = build_document_search_entry(
        DocumentInput(
            document_id=document.document_id,
            filename=document.original_filename,
            declared_media_type=document.media_type,
            content=document.content,
        ),
        document.extraction,
        company_id=tenant,
    )
    search_index_id = f"search-{uuid.uuid4()}"
    storage.store_document_search_index(
        tenant,
        "search-worker",
        document.document_id,
        entry,
        "search-conflict-first",
        "2" * 64,
        search_index_id=search_index_id,
    )

    with pytest.raises(DocumentSearchIndexConflict):
        storage.store_document_search_index(
            tenant,
            "search-worker",
            document.document_id,
            entry,
            "search-conflict-second",
            "3" * 64,
            search_index_id=search_index_id,
        )


def test_model_brand_catalog_is_reviewed_audited_immutable_and_tenant_scoped(
    storage: PostgresStorage,
) -> None:
    tenant = f"tenant-{uuid.uuid4()}"
    other = f"tenant-{uuid.uuid4()}"
    actor = "catalog-reviewer"
    artifact_id = f"artifact-{uuid.uuid4()}"
    source_pdf = _blank_pdf()
    retrieved_at = datetime.now(UTC)
    artifact = storage.store_model_brand_source_artifact(
        tenant,
        actor,
        artifact_id,
        "Official installation manual",
        "https://manufacturer.example/manual.pdf",
        "official-manual.pdf",
        source_pdf,
        retrieved_at,
        "artifact-first",
        "3" * 64,
    )
    assert artifact.sha256 == hashlib.sha256(source_pdf).hexdigest()
    assert storage.get_model_brand_source_artifact(tenant, artifact_id) == artifact
    assert storage.get_model_brand_source_artifact(other, artifact_id) is None
    assert (
        storage.store_model_brand_source_artifact(
            tenant,
            actor,
            "ignored-on-replay",
            "Official installation manual",
            "https://manufacturer.example/manual.pdf",
            "official-manual.pdf",
            source_pdf,
            retrieved_at,
            "artifact-first",
            "3" * 64,
        )
        == artifact
    )
    with pytest.raises(ModelBrandSourceArtifactConflict):
        storage.store_model_brand_source_artifact(
            tenant,
            actor,
            artifact_id,
            "Official installation manual",
            "https://manufacturer.example/manual.pdf",
            "official-manual.pdf",
            source_pdf,
            retrieved_at,
            "artifact-second",
            "7" * 64,
        )

    catalog = _model_brand_catalog(
        tenant,
        actor,
        f"catalog-{uuid.uuid4()}",
        artifact_id=artifact_id,
        source_sha256=artifact.sha256,
        retrieved_at=artifact.retrieved_at,
    )

    stored = storage.store_model_brand_catalog(
        tenant,
        actor,
        catalog,
        "catalog-first",
        "4" * 64,
    )

    assert stored.catalog == catalog
    assert storage.get_latest_model_brand_catalog(tenant, at=datetime.now(UTC)) == stored
    assert storage.get_latest_model_brand_catalog(other, at=datetime.now(UTC)) is None
    replay = storage.store_model_brand_catalog(
        tenant,
        actor,
        catalog,
        "catalog-first",
        "4" * 64,
    )
    assert replay == stored
    with pytest.raises(IdempotencyConflict):
        storage.store_model_brand_catalog(
            tenant,
            actor,
            catalog,
            "catalog-first",
            "5" * 64,
        )
    with pytest.raises(ModelBrandCatalogConflict):
        storage.store_model_brand_catalog(
            tenant,
            actor,
            catalog,
            "catalog-second",
            "6" * 64,
        )

    events = storage.list_audit_events(
        tenant,
        object_type="model_brand_catalog",
        object_id=catalog.catalog_id,
    )
    assert [event.event_type for event in events] == ["model_brand.catalog_stored"]
    assert events[0].basis["entries_total"] == 1
    mismatched_catalog = catalog.model_copy(
        update={
            "catalog_id": f"catalog-{uuid.uuid4()}",
            "entries": (
                catalog.entries[0].model_copy(
                    update={
                        "evidence": catalog.entries[0].evidence.model_copy(
                            update={"source_document_sha256": "a" * 64}
                        )
                    }
                ),
            ),
        }
    )
    with pytest.raises(ModelBrandSourceEvidenceMismatch):
        storage.store_model_brand_catalog(
            tenant,
            actor,
            mismatched_catalog,
            "catalog-mismatch",
            "8" * 64,
        )
    with storage.transaction(other) as connection:
        count = connection.execute(
            select(func.count()).select_from(model_brand_catalogs)
        ).scalar_one()
        assert count == 0
        artifact_count = connection.execute(
            select(func.count()).select_from(model_brand_source_artifacts)
        ).scalar_one()
        assert artifact_count == 0
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            model_brand_catalogs.update()
            .where(model_brand_catalogs.c.catalog_id == catalog.catalog_id)
            .values(catalog_id="mutated")
        )
    with storage.transaction(tenant) as connection, pytest.raises(DBAPIError):
        connection.execute(
            model_brand_source_artifacts.update()
            .where(model_brand_source_artifacts.c.artifact_id == artifact_id)
            .values(source_title="mutated")
        )

    workflow_catalog = _model_brand_catalog(
        tenant,
        "ignored-until-approval",
        f"workflow-catalog-{uuid.uuid4()}",
        artifact_id=artifact_id,
        source_sha256=artifact.sha256,
        retrieved_at=artifact.retrieved_at,
    )
    proposal = ModelBrandReviewProposal(
        proposal_id=f"proposal-{uuid.uuid4()}",
        company_id=tenant,
        submitted_by="catalog-author",
        submitted_at=datetime.now(UTC),
        catalog_id=workflow_catalog.catalog_id,
        approval_reference="review/workflow-postgres",
        entries=workflow_catalog.entries,
    )
    draft = storage.store_model_brand_review_proposal(
        tenant,
        "catalog-author",
        proposal,
        "review-submit",
        "9" * 64,
    )
    assert draft.status is ModelBrandReviewStatus.DRAFT
    assert storage.list_model_brand_reviews(other) == ()
    approved = storage.decide_model_brand_review(
        tenant,
        actor,
        proposal.proposal_id,
        ModelBrandReviewStatus.APPROVED,
        "Exact PDF evidence verified",
        "review-decision",
        "b" * 64,
    )
    assert approved.status is ModelBrandReviewStatus.APPROVED
    assert approved.decision is not None
    assert approved.decision.catalog_id == proposal.catalog_id
    with pytest.raises(ModelBrandReviewConflict):
        storage.decide_model_brand_review(
            tenant,
            actor,
            proposal.proposal_id,
            ModelBrandReviewStatus.REJECTED,
            "Cannot change terminal decision",
            "review-decision-second",
            "c" * 64,
        )
    with storage.transaction(other) as connection:
        assert (
            connection.execute(
                select(func.count()).select_from(model_brand_review_proposals)
            ).scalar_one()
            == 0
        )
        assert (
            connection.execute(
                select(func.count()).select_from(model_brand_review_decisions)
            ).scalar_one()
            == 0
        )
