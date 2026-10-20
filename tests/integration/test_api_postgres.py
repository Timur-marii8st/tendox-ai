"""End-to-end HTTP checks through the real PostgreSQL adapters."""

from __future__ import annotations

import asyncio
import hashlib
import json
import os
import uuid
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile

import httpx
import pytest
from alembic import command
from alembic.config import Config
from docx import Document
from docx.shared import Inches
from PIL import Image
from pypdf import PdfWriter
from sqlalchemy import create_engine, text

from procurement.api import create_app
from procurement.api.postgres import PostgresAuthenticator, PostgresCaseService
from procurement.documents.ai_requirements import (
    OpenRouterJsonResponse,
    analyze_requirements_with_model,
)
from procurement.documents.extraction import ExtractionResult
from procurement.storage import AiTransferConflict, PostgresStorage

PAPER_CASE = Path(__file__).parents[2] / "examples" / "paper_case.json"


@pytest.fixture(scope="module")
def database_urls() -> tuple[str, str]:
    runtime = os.environ.get("DATABASE_URL")
    admin = os.environ.get("DATABASE_ADMIN_URL")
    if not runtime or not admin:
        pytest.skip("DATABASE_URL and DATABASE_ADMIN_URL are required")
    config = Config("alembic.ini")
    config.set_main_option("sqlalchemy.url", admin.replace("%", "%%"))
    command.upgrade(config, "head")
    return runtime, admin


def _request(app: object, method: str, url: str, **kwargs: object) -> httpx.Response:
    async def send() -> httpx.Response:
        transport = httpx.ASGITransport(app=app, raise_app_exceptions=False)  # type: ignore[arg-type]
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            return await client.request(method, url, **kwargs)

    return asyncio.run(send())


def test_real_ai_transfer_request_and_independent_approval(database_urls) -> None:
    runtime_url, admin_url = database_urls
    tenant = f"transfer-api-{uuid.uuid4()}"
    preparer = _provision_key(admin_url, tenant, "preparer")
    approver = _provision_key(admin_url, tenant, "approver", "AI_TRANSFER_APPROVER")
    self_approver = _provision_key(admin_url, tenant, "preparer", "AI_TRANSFER_APPROVER")
    engine = create_engine(runtime_url, pool_pre_ping=True)
    try:
        storage = PostgresStorage(engine)
        app = create_app(PostgresCaseService(storage), PostgresAuthenticator(storage))
        doc = Document()
        doc.add_paragraph("Насос мощностью не менее 5 кВт.")
        content = BytesIO()
        doc.save(content)
        uploaded = _request(
            app,
            "POST",
            "/v1/documents",
            headers={"Authorization": f"Bearer {preparer}", "Idempotency-Key": "source"},
            data={"document_role": "TENDER_REQUIREMENTS"},
            files={
                "file": (
                    "requirements.docx",
                    content.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert uploaded.status_code == 201, uploaded.text
        path = f"/v1/documents/{uploaded.json()['document_id']}/ai-transfer-requests"
        headers = {"Authorization": f"Bearer {preparer}", "Idempotency-Key": "prepare"}
        body = {"model": "inclusionai/ling-3.0-flash-vl:free"}
        prepared = _request(app, "POST", path, headers=headers, json=body)
        assert prepared.status_code == 201, prepared.text
        assert prepared.headers["cache-control"] == "no-store"
        approval_path = f"/v1/ai-transfer-requests/{prepared.json()['transfer_id']}/approval"
        approval_body = {"request_payload_sha256": prepared.json()["payload_sha256"]}
        rejected = _request(
            app,
            "POST",
            approval_path,
            headers={"Authorization": f"Bearer {self_approver}", "Idempotency-Key": "self"},
            json=approval_body,
        )
        assert rejected.status_code == 409, rejected.text
        approval_headers = {"Authorization": f"Bearer {approver}", "Idempotency-Key": "approve"}
        approved = _request(
            app, "POST", approval_path, headers=approval_headers, json=approval_body
        )
        assert approved.status_code == 201, approved.text
        assert approved.headers["cache-control"] == "no-store"
        assert (
            _request(
                app, "POST", approval_path, headers=approval_headers, json=approval_body
            ).json()
            == approved.json()
        )
        assert _request(app, "POST", path, headers=headers, json=body).json() == prepared.json()
        detail = _request(
            app,
            "GET",
            approval_path.removesuffix("/approval"),
            headers={"Authorization": f"Bearer {approver}"},
        )
        assert detail.status_code == 200, detail.text
        assert detail.json() == {"request": prepared.json(), "approval": approved.json()}
        admission_path = approval_path.removesuffix("/approval") + "/execution-admissions"
        admission_headers = {"Authorization": f"Bearer {preparer}", "Idempotency-Key": "admit"}
        admission = _request(
            app,
            "POST",
            admission_path,
            headers=admission_headers,
            json=approval_body,
        )
        assert admission.status_code == 201, admission.text
        assert admission.headers["cache-control"] == "no-store"
        assert admission.json()["status"] == "ADMITTED"
        assert (
            _request(
                app,
                "POST",
                admission_path,
                headers=admission_headers,
                json=approval_body,
            ).json()
            == admission.json()
        )
        execution_id = admission.json()["execution_id"]
        execution_path = f"/v1/ai-transfer-executions/{execution_id}"
        admitted_detail = _request(
            app,
            "GET",
            execution_path,
            headers={"Authorization": f"Bearer {approver}"},
        )
        assert admitted_detail.status_code == 200, admitted_detail.text
        assert admitted_detail.json() == {
            "admission": admission.json(),
            "claim": None,
            "receipt": None,
            "raw_response": None,
        }
        claim = storage.claim_first_ai_transfer_dispatch(tenant, "worker", execution_id)
        assert claim.delivery_state == "DELIVERY_UNKNOWN"
        with pytest.raises(AiTransferConflict):
            storage.claim_first_ai_transfer_dispatch(tenant, "worker", execution_id)
        claimed_detail = _request(
            app,
            "GET",
            execution_path,
            headers={"Authorization": f"Bearer {preparer}"},
        )
        assert claimed_detail.status_code == 200, claimed_detail.text
        assert claimed_detail.json()["claim"] == claim.model_dump(mode="json")
    finally:
        engine.dispose()


def _case_payload(company_id: str, case_id: str) -> dict[str, Any]:
    payload = json.loads(PAPER_CASE.read_text(encoding="utf-8"))
    payload["company_id"] = company_id
    payload["case_id"] = case_id
    payload["profile"]["company_id"] = company_id
    payload["selected_quote"]["company_id"] = company_id
    return payload


def _provision_key(
    admin_url: str,
    company_id: str,
    actor_id: str,
    access_role: str = "OPERATOR",
) -> str:
    secret = f"integration-{uuid.uuid4()}"
    engine = create_engine(admin_url)
    try:
        with engine.begin() as connection:
            connection.execute(
                text(
                    "INSERT INTO api_keys "
                    "(key_id, company_id, key_digest, actor_id, access_role, enabled) "
                    "VALUES (:key_id, :company_id, :key_digest, :actor_id, :access_role, true)"
                ),
                {
                    "key_id": str(uuid.uuid4()),
                    "company_id": company_id,
                    "key_digest": hashlib.sha256(secret.encode()).hexdigest(),
                    "actor_id": actor_id,
                    "access_role": access_role,
                },
            )
    finally:
        engine.dispose()
    return secret


def _blank_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _bundle() -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("scan.pdf", _blank_pdf())
    return output.getvalue()


def _pictured_docx() -> bytes:
    image = BytesIO()
    Image.new("RGB", (80, 40), color="white").save(image, format="PNG")
    output = BytesIO()
    document = Document()
    document.add_picture(BytesIO(image.getvalue()), width=Inches(1))
    document.save(output)
    return output.getvalue()


def _requirements_docx() -> bytes:
    output = BytesIO()
    document = Document()
    document.add_paragraph("Гарантия не менее 24 мес.")
    document.save(output)
    return output.getvalue()


class _AiRequirementReviewClient:
    def complete_json(self, **kwargs: object) -> OpenRouterJsonResponse:
        model = kwargs["model"]
        assert isinstance(model, str)
        return OpenRouterJsonResponse(
            request_sha256="a" * 64,
            response_id="private-provider-response",
            model=model,
            content=json.dumps(
                {
                    "items": [],
                    "requirements": [
                        {
                            "item_name": None,
                            "property": "warranty",
                            "operator": ">=",
                            "value": "24",
                            "unit": "мес",
                            "mandatory": True,
                            "source_block_ids": ["b000001"],
                            "uncertainty": None,
                        }
                    ],
                    "unresolved_block_ids": [],
                    "document_notes": [],
                },
                ensure_ascii=False,
            ),
        )


def test_real_api_persists_independent_completeness_review(
    database_urls: tuple[str, str],
) -> None:
    runtime_url, admin_url = database_urls
    company = f"tenant-{uuid.uuid4()}"
    other_company = f"tenant-{uuid.uuid4()}"
    operator_token = _provision_key(admin_url, company, "operator")
    reviewer_token = _provision_key(
        admin_url,
        company,
        "completeness-reviewer",
        "COMPLETENESS_REVIEWER",
    )
    other_token = _provision_key(
        admin_url,
        other_company,
        "other-reviewer",
        "COMPLETENESS_REVIEWER",
    )
    engine = create_engine(runtime_url, pool_pre_ping=True)
    try:
        storage = PostgresStorage(engine)
        app = create_app(PostgresCaseService(storage), PostgresAuthenticator(storage))
        operator_auth = {"Authorization": f"Bearer {operator_token}"}
        reviewer_auth = {"Authorization": f"Bearer {reviewer_token}"}
        other_auth = {"Authorization": f"Bearer {other_token}"}
        content = _requirements_docx()
        uploaded = _request(
            app,
            "POST",
            "/v1/documents",
            headers={**operator_auth, "Idempotency-Key": "completeness-source"},
            files={
                "file": (
                    "requirements.docx",
                    content,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            data={"document_role": "TENDER_REQUIREMENTS"},
        )
        assert uploaded.status_code == 201
        document_id = uploaded.json()["document_id"]
        create_path = f"/v1/documents/{document_id}/completeness-reviews"
        create_headers = {**operator_auth, "Idempotency-Key": "completeness-review"}
        created = _request(app, "POST", create_path, headers=create_headers)
        created_replay = _request(app, "POST", create_path, headers=create_headers)
        assert created.status_code == created_replay.status_code == 201
        assert created.json() == created_replay.json()
        snapshot = created.json()["snapshot"]
        review_id = snapshot["review_id"]
        assert snapshot["mapping"]["completeness"]["status"] == "UNKNOWN"
        assert snapshot["source_document_sha256"] == hashlib.sha256(content).hexdigest()

        decision_path = f"/v1/completeness-reviews/{review_id}/decision"
        decision_body = {
            "evidence_sha256": snapshot["evidence_sha256"],
            "status": "PASS",
            "scope": "Весь native-документ и все технические требования",
            "reason": "Все извлечённые блоки проверены",
            "block_decisions": [],
        }
        assert (
            _request(
                app,
                "POST",
                decision_path,
                headers={**operator_auth, "Idempotency-Key": "operator-cannot-decide"},
                json=decision_body,
            ).status_code
            == 403
        )
        decision_headers = {**reviewer_auth, "Idempotency-Key": "completeness-decision"}
        decided = _request(
            app,
            "POST",
            decision_path,
            headers=decision_headers,
            json=decision_body,
        )
        decided_replay = _request(
            app,
            "POST",
            decision_path,
            headers=decision_headers,
            json=decision_body,
        )
        assert decided.status_code == decided_replay.status_code == 200
        assert decided.json() == decided_replay.json()
        assert decided.json()["decision"]["status"] == "PASS"
        assert decided.json()["decision"]["decided_by"] == "completeness-reviewer"
        assert (
            _request(
                app,
                "GET",
                f"/v1/completeness-reviews/{review_id}",
                headers=other_auth,
            ).status_code
            == 404
        )
        listed = _request(app, "GET", "/v1/completeness-reviews", headers=reviewer_auth)
        assert listed.status_code == 200
        assert any(item["snapshot"]["review_id"] == review_id for item in listed.json())
        assert "content" not in listed.text
    finally:
        engine.dispose()


def test_real_api_lists_and_decides_sanitized_ai_requirement_runs(
    database_urls: tuple[str, str],
) -> None:
    runtime_url, admin_url = database_urls
    company = f"tenant-{uuid.uuid4()}"
    operator_token = _provision_key(admin_url, company, "ai-submitter")
    reviewer_token = _provision_key(admin_url, company, "ai-reviewer", "AI_REQUIREMENTS_REVIEWER")
    engine = create_engine(runtime_url, pool_pre_ping=True)
    try:
        storage = PostgresStorage(engine)
        app = create_app(PostgresCaseService(storage), PostgresAuthenticator(storage))
        operator_auth = {"Authorization": f"Bearer {operator_token}"}
        reviewer_auth = {"Authorization": f"Bearer {reviewer_token}"}
        content = _requirements_docx()
        uploaded = _request(
            app,
            "POST",
            "/v1/documents",
            headers={**operator_auth, "Idempotency-Key": "ai-source"},
            files={
                "file": (
                    "requirements.docx",
                    content,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            data={"document_role": "TENDER_REQUIREMENTS"},
        )
        assert uploaded.status_code == 201
        document = uploaded.json()
        extraction = ExtractionResult.model_validate(document["extraction"])
        report = analyze_requirements_with_model(
            extraction,
            company_id=company,
            source_sha256=document["sha256"],
            authorization_reference="private-import-approval",
            client=_AiRequirementReviewClient(),
            created_at=datetime.now(UTC),
        )
        import_path = f"/v1/documents/{document['document_id']}/ai-requirement-runs"
        imported = _request(
            app,
            "POST",
            import_path,
            headers={**operator_auth, "Idempotency-Key": "ai-import"},
            json={"report": report.model_dump(mode="json")},
        )
        replayed_import = _request(
            app,
            "POST",
            import_path,
            headers={**operator_auth, "Idempotency-Key": "ai-import"},
            json={"report": report.model_dump(mode="json")},
        )
        assert imported.status_code == replayed_import.status_code == 201
        assert imported.json() == replayed_import.json()
        assert imported.headers["cache-control"] == "no-store"
        run_id = imported.json()["run_id"]
        listed = _request(app, "GET", "/v1/ai-requirement-runs", headers=operator_auth)
        assert listed.status_code == 200
        assert listed.headers["cache-control"] == "no-store"
        assert listed.json()[0]["run_id"] == run_id
        assert "authorization_reference" not in listed.text
        detail = _request(app, "GET", f"/v1/ai-requirement-runs/{run_id}", headers=reviewer_auth)
        assert detail.status_code == 200
        assert detail.json()["requirement_candidates"][0]["evidence"]
        assert "response_content" not in detail.text
        decision = _request(
            app,
            "POST",
            f"/v1/ai-requirement-runs/{run_id}/decision",
            headers={**reviewer_auth, "Idempotency-Key": "ai-review-decision"},
            json={
                "evidence_sha256": detail.json()["evidence_sha256"],
                "candidate_decisions": [
                    {
                        "candidate_kind": "REQUIREMENT",
                        "candidate_id": detail.json()["requirement_candidates"][0]["candidate_id"],
                        "disposition": "ACCEPTED",
                        "reason": "Exact evidence checked.",
                    }
                ],
            },
        )
        assert decision.status_code == 200
        assert decision.json()["completeness_status"] == "UNKNOWN"
    finally:
        engine.dispose()


def test_real_api_auth_rls_idempotency_and_atomic_evaluation(
    database_urls: tuple[str, str],
) -> None:
    runtime_url, admin_url = database_urls
    company_a = f"tenant-{uuid.uuid4()}"
    company_b = f"tenant-{uuid.uuid4()}"
    case_id = f"case-{uuid.uuid4()}"
    token_a = _provision_key(admin_url, company_a, "operator-a")
    token_b = _provision_key(admin_url, company_b, "operator-b")
    engine = create_engine(runtime_url, pool_pre_ping=True)
    try:
        storage = PostgresStorage(engine)
        app = create_app(PostgresCaseService(storage), PostgresAuthenticator(storage))
        auth_a = {"Authorization": f"Bearer {token_a}"}
        auth_b = {"Authorization": f"Bearer {token_b}"}
        payload = _case_payload(company_a, case_id)

        first = _request(
            app,
            "POST",
            "/v1/cases",
            headers={**auth_a, "Idempotency-Key": "create-1"},
            json=payload,
        )
        replay = _request(
            app,
            "POST",
            "/v1/cases",
            headers={**auth_a, "Idempotency-Key": "create-1"},
            json=payload,
        )
        conflict = _request(
            app,
            "POST",
            "/v1/cases",
            headers={**auth_a, "Idempotency-Key": "create-2"},
            json=payload,
        )

        assert first.status_code == replay.status_code == 201
        assert first.json() == replay.json()
        assert conflict.status_code == 409
        assert (
            _request(
                app,
                "GET",
                f"/v1/cases/{case_id}/versions/1",
                headers=auth_b,
            ).status_code
            == 404
        )
        malformed_cursor = _request(
            app,
            "GET",
            "/v1/cases?cursor=not-a-real-cursor",
            headers=auth_a,
        )
        assert malformed_cursor.status_code == 422
        assert malformed_cursor.json() == {"detail": "validation_error"}

        evaluation_path = f"/v1/cases/{case_id}/versions/1/evaluations"
        latest_evaluation_path = f"{evaluation_path}/latest"
        no_evaluation = _request(app, "GET", latest_evaluation_path, headers=auth_a)
        assert no_evaluation.status_code == 404
        evaluation_headers = {**auth_a, "Idempotency-Key": "evaluate-1"}
        evaluation_body = {"effective_at": "2026-09-08T12:00:00+03:00"}
        evaluated = _request(
            app,
            "POST",
            evaluation_path,
            headers=evaluation_headers,
            json=evaluation_body,
        )
        evaluated_replay = _request(
            app,
            "POST",
            evaluation_path,
            headers=evaluation_headers,
            json=evaluation_body,
        )

        assert evaluated.status_code == evaluated_replay.status_code == 201
        assert evaluated.json() == evaluated_replay.json()
        evaluation_id = evaluated.json()["evaluation_id"]
        assert evaluated.json()["audit_event_id"] == evaluation_id
        latest = _request(app, "GET", latest_evaluation_path, headers=auth_a)
        cross_tenant_latest = _request(app, "GET", latest_evaluation_path, headers=auth_b)
        missing_latest = _request(
            app,
            "GET",
            "/v1/cases/missing/versions/1/evaluations/latest",
            headers=auth_b,
        )
        assert latest.status_code == 200
        assert latest.json() == evaluated.json()
        assert cross_tenant_latest.status_code == missing_latest.status_code == 404
        assert cross_tenant_latest.json() == missing_latest.json() == {"detail": "not_found"}
        assert len(storage.list_audit_events(company_a, object_id=case_id)) == 1
        assert storage.list_audit_events(company_b, object_id=case_id) == []

        document_headers = {**auth_a, "Idempotency-Key": "document-1"}
        document_file = {"file": ("scan.pdf", _blank_pdf(), "application/pdf")}
        uploaded = _request(
            app,
            "POST",
            "/v1/documents",
            headers=document_headers,
            files=document_file,
            data={"document_role": "UNKNOWN"},
        )
        uploaded_replay = _request(
            app,
            "POST",
            "/v1/documents",
            headers=document_headers,
            files=document_file,
            data={"document_role": "UNKNOWN"},
        )
        assert uploaded.status_code == uploaded_replay.status_code == 201
        assert uploaded.json() == uploaded_replay.json()
        document_id = uploaded.json()["document_id"]
        assert uploaded.json()["document_role"] == "UNKNOWN"
        assert uploaded.json()["extraction"]["status"] == "UNKNOWN"
        assert "content" not in uploaded.json()
        assert (
            _request(
                app,
                "GET",
                f"/v1/documents/{document_id}",
                headers=auth_b,
            ).status_code
            == 404
        )
        document_events = storage.list_audit_events(company_a, object_id=document_id)
        assert len(document_events) == 1
        assert document_events[0].basis["document_role"] == "UNKNOWN"

        docx_content = _pictured_docx()
        docx_upload = _request(
            app,
            "POST",
            "/v1/documents",
            headers={**auth_a, "Idempotency-Key": "docx-media-source-1"},
            files={
                "file": (
                    "requirements.docx",
                    docx_content,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            data={"document_role": "TENDER_REQUIREMENTS"},
        )
        assert docx_upload.status_code == 201
        docx_id = docx_upload.json()["document_id"]
        docx_ocr_path = f"/v1/documents/{docx_id}/media/0/ocr"
        docx_media_ocr = _request(
            app,
            "POST",
            docx_ocr_path,
            headers={**auth_a, "Idempotency-Key": "docx-media-ocr-1"},
            json={"threshold": 0.75},
        )
        assert docx_media_ocr.status_code == 201
        assert docx_media_ocr.json()["result"]["status"] == "UNKNOWN"
        assert (
            docx_media_ocr.json()["source_document_sha256"]
            == hashlib.sha256(docx_content).hexdigest()
        )
        assert "content" not in docx_media_ocr.text
        assert (
            _request(
                app,
                "GET",
                f"{docx_ocr_path}/latest",
                headers=auth_b,
            ).status_code
            == 404
        )
        assert {
            event.event_type for event in storage.list_audit_events(company_a, object_id=docx_id)
        } == {"document.stored", "document.docx_media_ocr_stored"}

        bundle_headers = {**auth_a, "Idempotency-Key": "bundle-1"}
        bundle_file = {"file": ("tender.zip", _bundle(), "application/zip")}
        bundle = _request(
            app,
            "POST",
            "/v1/document-bundles",
            headers=bundle_headers,
            files=bundle_file,
            data={"document_role": "TENDER_REQUIREMENTS"},
        )
        bundle_replay = _request(
            app,
            "POST",
            "/v1/document-bundles",
            headers=bundle_headers,
            files=bundle_file,
            data={"document_role": "TENDER_REQUIREMENTS"},
        )
        assert bundle.status_code == bundle_replay.status_code == 201
        assert bundle.json() == bundle_replay.json()
        bundle_id = bundle.json()["bundle_id"]
        assert bundle.json()["result"]["status"] == "UNKNOWN"
        assert bundle.json()["result"]["reasons"] == ["MEMBER_EXTRACTION_UNKNOWN"]
        assert "content" not in bundle.json()
        assert (
            _request(
                app,
                "GET",
                f"/v1/document-bundles/{bundle_id}",
                headers=auth_b,
            ).status_code
            == 404
        )
        bundle_events = storage.list_audit_events(company_a, object_id=bundle_id)
        assert len(bundle_events) == 1
        assert bundle_events[0].event_type == "document_bundle.stored"

        member_ocr_path = f"/v1/document-bundles/{bundle_id}/members/0/ocr"
        member_ocr_headers = {**auth_a, "Idempotency-Key": "bundle-member-ocr-1"}
        member_ocr = _request(
            app,
            "POST",
            member_ocr_path,
            headers=member_ocr_headers,
            json={"threshold": 0.75},
        )
        member_ocr_replay = _request(
            app,
            "POST",
            member_ocr_path,
            headers=member_ocr_headers,
            json={"threshold": 0.75},
        )
        assert member_ocr.status_code == member_ocr_replay.status_code == 201
        assert member_ocr.json() == member_ocr_replay.json()
        assert member_ocr.json()["result"]["status"] == "UNKNOWN"
        assert member_ocr.json()["source_bundle_id"] == bundle_id
        assert (
            _request(
                app,
                "GET",
                f"{member_ocr_path}/latest",
                headers=auth_b,
            ).status_code
            == 404
        )
        bundle_events = storage.list_audit_events(company_a, object_id=bundle_id)
        assert {event.event_type for event in bundle_events} == {
            "document_bundle.stored",
            "document_bundle.member_ocr_stored",
        }
    finally:
        engine.dispose()
