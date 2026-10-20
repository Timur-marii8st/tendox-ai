"""Only an explicitly isolated PostgreSQL DB and synthetic 44/223 sources."""

import os
from datetime import UTC, datetime
from uuid import uuid4

import pytest
from alembic import command
from alembic.config import Config
from sqlalchemy import create_engine, func, select, text
from test_eis_document_queue_pg import admit

from procurement.sources.eis_223_documents import manifest223
from procurement.storage.eis_223_index import ingest_notice223_window, search_notice223_index
from procurement.storage.eis_document_store import get_document_status
from procurement.storage.eis_evidence_bridge import get_eis_evidence_bridge
from procurement.storage.eis_mixed_comparison import MixedComparisonRequest, get_mixed_comparison
from procurement.storage.models import (
    audit_events,
    eis_223_document_attempts,
    eis_223_document_jobs,
    eis_223_document_receipts,
    eis_223_notice_versions,
    eis_223_schema_checks,
    eis_document_analyses,
    eis_document_jobs,
    eis_document_observations,
    eis_notice_position_projections,
    eis_notice_versions,
)
from procurement.storage.repository import PostgresStorage
from test_eis_223_archive import xml, zip_bytes
from test_eis_223_index import window

NOW = datetime(2026, 10, 2, tzinfo=UTC)


def test_source_estimates_cached_exact_lot_tenant_and_no_writes(storage, monkeypatch):
    from decimal import Decimal

    from procurement.sources import eis_223_archive, eis_feed_archive
    from procurement.storage.eis_source_estimate import (
        SourceEstimateRequest,
        SourceLotError,
        get_source_estimate,
    )

    company = "test-estimate-" + uuid4().hex
    ref44 = admit(storage, company, 0, listed=False, deadline="2026-10-09T10:00:00Z")
    second = """<t:lot><t:ordinalNumber>7</t:ordinalNumber><t:cancelled>false</t:cancelled>
      <t:lotData><t:subject>Бумага, второй лот</t:subject>
      <t:currency><t:code>RUB</t:code><t:digitalCode>643</t:digitalCode></t:currency>
      <t:initialSum>240</t:initialSum></t:lotData></t:lot>"""
    ref223 = admit223(storage, company, xml(second_lot=second))
    costs = {
        "goods": "50",
        "expenses": "10",
        "vat_rate": "0.20",
        "basis": "Синтетические затраты для проверки, не КП",
        "costs_complete_ack": True,
    }

    def forbidden(*args, **kwargs):
        raise AssertionError("Source estimates cannot parse or download")

    monkeypatch.setattr(eis_223_archive, "parse_notice223_xml", forbidden)
    monkeypatch.setattr(eis_feed_archive, "_notice", forbidden)
    before = counts(storage, company)
    payload = SourceEstimateRequest.model_validate(
        {
            "notice": {"law": "223", **ref223.model_dump()},
            "lot_index": 2,
            "costs": costs,
        }
    )
    result = get_source_estimate(storage, company, payload, actor_id="test", now=NOW)
    replay = get_source_estimate(storage, company, payload, actor_id="test", now=NOW)
    assert replay == result
    assert result.ceiling_rub == Decimal("240") and result.lot_ordinal.value == "7"
    assert result.ceiling.provenance.locator.endswith("/t:lot[2]/t:lotData/t:initialSum")
    assert result.estimate.scenarios[0].profit == Decimal("140")
    assert result.eligibility == result.quote_status == "UNKNOWN"
    with pytest.raises(LookupError):
        get_source_estimate(storage, company + "-other", payload, actor_id="test", now=NOW)
    with pytest.raises(ValueError):
        get_source_estimate(
            storage,
            company,
            payload.model_copy(
                update={"notice": payload.notice.model_copy(update={"notice_xml_sha256": "f" * 64})}
            ),
            actor_id="test",
            now=NOW,
        )
    with pytest.raises(SourceLotError):
        get_source_estimate(
            storage, company, payload.model_copy(update={"lot_index": 3}), actor_id="test", now=NOW
        )
    request44 = SourceEstimateRequest.model_validate(
        {"notice": {"law": "44", **ref44.model_dump()}, "costs": costs}
    )
    missing = get_source_estimate(storage, company, request44, actor_id="test", now=NOW)
    assert missing.ceiling_rub is None and missing.estimate.status == "UNKNOWN"
    assert counts(storage, company) == before
    # A new saved version blocks the old one, even if illustrative costs are zero.
    monkeypatch.undo()
    admit(storage, company, 0, listed=False, version=2, deadline="2026-10-09T10:00:00Z")
    monkeypatch.setattr(eis_223_archive, "parse_notice223_xml", forbidden)
    monkeypatch.setattr(eis_feed_archive, "_notice", forbidden)
    before = counts(storage, company)
    stale = get_source_estimate(storage, company, request44, actor_id="test", now=NOW)
    assert stale.estimate.status == "BLOCKED" and "NOTICE_VERSION_STALE" in stale.blocking_reasons
    assert counts(storage, company) == before


def test_position_facts_native_tables_and_links_share_read_only_source_snapshot(
    storage, monkeypatch
):
    from io import BytesIO

    from docx import Document
    from test_eis_document_queue_pg import UID, URL

    from procurement.storage import eis_223_documents as docs223
    from procurement.storage import eis_document_store as docs44
    from procurement.storage.eis_document_store import (
        NoticeDocumentRef,
        plan_document_job,
        record_document_outcome,
    )
    from procurement.storage.eis_notice_index import EisReadyWindow, ingest_ready_window
    from test_eis_feed_archive import _archive, _notice

    company = "test-positions-" + uuid4().hex
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Характеристика"
    table.cell(0, 1).text = "Значение"
    table.cell(1, 0).text = "Листов в пачке"
    table.cell(1, 1).text = "Не менее 500"
    output = BytesIO()
    doc.save(output)
    content = output.getvalue()
    number = "0848300054826000999"
    raw = _notice(number).decode().replace("182420.00", "30.00")
    raw = raw.replace(
        "<notificationInfo>",
        f"""<attachmentsInfo><attachmentInfo>
      <publishedContentId>{UID}</publishedContentId><fileName>spec.docx</fileName>
      <fileSize>{len(content)}</fileSize><url>{URL}</url></attachmentInfo></attachmentsInfo>
      <notificationInfo>""",
    )
    raw = raw.replace(
        "</notificationInfo>",
        """<purchaseObjectsInfo><notDrugPurchaseObjectsInfo>
      <purchaseObject><name>Бумага А4 80</name><quantity><value>3</value></quantity>
      <OKEI><name>Штука</name></OKEI><price>10</price><sum>30</sum>
      <KTRU><code>17.12.14.110-00000001</code><OKPD2><OKPDCode>17.12.14.110</OKPDCode></OKPD2>
      <characteristics><characteristicsUsingReferenceInfo><name>Листов в пачке</name>
      <values><value><rangeSet><valueRange><min>500</min><minMathNotation>greaterOrEqual</minMathNotation>
      </valueRange></rangeSet><OKEI><name>Штука</name></OKEI></value></values>
      </characteristicsUsingReferenceInfo></characteristics></KTRU></purchaseObject>
      </notDrugPurchaseObjectsInfo></purchaseObjectsInfo></notificationInfo>""",
    )
    ingest_ready_window(
        storage,
        company,
        EisReadyWindow(
            "getDocsIP/PRIZ/50/2026-10-02+03:00",
            "a" * 64,
            1,
            (_archive(("notice.xml", raw.encode())),),
        ),
        actor_id="test",
    )
    with storage.transaction(company) as conn:
        digest = conn.scalar(select(eis_notice_versions.c.xml_sha256))
    ref44 = NoticeDocumentRef(purchase_number=number, version_number=1, notice_xml_sha256=digest)
    plan_document_job(storage, company, ref44, actor_id="test")
    record_document_outcome(storage, company, ref44, UID, actor_id="test", content=content)
    ref223 = admit223(storage, company, xml(code="USD", numeric="840"))
    monkeypatch.setattr(docs223, "resource_gate", lambda *a, **k: None)
    docs223.collect_document223(
        storage, company, ref223, UID, actor_id="test", downloader=lambda *a: content
    )

    def forbidden(*a, **k):
        raise AssertionError("Read-only position comparison cannot parse or download")

    monkeypatch.setattr(docs223, "analyze_document", forbidden)
    monkeypatch.setattr(docs44, "analyze_document", forbidden)
    request = MixedComparisonRequest.model_validate(
        {"notices": [{"law": "44", **ref44.model_dump()}, {"law": "223", **ref223.model_dump()}]}
    )
    before = counts(storage, company)
    result = get_mixed_comparison(storage, company, request, now=NOW)
    p44 = result.rows[0].evidence.position_set
    p223 = result.rows[1].position_set
    assert p44.positions[0].characteristics[0].values[0].minimum.value == 500
    assert p44.positions[0].characteristics[0].values[0].minimum_operator.value == "greaterOrEqual"
    assert p223.positions[0].currency.value == "USD" and p223.positions[0].unit_price.value is None
    for p in (p44, p223):
        assert p.native_tables[0].rows[1].cells[1].text == "Не менее 500"
        assert p.native_tables[0].position_association == "UNKNOWN"
        assert p.positions[0].quantity.provenance.document_id == p.notice_xml_sha256
    assert result.position_links and result.position_links[0].equivalence == "UNKNOWN"
    assert result.profit_order_status == result.economic_comparability == "UNKNOWN"
    assert counts(storage, company) == before


@pytest.fixture(scope="module")
def storage():
    app_url, admin_url = os.getenv("DATABASE_URL"), os.getenv("DATABASE_ADMIN_URL")
    if not app_url or not admin_url:
        pytest.skip("isolated PostgreSQL required")
    if "procurement_eis_test_" not in app_url or "procurement_eis_test_" not in admin_url:
        pytest.fail("mixed gate requires explicitly isolated test database")
    config = Config("alembic.ini")
    config.set_main_option("sqlalchemy.url", admin_url.replace("%", "%%"))
    command.upgrade(config, "head")
    engine = create_engine(app_url)
    yield PostgresStorage(engine)
    engine.dispose()


def counts(storage, company):
    with storage.transaction(company, consistent_read=True) as conn:
        return tuple(
            conn.scalar(select(func.count()).select_from(table))
            for table in (
                audit_events,
                eis_223_notice_versions,
                eis_223_schema_checks,
                eis_223_document_receipts,
                eis_223_document_jobs,
                eis_223_document_attempts,
                eis_notice_versions,
                eis_notice_position_projections,
                eis_document_jobs,
                eis_document_observations,
                eis_document_analyses,
            )
        )


def admit223(storage, company, raw):
    ingest_notice223_window(
        storage, company, window(zip_bytes([("notice.xml", raw)])), actor_id="test"
    )
    notice = search_notice223_index(storage, company, "бумага", now=NOW).matched[0].notice
    return manifest223(notice).notice


def test_legacy_position_projection_immutable_tenant_bound_replay_and_fast_read(
    storage, monkeypatch
):
    from dataclasses import replace

    from sqlalchemy import insert, update
    from sqlalchemy.exc import DBAPIError

    from procurement.storage import eis_notice_index as intake
    from procurement.storage import eis_position_projection as projection
    from procurement.storage.eis_document_store import NoticeDocumentRef
    from test_eis_feed_archive import _archive, numeric_notice

    company = "test-projection-" + uuid4().hex
    raw = numeric_notice()
    window44 = intake.EisReadyWindow(
        "getDocsIP/PRIZ/50/2026-10-02+03:00", "a" * 64, 1, (_archive(("numeric.xml", raw)),)
    )
    preflight = intake.preflight_ready_window(window44)
    notice = preflight.inspections[0].records[0]
    old = replace(
        notice,
        lot_items=tuple(
            replace(
                i,
                characteristics=tuple(
                    replace(
                        c,
                        values=tuple(
                            replace(v, concrete_values=(), minimum=None) for v in c.values
                        ),
                    )
                    for c in i.characteristics
                ),
            )
            for i in notice.lot_items
        ),
    )
    legacy = replace(preflight, inspections=(replace(preflight.inspections[0], records=(old,)),))
    monkeypatch.setattr(intake, "preflight_ready_window", lambda *a: legacy)
    intake.ingest_ready_window(storage, company, window44, actor_id="test")
    ref44 = NoticeDocumentRef(
        purchase_number=notice.purchase_number,
        version_number=1,
        notice_xml_sha256=notice.xml_sha256,
    )
    ref223 = admit223(storage, company, xml())
    with storage.transaction(company, consistent_read=True) as conn:
        original = conn.execute(select(eis_notice_versions)).mappings().one()
        assert not original["lot_items"][0]["characteristics"][1]["values"][0]["concrete_values"]
    assert projection.project_notice_positions(storage, company, ref44, actor_id="test")
    before = counts(storage, company)
    with storage.transaction(company, consistent_read=True) as conn:
        row = conn.execute(select(eis_notice_position_projections)).mappings().one()
        assert conn.execute(select(eis_notice_versions)).mappings().one() == original
    with pytest.raises(LookupError, match="NOTICE_NOT_FOUND"):
        projection.project_notice_positions(storage, "other-" + company, ref44, actor_id="test")
    with pytest.raises(ValueError, match="NOTICE_DIGEST_CONFLICT"):
        projection.project_notice_positions(
            storage,
            company,
            ref44.model_copy(update={"notice_xml_sha256": "f" * 64}),
            actor_id="test",
        )
    with storage.transaction("other-" + company, consistent_read=True) as conn:
        assert conn.scalar(select(func.count()).select_from(eis_notice_position_projections)) == 0
    bad = dict(row)
    bad["parser_version"] = "synthetic-conflict"
    bad["notice_xml_sha256"] = "f" * 64
    bad["position_set"] = {**bad["position_set"], "notice_xml_sha256": "f" * 64}
    with pytest.raises(DBAPIError), storage.transaction(company) as conn:
        conn.execute(insert(eis_notice_position_projections).values(**bad))
    with pytest.raises(DBAPIError), storage.transaction(company) as conn:
        conn.execute(update(eis_notice_position_projections).values(parser_version="changed"))

    def forbidden(*a, **k):
        pytest.fail("Cache replay/comparison cannot parse source bytes")

    monkeypatch.setattr(projection, "_notice", forbidden)
    monkeypatch.setattr(projection, "ZipFile", forbidden)
    assert not projection.project_notice_positions(storage, company, ref44, actor_id="test")
    request = MixedComparisonRequest.model_validate(
        {"notices": [{"law": "44", **ref44.model_dump()}, {"law": "223", **ref223.model_dump()}]}
    )
    result = get_mixed_comparison(storage, company, request, now=NOW)
    value = result.rows[0].evidence.position_set.positions[0].characteristics[1].values[0]
    assert value.minimum.value == -40 and value.maximum.value == 0
    assert [x.value for x in value.concrete_values] == [30, 400]
    assert value.concrete_values[1].provenance.locator.endswith("valueSet/concreteValue[2]")
    assert result.rows[0].evidence.position_set.review_status == "UNKNOWN"
    assert counts(storage, company) == before


def test_mixed_one_snapshot_currency_lots_criteria_unknown_and_no_mutation(storage, monkeypatch):
    from procurement.storage import eis_evidence_bridge as bridge
    from procurement.storage import eis_mixed_comparison as mixed

    company = "test-mixed-" + uuid4().hex
    from test_eis_document_queue_pg import UID, docx_bytes

    from procurement.storage import eis_223_documents as documents223
    from procurement.storage.eis_document_store import plan_document_job, record_document_outcome

    content = docx_bytes()
    ref44 = admit(
        storage, company, len(content), competition=True, deadline="2026-10-09T08:00:00+03:00"
    )
    plan_document_job(storage, company, ref44, actor_id="test")
    record_document_outcome(storage, company, ref44, UID, actor_id="test", content=content)
    # Two different lot prices/currencies must remain separate, never added or converted.
    raw = xml(code="USD", numeric="840")
    start, end = raw.index(b"<t:lot>"), raw.index(b"</t:lot>") + len(b"</t:lot>")
    lot2 = raw[start:end].replace(b"<t:ordinalNumber>1", b"<t:ordinalNumber>2")
    lot2 = lot2.replace(b"100.25", b"900.37").replace(b"USD", b"RUB").replace(b"840", b"643")
    ref223 = admit223(storage, company, raw[:end] + lot2 + raw[end:])
    monkeypatch.setattr(documents223, "resource_gate", lambda *a, **k: None)
    documents223.collect_document223(
        storage, company, ref223, UID, actor_id="test", downloader=lambda *a: content
    )
    before = counts(storage, company)
    reads = []
    original44, original223 = mixed.read_eis_evidence_bridge, mixed.read_comparison223_row

    def observe(fn):
        def read(conn, *args, **kwargs):
            assert conn.scalar(text("SHOW transaction_read_only")) == "on"
            assert conn.scalar(text("SHOW transaction_isolation")) == "repeatable read"
            reads.append(id(conn))
            return fn(conn, *args, **kwargs)

        return read

    monkeypatch.setattr(mixed, "read_eis_evidence_bridge", observe(original44))
    monkeypatch.setattr(mixed, "read_comparison223_row", observe(original223))
    monkeypatch.setattr(
        bridge, "_stored_competition_criteria", lambda *a: pytest.fail("archive reread")
    )
    request = MixedComparisonRequest.model_validate(
        {
            "notices": [
                {"law": "223", **ref223.model_dump()},
                {"law": "44", **ref44.model_dump()},
            ]
        }
    )
    result = get_mixed_comparison(storage, company, request, now=NOW)
    assert len(reads) == 2 and len(set(reads)) == 1
    assert [r.law for r in result.rows] == ["223", "44"] and result.checked_at == NOW
    row223, row44 = result.rows
    assert [str(lot.initial_sum.value) for lot in row223.notice.lots] == ["100.25", "900.37"]
    assert [lot.currency.value for lot in row223.notice.lots] == ["USD", "RUB"]
    assert row223.notice.lots[0].initial_sum_rub is None
    assert row44.evidence.competition_criteria.source_status == "UNKNOWN"
    assert "CRITERIA_NOT_READ_IN_FAST_COMPARISON" in row44.evidence.competition_criteria.reasons
    assert (
        row44.readiness_status == "NEEDS_ANALYSIS" and row223.readiness_status == "NEEDS_ANALYSIS"
    )
    assert result.profit_order_status == result.economic_comparability == "UNKNOWN"
    assert row44.evidence.documents[0].conditions
    assert row223.documents[0].snippets and row223.documents[0].review_status == "UNKNOWN"
    assert counts(storage, company) == before
    assert get_document_status(storage, company, ref44).notice == ref44
    assert counts(storage, company) == before


def test_mixed_stale_expired_conflict_and_tenant_never_partial(storage):
    company = "test-mixed-" + uuid4().hex
    old = admit(storage, company, 1, version=1)
    admit(storage, company, 1, version=2)
    expired = admit223(storage, company, xml(deadline="2026-01-01T08:00:00+03:00"))
    request = MixedComparisonRequest.model_validate(
        {
            "notices": [
                {"law": "44", **old.model_dump()},
                {"law": "223", **expired.model_dump()},
            ]
        }
    )
    before = counts(storage, company)
    result = get_mixed_comparison(storage, company, request, now=NOW)
    assert result.rows[0].blocking_reasons == ("NOTICE_VERSION_STALE",)
    assert result.rows[1].blocking_reasons == ("DEADLINE_EXPIRED",)
    assert all(row.readiness_status == "BLOCKED" for row in result.rows)
    with pytest.raises(LookupError):
        get_mixed_comparison(storage, "other-" + company, request, now=NOW)
    for i in (0, 1):
        refs = list(request.notices)
        refs[i] = refs[i].model_copy(update={"notice_xml_sha256": "f" * 64})
        with pytest.raises(ValueError, match="NOTICE_DIGEST_CONFLICT"):
            get_mixed_comparison(
                storage, company, MixedComparisonRequest(notices=tuple(refs)), now=NOW
            )
    assert counts(storage, company) == before


def test_existing44_bridge_still_extracts_competition_criteria(storage):
    company = "test-mixed-" + uuid4().hex
    ref = admit(storage, company, 1, competition=True)
    before = counts(storage, company)
    result = get_eis_evidence_bridge(storage, company, ref, now=NOW)
    assert result.competition_criteria.source_status == "EXTRACTED"
    assert counts(storage, company) == before


def test_mixed_official44_cancellation_remains_blocked(storage):
    from test_eis_notice_index_pg import event_zip

    from procurement.storage.eis_notice_index import EisReadyWindow, ingest_ready_window

    company = "test-mixed-" + uuid4().hex
    ref44 = admit(storage, company, 1, deadline="2026-10-09T08:00:00+03:00")
    ref223 = admit223(storage, company, xml())
    event_window = EisReadyWindow(
        "getDocsIP/PRIZ/50/2026-09-25/03+03:00/epNotificationCancel",
        "c" * 64,
        1,
        (event_zip(),),
    )
    ingest_ready_window(storage, company, event_window, actor_id="test")
    before = counts(storage, company)
    request = MixedComparisonRequest.model_validate(
        {
            "notices": [
                {"law": "223", **ref223.model_dump()},
                {"law": "44", **ref44.model_dump()},
            ]
        }
    )
    result = get_mixed_comparison(storage, company, request, now=NOW)
    assert result.rows[1].readiness_status == "BLOCKED"
    assert result.rows[1].blocking_reasons == ("PROCEDURE_BLOCKED",)
    assert result.rows[1].evidence.application_status == "FAIL"
    assert any("event.xml" in reason for reason in result.rows[1].evidence.application_reasons)
    assert result.profit_order_status == "UNKNOWN" and counts(storage, company) == before
