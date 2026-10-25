"""Persisted exact-version jobs never turn missing files or extraction into admission."""

from __future__ import annotations

import os
import uuid
from datetime import UTC, datetime, timedelta
from io import BytesIO
from zipfile import ZipFile

import pytest
from alembic import command
from alembic.config import Config
from docx import Document
from sqlalchemy import create_engine, func, select, update
from sqlalchemy.exc import DBAPIError

from procurement.sources.eis_document_worker import run_document_step
from procurement.sources.eis_public import EisPublicError
from procurement.storage.eis_document_store import (
    DocumentSelectionRequest,
    NoticeDocumentRef,
    document_read_gate,
    get_document_analysis,
    get_document_status,
    plan_document_job,
    plan_top_document_jobs,
    reanalyze_received_document,
    record_document_outcome,
)
from procurement.storage.eis_evidence_bridge import get_eis_evidence_bridge
from procurement.storage.eis_notice_index import EisReadyWindow, ingest_ready_window
from procurement.storage.models import (
    audit_events,
    eis_document_analyses,
    eis_document_blobs,
    eis_document_jobs,
    eis_document_observations,
    eis_notice_versions,
)
from procurement.storage.repository import PostgresStorage

NUMBER = "0848300054826000154"
UID = "A" * 32
URL = f"https://zakupki.gov.ru/44fz/filestore/public/1.0/download/priz/file.html?uid={UID}"


def test_slash_source_filename_job_blocks_before_download(storage: PostgresStorage) -> None:
    company = f"slash-job-{uuid.uuid4()}"
    ref = admit(storage, company, 6, filename="Часть 1/2.pdf")
    plan_document_job(storage, company, ref, actor_id="test")

    def forbidden(_: str) -> bytes:
        raise AssertionError("display label must be reviewed before download")

    probe = run_document_step(
        storage, company, actor_id="probe", downloader=forbidden, min_free_bytes=1, probe_only=True
    )
    assert probe == {
        "status": "BLOCKED",
        "reason": "ATTACHMENT_FILENAME_REQUIRES_REVIEW",
        "attempted": 0,
    }
    report = get_document_status(storage, company, ref)
    assert report.documents[0].fetch_status == "PENDING"
    result = run_document_step(
        storage, company, actor_id="test", downloader=forbidden, min_free_bytes=1
    )
    assert result == probe
    report = get_document_status(storage, company, ref)
    assert report.documents[0].filename == "Часть 1/2.pdf"
    assert report.documents[0].reasons == ("ATTACHMENT_FILENAME_REQUIRES_REVIEW",)
    assert report.extraction_status == "UNKNOWN"
    with storage.transaction(company) as conn:
        assert conn.scalar(select(func.count()).select_from(eis_document_blobs)) == 0
        assert conn.scalar(select(func.count()).select_from(eis_document_analyses)) == 0


def test_slash_label_cached_reanalysis_is_unknown(
    storage: PostgresStorage, monkeypatch: pytest.MonkeyPatch
) -> None:
    from procurement.storage import eis_document_store

    company = f"slash-reanalysis-{uuid.uuid4()}"
    ref = admit(storage, company, 6, filename="Часть 1/2.pdf")
    plan_document_job(storage, company, ref, actor_id="test")
    record_document_outcome(storage, company, ref, UID, actor_id="test", content=b"source")
    report = get_document_status(storage, company, ref)
    assert report.extraction_status == "UNKNOWN"
    assert report.documents[0].reasons == ("ATTACHMENT_FILENAME_REQUIRES_REVIEW",)
    monkeypatch.setattr(eis_document_store, "ANALYZER_VERSION", "eis-native-v2-xls")
    assert reanalyze_received_document(storage, company, ref, UID, actor_id="test")
    report = get_document_status(storage, company, ref)
    assert report.documents[0].filename == "Часть 1/2.pdf"
    assert report.documents[0].reasons == ("ATTACHMENT_FILENAME_REQUIRES_REVIEW",)
    assert report.extraction_status == "UNKNOWN"


@pytest.fixture(scope="module")
def storage() -> PostgresStorage:
    database_url = os.environ.get("DATABASE_URL")
    admin_url = os.environ.get("DATABASE_ADMIN_URL")
    if not database_url or not admin_url:
        pytest.skip("Separate PostgreSQL gate DB required")
    config = Config("alembic.ini")
    config.set_main_option("sqlalchemy.url", admin_url.replace("%", "%%"))
    command.upgrade(config, "head")
    engine = create_engine(database_url, pool_pre_ping=True)
    try:
        yield PostgresStorage(engine)
    finally:
        engine.dispose()


def docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Поставка в Москву. Оплата в течение 30 дней после приёмки.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def admit(
    storage: PostgresStorage,
    company: str,
    size: int,
    *,
    version: int = 1,
    listed: bool = True,
    filename: str = "spec.docx",
    deadline: str | None = None,
    number: str = NUMBER,
    competition: bool = False,
    price_weight: str = "60",
    qualification_weight: str = "40",
    ambiguous_attachment: bool = False,
) -> NoticeDocumentRef:
    attachments = (
        f"""<attachmentsInfo><attachmentInfo><publishedContentId>{UID}</publishedContentId>
<fileName>{filename}</fileName><fileSize>{size}</fileSize><url>{URL}</url>
</attachmentInfo></attachmentsInfo>"""
        if listed
        else ""
    )
    if ambiguous_attachment:
        entry = attachments.removeprefix("<attachmentsInfo>").removesuffix("</attachmentsInfo>")
        attachments = attachments.replace(
            "</attachmentsInfo>", entry.replace(filename, "other.pdf") + "</attachmentsInfo>"
        )
    xml = f"""<export xmlns="http://zakupki.gov.ru/oos/export/1"><epNotificationEF2020>
<versionNumber>{version}</versionNumber><commonInfo><purchaseNumber>{number}</purchaseNumber>
<purchaseObjectInfo>Поставка бумаги</purchaseObjectInfo></commonInfo>{attachments}
<notificationInfo><procedureInfo><collectingInfo>
{f"<endDT>{deadline}</endDT>" if deadline else ""}
</collectingInfo></procedureInfo></notificationInfo></epNotificationEF2020></export>"""
    if competition:
        xml = xml.replace("epNotificationEF2020", "epNotificationEOK2020")
        xml = xml.replace(
            "</notificationInfo>",
            '<p:criteriaInfo xmlns:p="http://zakupki.gov.ru/oos/EPtypes/1">'
            "<p:criterionInfo><p:costCriterionInfo><p:code>CP</p:code><p:valueInfo>"
            f"<p:value>{price_weight}</p:value></p:valueInfo></p:costCriterionInfo></p:criterionInfo>"
            "<p:criterionInfo><p:qualitativeCriterionInfo><p:code>QO</p:code><p:valueInfo>"
            f"<p:value>{qualification_weight}</p:value></p:valueInfo>"
            "</p:qualitativeCriterionInfo></p:criterionInfo></p:criteriaInfo></notificationInfo>",
        )
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr("notice.xml", xml)
    ingest_ready_window(
        storage,
        company,
        EisReadyWindow(
            f"getDocsIP/PRIZ/50/2026-09-25/{(int(number[-2:]) + version) % 24:02d}+03:00"
            + ("/epNotificationEOK2020" if competition else ""),
            "a" * 64,
            1,
            (output.getvalue(),),
        ),
        actor_id="test",
    )
    with storage.transaction(company) as connection:
        digest = connection.scalar(
            select(eis_notice_versions.c.xml_sha256).where(
                (eis_notice_versions.c.version_number == version)
                & (eis_notice_versions.c.purchase_number == number)
            )
        )
    return NoticeDocumentRef(
        purchase_number=number, version_number=version, notice_xml_sha256=digest
    )


def test_document_worker_rechecks_expiry_after_job_planning(storage: PostgresStorage) -> None:
    company = f"document-expiry-{uuid.uuid4()}"
    content = docx_bytes()
    ref = admit(storage, company, len(content), deadline="2026-09-30T12:00:00+03:00")
    plan_document_job(storage, company, ref, actor_id="test")
    calls: list[str] = []

    def download(url: str) -> bytes:
        calls.append(url)
        return content

    ready = run_document_step(
        storage,
        company,
        actor_id="probe",
        now=datetime(2026, 10, 1, tzinfo=UTC),
        downloader=download,
        min_free_bytes=1,
        probe_only=True,
    )
    assert ready["status"] == "IDLE"
    assert calls == []
    with storage.transaction(company) as connection:
        assert (
            connection.scalar(
                select(func.count())
                .select_from(audit_events)
                .where(audit_events.c.event_type == "EIS_DOCUMENT_READ_BLOCKED")
            )
            == 0
        )

    result = run_document_step(
        storage,
        company,
        actor_id="test",
        now=datetime(2026, 10, 1, tzinfo=UTC),
        downloader=download,
        min_free_bytes=1,
    )
    assert calls == [], "Expired jobs must be blocked before source I/O"
    assert result["attempted"] == 0
    assert get_document_status(storage, company, ref).received_count == 0


def test_blocked_document_job_does_not_starve_next_and_audit_replays(
    storage: PostgresStorage,
) -> None:
    company = f"document-fair-{uuid.uuid4()}"
    content = docx_bytes()
    expired = admit(storage, company, len(content), deadline="2026-09-30T12:00:00+03:00")
    plan_document_job(storage, company, expired, actor_id="test")
    valid = admit(storage, company, len(content), number=NUMBER[:-1] + "5")
    plan_document_job(storage, company, valid, actor_id="test")
    now = datetime(2026, 10, 1, tzinfo=UTC)
    result = run_document_step(
        storage,
        company,
        actor_id="test",
        now=now,
        downloader=lambda _: content,
        min_free_bytes=1,
    )
    assert result["status"] == "PROCESSED" and result["purchase_number"] == valid.purchase_number
    assert result["skipped_jobs"] == 1 and result["skip_reasons"] == {"DEADLINE_EXPIRED": 1}
    assert get_document_status(storage, company, expired).received_count == 0
    report = get_document_status(storage, company, expired, now=now)
    assert report.read_gate and report.read_gate.blocked_reasons == ("DEADLINE_EXPIRED",)
    # Repeated GET-equivalent status reads do not create decision audit events.
    get_document_status(storage, company, expired, now=now + timedelta(seconds=1))
    for offset in range(2):
        gate = document_read_gate(
            storage, company, expired, actor_id="test", now=now + timedelta(seconds=offset)
        )
        assert gate.blocked_reasons == ("DEADLINE_EXPIRED",)
    with storage.transaction(company) as connection:
        events = (
            connection.execute(
                select(audit_events).where(audit_events.c.event_type == "EIS_DOCUMENT_READ_BLOCKED")
            )
            .mappings()
            .all()
        )
    assert len(events) == 1
    assert events[0]["basis"]["notice"]["notice_xml_sha256"] == expired.notice_xml_sha256
    assert any("collectingInfo/endDT" in value for value in events[0]["basis"]["evidence"])
    with pytest.raises(LookupError, match="NOTICE_NOT_FOUND"):
        document_read_gate(storage, "other-document-gate", expired, actor_id="test", now=now)
    with storage.transaction("other-document-gate") as connection:
        assert connection.scalar(select(func.count()).select_from(audit_events)) == 0


def test_read_only_ready_probe_skips_blocked_and_deferred_without_mutations(
    storage: PostgresStorage,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from procurement.sources import eis_document_worker as worker

    company = f"document-probe-{uuid.uuid4()}"
    ref = admit(storage, company, 3, filename="instructions.txt")
    plan_document_job(storage, company, ref, actor_id="test")
    now = datetime.now(UTC)

    def forbidden(*args: object, **kwargs: object) -> None:
        raise AssertionError("Ready probe must not download, analyze or audit")

    for name in ("document_read_gate", "reanalyze_received_document", "record_document_outcome"):
        monkeypatch.setattr(worker, name, forbidden)

    def counts():
        with storage.transaction(company, consistent_read=True) as conn:
            return tuple(
                conn.scalar(select(func.count()).select_from(t))
                for t in (
                    audit_events,
                    eis_document_jobs,
                    eis_document_observations,
                    eis_document_blobs,
                    eis_document_analyses,
                )
            )

    def probe(tenant=company, clock=now):
        return run_document_step(
            storage,
            tenant,
            actor_id="probe",
            now=clock,
            downloader=forbidden,
            min_free_bytes=1,
            probe_only=True,
        )

    before = counts()
    assert probe()["status"] == "READY"
    assert probe("other-document-probe")["status"] == "IDLE"
    assert counts() == before
    record_document_outcome(
        storage, company, ref, UID, actor_id="test", failure_code="ATTACHMENT_UNAVAILABLE"
    )
    before = counts()
    assert probe()["status"] == "IDLE"
    assert probe(clock=now + timedelta(minutes=31))["status"] == "READY"
    assert counts() == before
    admit(storage, company, 3, filename="instructions.txt", version=2)
    before = counts()
    assert probe()["status"] == "IDLE"
    assert counts() == before


def test_read_only_probe_expiry_without_audit(
    storage: PostgresStorage,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from procurement.sources import eis_document_worker as worker

    company = f"document-probe-expired-{uuid.uuid4()}"
    ref = admit(
        storage, company, 3, filename="instructions.txt", deadline="2030-01-01T12:00:00+03:00"
    )
    plan_document_job(storage, company, ref, actor_id="test")

    def forbidden(*args: object, **kwargs: object) -> None:
        raise AssertionError("No mutation during probe")

    monkeypatch.setattr(worker, "document_read_gate", forbidden)
    report = run_document_step(
        storage,
        company,
        actor_id="probe",
        now=datetime(2031, 1, 1, tzinfo=UTC),
        downloader=forbidden,
        min_free_bytes=1,
        probe_only=True,
    )
    assert report["status"] == "IDLE" and report["skip_reasons"] == {"DEADLINE_EXPIRED": 1}
    with storage.transaction(company) as conn:
        assert (
            conn.scalar(
                select(func.count())
                .select_from(audit_events)
                .where(audit_events.c.event_type == "EIS_DOCUMENT_READ_BLOCKED")
            )
            == 0
        )


@pytest.mark.parametrize(
    "published,blocked",
    [
        ("2026-09-30T10:00:00+03:00", True),
        ("2099-01-01T10:00:00+03:00", False),
        ("", False),
    ],
)
def test_document_worker_rechecks_cancel_but_keeps_ambiguous_unknown(
    storage: PostgresStorage,
    published: str,
    blocked: bool,
) -> None:
    company = f"document-cancel-{uuid.uuid4()}"
    content = docx_bytes()
    ref = admit(storage, company, len(content))
    plan_document_job(storage, company, ref, actor_id="test")
    output = BytesIO()
    with ZipFile(output, "w") as archive:
        archive.writestr(
            "cancel.xml",
            f"""<export xmlns="http://zakupki.gov.ru/oos/export/1">
<epNotificationCancel><id>1</id><commonInfo><purchaseNumber>{NUMBER}</purchaseNumber>
<docPublishDTInEIS>{published}</docPublishDTInEIS></commonInfo><cancelReasonInfo/>
</epNotificationCancel></export>""",
        )
    ingest_ready_window(
        storage,
        company,
        EisReadyWindow(
            "getDocsIP/PRIZ/50/2026-09-30+03:00/epNotificationCancel",
            "b" * 64,
            1,
            (output.getvalue(),),
        ),
        actor_id="test",
    )
    calls: list[str] = []

    def download(url: str) -> bytes:
        calls.append(url)
        return content

    ready = run_document_step(
        storage,
        company,
        actor_id="probe",
        now=datetime(2026, 10, 1, tzinfo=UTC),
        downloader=download,
        min_free_bytes=1,
        probe_only=True,
    )
    assert ready["status"] == ("IDLE" if blocked else "READY")
    assert calls == []
    with storage.transaction(company) as connection:
        assert (
            connection.scalar(
                select(func.count())
                .select_from(audit_events)
                .where(audit_events.c.event_type == "EIS_DOCUMENT_READ_BLOCKED")
            )
            == 0
        )

    result = run_document_step(
        storage,
        company,
        actor_id="test",
        now=datetime(2026, 10, 1, tzinfo=UTC),
        downloader=download,
        min_free_bytes=1,
    )
    assert result["attempted"] == (0 if blocked else 1)
    assert len(calls) == (0 if blocked else 1)
    assert get_document_status(storage, company, ref).review_status == "UNKNOWN"
    if blocked:
        with storage.transaction(company) as connection:
            basis = connection.scalar(
                select(audit_events.c.basis).where(
                    audit_events.c.event_type == "EIS_DOCUMENT_READ_BLOCKED"
                )
            )
        assert basis["blocked_reasons"] == ["PROCEDURE_BLOCKED"]
        assert any(
            "epNotificationCancel" in value and "XML" in value for value in basis["evidence"]
        )


def test_stale_document_job_records_block_before_fetch(
    storage: PostgresStorage,
) -> None:
    company = f"document-stale-{uuid.uuid4()}"
    content = docx_bytes()
    ref = admit(storage, company, len(content))
    plan_document_job(storage, company, ref, actor_id="test")
    admit(storage, company, len(content), version=2)

    def forbidden(_: str) -> bytes:
        raise AssertionError("stale version must not fetch")

    result = run_document_step(
        storage, company, actor_id="test", downloader=forbidden, min_free_bytes=1
    )
    assert result["attempted"] == 0 and result["skip_reasons"] == {"NOTICE_VERSION_STALE": 1}
    gate = document_read_gate(storage, company, ref, actor_id="test", now=datetime.now(UTC))
    assert gate.blocked_reasons == ("NOTICE_VERSION_STALE",)


@pytest.mark.parametrize("new_version", ["eis-native-v2-xls", "eis-native-v3-doc"])
def test_versioned_cached_analysis_reuses_bytes_and_rejects_stale_or_wrong_tenant(
    storage: PostgresStorage,
    monkeypatch: pytest.MonkeyPatch,
    new_version: str,
) -> None:
    from procurement.storage import eis_document_store

    company = f"document-reanalysis-{uuid.uuid4()}"
    content = docx_bytes()
    ref = admit(storage, company, len(content))
    plan_document_job(storage, company, ref, actor_id="test")
    monkeypatch.setattr(eis_document_store, "ANALYZER_VERSION", "eis-native-v1")
    record_document_outcome(storage, company, ref, UID, actor_id="test", content=content)
    monkeypatch.setattr(eis_document_store, "ANALYZER_VERSION", new_version)
    report = get_document_status(storage, company, ref)
    assert report.documents[0].reasons == ("ANALYSIS_NOT_AVAILABLE",)
    assert report.analyzer_version == new_version
    assert report.analysis_pending_count == 1

    from procurement.sources import eis_document_worker as worker

    with storage.transaction(company) as conn:
        before = tuple(
            conn.scalar(select(func.count()).select_from(t))
            for t in (
                audit_events,
                eis_document_analyses,
                eis_document_observations,
            )
        )

    def no_probe_analysis(*args: object, **kwargs: object) -> None:
        raise AssertionError("Ready probe must not parse native bytes")

    with monkeypatch.context() as probe_patch:
        probe_patch.setattr(worker, "reanalyze_received_document", no_probe_analysis)
        ready = run_document_step(
            storage, company, actor_id="probe", min_free_bytes=1, probe_only=True
        )
        assert ready["status"] == "READY" and ready["operation"] == "ANALYSIS"
    with storage.transaction(company) as conn:
        assert before == tuple(
            conn.scalar(select(func.count()).select_from(t))
            for t in (
                audit_events,
                eis_document_analyses,
                eis_document_observations,
            )
        )

    def forbidden(_: str) -> bytes:
        raise AssertionError("received bytes must not be downloaded again")

    result = run_document_step(
        storage, company, actor_id="test", downloader=forbidden, min_free_bytes=1
    )
    assert result["status"] == "REANALYZED" and result["attempted"] == 0
    assert get_document_status(storage, company, ref).extraction_status == "EXTRACTED"
    assert get_document_status(storage, company, ref).analysis_pending_count == 0

    def no_extraction(*_: object, **__: object) -> None:
        raise AssertionError("parser-version replay must use its immutable cache")

    monkeypatch.setattr(eis_document_store, "analyze_document", no_extraction)
    assert not reanalyze_received_document(storage, company, ref, UID, actor_id="test")
    assert (
        run_document_step(
            storage, company, actor_id="test", downloader=forbidden, min_free_bytes=1
        )["status"]
        == "IDLE"
    )
    with storage.transaction(company) as connection:
        assert connection.scalar(select(func.count()).select_from(eis_document_blobs)) == 1
        assert connection.scalar(select(func.count()).select_from(eis_document_observations)) == 1
        assert connection.scalar(select(func.count()).select_from(eis_document_analyses)) == 2
        assert (
            connection.scalar(
                select(func.count())
                .select_from(audit_events)
                .where(audit_events.c.event_type == "EIS_DOCUMENT_REANALYZED")
            )
            == 1
        )
    with pytest.raises(LookupError, match="NOTICE_NOT_FOUND"):
        reanalyze_received_document(storage, "other-reanalysis-tenant", ref, UID, actor_id="test")
    admit(storage, company, len(content), version=2)
    with pytest.raises(ValueError, match="NOTICE_VERSION_STALE"):
        reanalyze_received_document(storage, company, ref, UID, actor_id="test")


def test_cached_reanalysis_verifies_byte_digest_before_parser(
    storage: PostgresStorage,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from sqlalchemy import insert

    from procurement.storage import eis_document_store

    company = f"document-corrupt-cache-{uuid.uuid4()}"
    content = docx_bytes()
    ref = admit(storage, company, len(content))
    plan_document_job(storage, company, ref, actor_id="test")
    # Simulate corrupted internal evidence using only a synthetic tenant.
    with storage.transaction(company) as connection:
        connection.execute(
            insert(eis_document_blobs).values(
                company_id=company,
                sha256="f" * 64,
                content=content,
                byte_count=len(content),
                received_at=datetime.now(UTC),
            )
        )
        connection.execute(
            insert(eis_document_observations).values(
                company_id=company,
                purchase_number=ref.purchase_number,
                version_number=ref.version_number,
                observation_id=str(uuid.uuid4()),
                content_uid=UID,
                sha256="f" * 64,
                fetch_status="RECEIVED",
                reason_code=None,
                observed_at=datetime.now(UTC),
            )
        )

    def forbidden(*_: object, **__: object) -> None:
        raise AssertionError("invalid digest must be rejected before parser")

    monkeypatch.setattr(eis_document_store, "analyze_document", forbidden)
    with pytest.raises(ValueError, match="DOCUMENT_BYTES_NOT_VERIFIABLE"):
        reanalyze_received_document(storage, company, ref, UID, actor_id="test")


def test_document_queue_replay_cache_stale_and_unknown_review(
    storage: PostgresStorage,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from procurement.storage import eis_document_store

    company = f"documents-{uuid.uuid4()}"
    content = docx_bytes()
    ref = admit(storage, company, len(content))
    assert get_document_status(storage, company, ref).download_status == "NOT_REQUESTED"
    assert plan_document_job(storage, company, ref, actor_id="test").download_status == "QUEUED"
    plan_document_job(storage, company, ref, actor_id="test")
    with storage.transaction(company) as connection:
        assert connection.scalar(select(func.count()).select_from(eis_document_jobs)) == 1

    calls = []

    def download(url: str) -> bytes:
        calls.append(url)
        return content

    result = run_document_step(
        storage, company, actor_id="test", downloader=download, min_free_bytes=1
    )
    assert result["status"] == "PROCESSED" and calls == [URL]
    analysis = get_document_analysis(storage, company, ref, UID)
    assert analysis.status == "PASS" and len(analysis.snippets) >= 2

    def forbidden(*args: object, **kwargs: object) -> None:
        raise AssertionError("Cached bytes must not be re-extracted")

    monkeypatch.setattr(eis_document_store, "analyze_document", forbidden)
    assert not record_document_outcome(storage, company, ref, UID, actor_id="test", content=content)
    report = get_document_status(storage, company, ref)
    assert (
        report.download_status == "COMPLETE_LISTED_SET" and report.extraction_status == "EXTRACTED"
    )
    assert report.review_status == report.source_manifest_coverage == "UNKNOWN"
    assert (
        run_document_step(storage, company, actor_id="test", downloader=download, min_free_bytes=1)[
            "status"
        ]
        == "IDLE"
    )
    assert calls == [URL]
    new = admit(storage, company, len(content), version=2)
    assert get_document_status(storage, company, ref).version_status == "STALE"
    with pytest.raises(ValueError, match="NOTICE_VERSION_STALE"):
        plan_document_job(storage, company, ref, actor_id="test")
    assert plan_document_job(storage, company, new, actor_id="test").received_count == 0
    # Fresh bytes are still fetched for the new XML, then their exact analysis is reused.
    run_document_step(storage, company, actor_id="test", downloader=download, min_free_bytes=1)
    assert calls == [URL, URL]
    assert get_document_analysis(storage, company, new, UID) == analysis
    with storage.transaction(company) as connection:
        assert connection.scalar(select(func.count()).select_from(eis_document_blobs)) == 1
        assert connection.scalar(select(func.count()).select_from(eis_document_analyses)) == 1


def test_document_failure_retry_size_mismatch_and_budget_precede_network(
    storage: PostgresStorage,
) -> None:
    company = f"documents-{uuid.uuid4()}"
    ref = admit(storage, company, 3, filename="instructions.txt")
    plan_document_job(storage, company, ref, actor_id="test")
    calls = 0

    def failed(url: str) -> bytes:
        nonlocal calls
        calls += 1
        raise EisPublicError("ATTACHMENT_UNAVAILABLE")

    result = run_document_step(
        storage, company, actor_id="test", downloader=failed, max_database_bytes=1, min_free_bytes=1
    )
    assert result["attempted"] == calls == 0
    result = run_document_step(
        storage, company, actor_id="test", downloader=failed, min_free_bytes=1
    )
    assert result["status"] == "BLOCKED" and calls == 1
    report = get_document_status(storage, company, ref)
    assert report.documents[0].fetch_status == "FAILED" and report.received_count == 0
    assert (
        run_document_step(storage, company, actor_id="test", downloader=failed, min_free_bytes=1)[
            "status"
        ]
        == "IDLE"
    )
    record_document_outcome(storage, company, ref, UID, actor_id="test", content=b"wrong")
    report = get_document_status(storage, company, ref)
    assert report.documents[0].fetch_status == "SIZE_MISMATCH"
    assert report.extraction_status == report.review_status == "UNKNOWN"
    run_document_step(
        storage,
        company,
        actor_id="test",
        now=datetime.now(UTC) + timedelta(minutes=31),
        downloader=lambda url: b"abc",
        min_free_bytes=1,
    )
    report = get_document_status(storage, company, ref)
    assert report.download_status == "COMPLETE_LISTED_SET"
    assert report.documents[0].extraction_status == "UNKNOWN"  # unsupported txt, never PASS.
    with pytest.raises(ValueError, match="DOCUMENT_CONTENT_CONFLICT"):
        record_document_outcome(storage, company, ref, UID, actor_id="test", content=b"bad")


def test_document_evidence_is_tenant_scoped_and_immutable(storage: PostgresStorage) -> None:
    company = f"documents-{uuid.uuid4()}"
    other = f"other-{uuid.uuid4()}"
    ref = admit(storage, company, 3, filename="instructions.txt")
    plan_document_job(storage, company, ref, actor_id="test")
    record_document_outcome(storage, company, ref, UID, actor_id="test", content=b"abc")
    with pytest.raises(LookupError, match="NOTICE_NOT_FOUND"):
        get_document_status(storage, other, ref)
    for table in (
        eis_document_jobs,
        eis_document_blobs,
        eis_document_analyses,
        eis_document_observations,
    ):
        with storage.transaction(other) as connection:
            assert connection.scalar(select(func.count()).select_from(table)) == 0
        with pytest.raises(DBAPIError), storage.transaction(company) as connection:
            connection.execute(update(table).values(company_id=company))
    with pytest.raises(ValueError, match="DOCUMENT_NOT_IN_MANIFEST"):
        record_document_outcome(storage, company, ref, "B" * 32, actor_id="test", content=b"abc")


def test_empty_manifest_and_resource_limit_are_not_complete(storage: PostgresStorage) -> None:
    company = f"documents-{uuid.uuid4()}"
    ref = admit(storage, company, 0, listed=False)
    report = plan_document_job(storage, company, ref, actor_id="test")
    assert report.download_status == "UNKNOWN_MANIFEST" and report.extraction_status == "UNKNOWN"
    new = admit(storage, company, 20 * 1024**2 + 1, version=2)
    with pytest.raises(ValueError, match="DOCUMENT_SET_RESOURCE_LIMIT"):
        plan_document_job(storage, company, new, actor_id="test")


def test_top_documents_filters_before_limit_and_never_selects_expired(
    storage: PostgresStorage,
) -> None:
    company = f"selection-{uuid.uuid4()}"
    output = BytesIO()
    eligible_number = str(9000000000000000000 + 105)
    with ZipFile(output, "w") as archive:
        for index in range(107):
            number = str(9000000000000000000 + index)
            attachment = (
                (
                    f"<attachmentsInfo><attachmentInfo><publishedContentId>{UID}</publishedContentId>"
                    f"<fileName>spec.docx</fileName><fileSize>3</fileSize><url>{URL}</url>"
                    "</attachmentInfo></attachmentsInfo>"
                )
                if index >= 105
                else ""
            )
            deadline = "2000-01-01T12:00:00+03:00" if index == 106 else "2099-01-01T12:00:00+03:00"
            xml = f"""<export xmlns="http://zakupki.gov.ru/oos/export/1"><epNotificationEF2020>
<versionNumber>1</versionNumber><commonInfo><purchaseNumber>{number}</purchaseNumber>
<purchaseObjectInfo>Бумага</purchaseObjectInfo></commonInfo>{attachment}
<notificationInfo><procedureInfo><collectingInfo><endDT>{deadline}</endDT></collectingInfo>
</procedureInfo></notificationInfo></epNotificationEF2020></export>"""
            archive.writestr(f"{number}.xml", xml)
    ingest_ready_window(
        storage,
        company,
        EisReadyWindow("getDocsIP/PRIZ/50/2026-09-25+03:00", "a" * 64, 1, (output.getvalue(),)),
        actor_id="test",
    )
    request = DocumentSelectionRequest(query="бумага", max_candidates=3)
    result = plan_top_document_jobs(storage, company, request, actor_id="test")
    assert result.matched_count == 107 and result.eligible_count == 1
    assert [job.notice.purchase_number for job in result.jobs] == [eligible_number]
    assert result.jobs[0].review_status == "UNKNOWN" and not result.rejected
    replay = plan_top_document_jobs(storage, company, request, actor_id="test")
    assert [job.model_dump(exclude={"read_gate": {"checked_at"}}) for job in result.jobs] == [
        job.model_dump(exclude={"read_gate": {"checked_at"}}) for job in replay.jobs
    ]
    with storage.transaction(company) as connection:
        assert connection.scalar(select(func.count()).select_from(eis_document_jobs)) == 1


def test_exact_evidence_bridge_links_cached_bytes_without_review_or_cross_tenant(
    storage: PostgresStorage,
) -> None:
    company = f"tenant-eis-evidence-bridge-{uuid.uuid4()}"
    content = docx_bytes()
    ref = admit(storage, company, len(content), filename="Техническое задание.docx")
    plan_document_job(storage, company, ref, actor_id="test")
    record_document_outcome(storage, company, ref, UID, actor_id="test", content=content)
    bridge = get_eis_evidence_bridge(storage, company, ref, now=datetime(2026, 9, 29, tzinfo=UTC))
    assert bridge.notice == ref
    assert bridge.collection.received_count == 1
    assert bridge.documents[0].role_hint == "TECHNICAL_SPECIFICATION"
    assert bridge.documents[0].role_status == "UNKNOWN"
    assert bridge.documents[0].conditions
    assert (
        bridge.documents[0].conditions[0].provenance.document_id
        == bridge.collection.documents[0].sha256
    )
    assert bridge.requirement_status == bridge.completeness_status == "UNKNOWN"
    assert bridge.economic_status == "UNKNOWN"
    assert get_eis_evidence_bridge(storage, company, ref).documents == bridge.documents
    with pytest.raises(LookupError, match="NOTICE_NOT_FOUND"):
        get_eis_evidence_bridge(storage, "tenant-eis-evidence-other", ref)
    admit(storage, company, len(content), version=2, filename="Техническое задание.docx")
    stale = get_eis_evidence_bridge(storage, company, ref)
    assert stale.collection.version_status == "STALE"
    assert stale.requirement_status == "UNKNOWN"


def test_evidence_context_reads_cached_blocks_without_changing_analysis_or_audit(
    storage: PostgresStorage,
) -> None:
    company = f"context-evidence-{uuid.uuid4()}"
    output = BytesIO()
    document = Document()
    document.add_paragraph("2. Место поставки:")
    document.add_paragraph("Синтетический город, тестовый адрес")
    document.add_paragraph("4. Срок поставки:")
    document.add_paragraph("В течение 10 рабочих дней после заявки")
    document.save(output)
    content = output.getvalue()
    ref = admit(storage, company, len(content), filename="Описание объекта закупки.docx")
    plan_document_job(storage, company, ref, actor_id="test")
    record_document_outcome(storage, company, ref, UID, actor_id="test", content=content)
    before = get_document_analysis(storage, company, ref, UID)
    with storage.transaction(company) as connection:
        audit_before = connection.scalar(select(func.count()).select_from(audit_events))
    for _ in range(2):
        bridge = get_eis_evidence_bridge(storage, company, ref)
        conditions = bridge.documents[0].conditions
        assert [item.context[0].provenance.locator for item in conditions] == [
            "paragraph 2",
            "paragraph 4",
        ]
        assert bridge.documents[0].conditions_total == 2
        assert not bridge.documents[0].conditions_truncated
        assert bridge.requirement_status == bridge.completeness_status == "UNKNOWN"
    assert get_document_analysis(storage, company, ref, UID) == before
    with storage.transaction(company) as connection:
        assert connection.scalar(select(func.count()).select_from(audit_events)) == audit_before
    with pytest.raises(LookupError, match="NOTICE_NOT_FOUND"):
        get_eis_evidence_bridge(storage, "other-context-tenant", ref)


def test_comparison_uses_saved_exact_tenant_refs_and_preserves_jobs_audits(
    storage: PostgresStorage,
) -> None:
    from procurement.storage.eis_comparison import EisComparisonRequest, get_eis_comparison

    company = f"compare-evidence-{uuid.uuid4()}"
    content = docx_bytes()
    first = admit(storage, company, len(content), deadline="2099-01-01T12:00:00+03:00")
    second = admit(
        storage,
        company,
        len(content),
        number="0848300054826000156",
        deadline="2000-01-01T12:00:00+03:00",
    )
    plan_document_job(storage, company, first, actor_id="test")
    record_document_outcome(storage, company, first, UID, actor_id="test", content=content)
    request = EisComparisonRequest(notices=(first, second))

    def counts():
        with storage.transaction(company) as connection:
            return tuple(
                connection.scalar(select(func.count()).select_from(table))
                for table in (
                    eis_document_jobs,
                    eis_document_blobs,
                    eis_document_analyses,
                    eis_document_observations,
                    audit_events,
                )
            )

    before = counts()
    result = get_eis_comparison(storage, company, request)
    assert result.rows[0].readiness_status == "NEEDS_ANALYSIS"
    assert result.rows[1].readiness_status == "BLOCKED"
    assert result.rows[1].blocking_reasons == ("DEADLINE_EXPIRED",)
    assert result.rows[0].evidence.collection.received_count == 1
    assert not result.rows[1].evidence.collection.requested
    assert result.profit_order_status == "UNKNOWN"
    assert all(row.evidence.checked_at == result.checked_at for row in result.rows)
    assert counts() == before
    replay = get_eis_comparison(storage, company, request)
    assert [
        (row.evidence.notice, row.blocking_reasons, row.evidence.documents) for row in replay.rows
    ] == [
        (row.evidence.notice, row.blocking_reasons, row.evidence.documents) for row in result.rows
    ]
    assert counts() == before
    with pytest.raises(LookupError, match="NOTICE_NOT_FOUND"):
        get_eis_comparison(storage, "other-compare-tenant", request)
    admit(storage, company, len(content), version=2, deadline="2099-01-01T12:00:00+03:00")
    stale = get_eis_comparison(storage, company, request)
    assert stale.rows[0].blocking_reasons == ("NOTICE_VERSION_STALE",)


def test_competition_criteria_read_exact_archive_tenant_versions_without_mutation(storage):
    from decimal import Decimal

    from procurement.storage.eis_comparison import EisComparisonRequest, get_eis_comparison

    company = f"criteria-evidence-{uuid.uuid4()}"
    first = admit(
        storage, company, 0, listed=False, competition=True, deadline="2099-01-01T12:00:00+03:00"
    )
    second = admit(
        storage,
        company,
        0,
        number="0848300054826000156",
        listed=False,
        deadline="2099-01-01T12:00:00+03:00",
    )
    with storage.transaction(company) as conn:
        before = [
            conn.scalar(select(func.count()).select_from(t))
            for t in (
                audit_events,
                eis_document_jobs,
                eis_document_observations,
                eis_document_analyses,
            )
        ]
    request = EisComparisonRequest(notices=(first, second))
    result = get_eis_comparison(storage, company, request)
    criteria = result.rows[0].evidence.competition_criteria
    assert criteria.weight_integrity == "CONSISTENT"
    assert [c.weights[0].value_percent for c in criteria.criteria] == [Decimal(60), Decimal(40)]
    assert all(c.provenance.document_id == first.notice_xml_sha256 for c in criteria.criteria)
    assert result.rows[1].evidence.competition_criteria.source_status == "NOT_APPLICABLE"
    assert result.profit_order_status == criteria.review_status == "UNKNOWN"
    assert (
        get_eis_comparison(storage, company, request).rows[0].evidence.competition_criteria
        == criteria
    )
    with pytest.raises(LookupError):
        get_eis_comparison(storage, "other-criteria-tenant", request)
    with storage.transaction(company) as conn:
        assert before == [
            conn.scalar(select(func.count()).select_from(t))
            for t in (
                audit_events,
                eis_document_jobs,
                eis_document_observations,
                eis_document_analyses,
            )
        ]
    current = admit(
        storage,
        company,
        0,
        listed=False,
        competition=True,
        version=2,
        price_weight="50",
        qualification_weight="50",
        deadline="2099-01-01T12:00:00+03:00",
    )
    stale = get_eis_evidence_bridge(storage, company, first)
    assert stale.collection.version_status == "STALE"
    assert stale.competition_criteria == criteria
    assert get_eis_evidence_bridge(storage, company, current).competition_criteria.criteria[
        0
    ].weights[0].value_percent == Decimal(50)


def test_ambiguous_reference_blocks_intake_analysis_and_probe(storage):
    company = f"ambiguous-reference-{uuid.uuid4()}"
    content = docx_bytes()
    ref = admit(storage, company, len(content), ambiguous_attachment=True)
    plan_document_job(storage, company, ref, actor_id="test")
    report = get_document_status(storage, company, ref)
    assert report.documents[0].reference_status == "UNKNOWN"
    assert len(report.documents[0].aliases) == 1
    assert report.documents[0].reasons == ("ATTACHMENT_REFERENCE_REQUIRES_REVIEW",)
    with pytest.raises(ValueError, match="ATTACHMENT_REFERENCE_REQUIRES_REVIEW"):
        record_document_outcome(storage, company, ref, UID, actor_id="test", content=content)

    def forbidden(*args):
        raise AssertionError("ambiguous source must not be downloaded")

    for probe in (True, False):
        result = run_document_step(
            storage,
            company,
            actor_id="test",
            downloader=forbidden,
            min_free_bytes=1,
            probe_only=probe,
        )
        assert result == {
            "status": "BLOCKED",
            "reason": "ATTACHMENT_REFERENCE_REQUIRES_REVIEW",
            "attempted": 0,
        }
    with storage.transaction(company) as conn:
        assert conn.scalar(select(func.count()).select_from(eis_document_blobs)) == 0
        assert conn.scalar(select(func.count()).select_from(eis_document_analyses)) == 0
    assert get_document_status(storage, company, ref).extraction_status == "UNKNOWN"
