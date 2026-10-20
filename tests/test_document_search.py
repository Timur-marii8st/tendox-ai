"""Hybrid discovery search keeps native and vision evidence visibly separate."""

from __future__ import annotations

import base64
import hashlib
import json
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from procurement.cli import main
from procurement.documents.document_search import (
    VISION_SEARCH_TAG_PROMPT_V1,
    DocumentSearchEntry,
    SearchSourceKind,
    VisionSearchTag,
    VisionSearchTagManifest,
    VisionTagReviewStatus,
    build_document_search_entry,
    search_documents,
)
from procurement.documents.extraction import (
    DocumentInput,
    ExtractionReason,
    ExtractionStatus,
    extract_document,
)

_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _docx_bytes(text: str, *, image: bool = False) -> bytes:
    document = Document()
    document.add_paragraph(text)
    if image:
        document.add_picture(BytesIO(_TINY_PNG))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _document(*, company_id: str = "tenant-a") -> tuple[DocumentInput, str]:
    content = _docx_bytes("Выполнение работ по ремонту пятого этажа", image=True)
    document = DocumentInput(
        document_id="document-42",
        filename="ТЗ_ремонт_кондиционеров_5_го_этажа.docx",
        declared_media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        content=content,
    )
    return document, company_id


def _vision_manifest(
    document: DocumentInput,
    *,
    company_id: str = "tenant-a",
    review_status: VisionTagReviewStatus = VisionTagReviewStatus.AI_SUGGESTED,
) -> VisionSearchTagManifest:
    return VisionSearchTagManifest(
        company_id=company_id,
        document_id=document.document_id,
        source_document_sha256=hashlib.sha256(document.content).hexdigest(),
        media_index=0,
        source_media="word/media/image1.png",
        source_media_sha256=hashlib.sha256(_TINY_PNG).hexdigest(),
        producer="vision-review-worker",
        model="approved-vision-model",
        prompt_sha256=hashlib.sha256(VISION_SEARCH_TAG_PROMPT_V1.encode()).hexdigest(),
        source_artifact_sha256="b" * 64,
        produced_at=datetime(2026, 9, 15, 12, tzinfo=UTC),
        tags=(
            VisionSearchTag(
                text="General Climate",
                confidence=0.97,
                locator='DOCX media "word/media/image1.png" :: title block',
                review_status=review_status,
                reviewed_by=(
                    "operator-7" if review_status is VisionTagReviewStatus.HUMAN_CONFIRMED else None
                ),
            ),
        ),
    )


def _entry(*, vision: bool = False, confirmed: bool = False) -> DocumentSearchEntry:
    document, company_id = _document()
    extraction = extract_document(document)
    manifest = (
        _vision_manifest(
            document,
            review_status=(
                VisionTagReviewStatus.HUMAN_CONFIRMED
                if confirmed
                else VisionTagReviewStatus.AI_SUGGESTED
            ),
        )
        if vision
        else None
    )
    return build_document_search_entry(
        document,
        extraction,
        company_id=company_id,
        title="Ремонт систем кондиционирования",
        vision_manifest=manifest,
    )


def test_unknown_document_still_exposes_safe_native_text_for_discovery() -> None:
    entry = _entry()

    assert entry.extraction_status is ExtractionStatus.UNKNOWN
    assert entry.extraction_reasons == (ExtractionReason.EMBEDDED_MEDIA_IGNORED.value,)
    result = search_documents((entry,), company_id="tenant-a", query="выполнение работ")[0]

    assert result.status is ExtractionStatus.PASS
    assert [hit.source_kind for hit in result.hits] == [SearchSourceKind.NATIVE_TEXT]
    assert result.hits[0].provenance.locator == "paragraph 1"


def test_russian_query_matches_related_filename_and_title_forms() -> None:
    result = search_documents((_entry(),), company_id="tenant-a", query="кондиционирование")[0]

    assert result.status is ExtractionStatus.PASS
    assert {hit.source_kind for hit in result.hits} == {
        SearchSourceKind.TITLE,
        SearchSourceKind.FILENAME,
    }


def test_ai_vision_only_match_remains_unknown() -> None:
    result = search_documents((_entry(vision=True),), company_id="tenant-a", query="General")[0]

    assert result.status is ExtractionStatus.UNKNOWN
    assert len(result.hits) == 1
    assert result.hits[0].source_kind is SearchSourceKind.VISION_TAG
    assert result.hits[0].status is ExtractionStatus.UNKNOWN
    assert result.hits[0].provenance.extractor.value == "llm"
    assert result.hits[0].vision_evidence is not None
    assert result.hits[0].vision_evidence.model == "approved-vision-model"
    assert (
        result.hits[0].vision_evidence.source_media_sha256 == hashlib.sha256(_TINY_PNG).hexdigest()
    )


def test_human_confirmed_vision_match_is_pass_for_discovery_only() -> None:
    result = search_documents(
        (_entry(vision=True, confirmed=True),),
        company_id="tenant-a",
        query="General Climate",
    )[0]

    assert result.status is ExtractionStatus.PASS
    assert result.hits[0].vision_review_status is VisionTagReviewStatus.HUMAN_CONFIRMED
    assert result.hits[0].vision_evidence is not None
    assert result.hits[0].vision_evidence.reviewed_by == "operator-7"


def test_search_is_tenant_scoped() -> None:
    entry = _entry()

    assert search_documents((entry,), company_id="tenant-b", query="ремонт") == ()


def test_vision_manifest_must_match_tenant_and_document_digest() -> None:
    document, _ = _document()
    extraction = extract_document(document)
    manifest = _vision_manifest(document, company_id="tenant-b")

    with pytest.raises(ValueError, match="does not match"):
        build_document_search_entry(
            document,
            extraction,
            company_id="tenant-a",
            vision_manifest=manifest,
        )


def test_vision_manifest_must_match_a_linked_media_digest() -> None:
    document, _ = _document()
    extraction = extract_document(document)
    payload = _vision_manifest(document).model_dump(mode="json")
    payload["source_media_sha256"] = "0" * 64
    manifest = VisionSearchTagManifest.model_validate(payload)

    with pytest.raises(ValueError, match="linked DOCX media"):
        build_document_search_entry(
            document,
            extraction,
            company_id="tenant-a",
            vision_manifest=manifest,
        )


def test_vision_manifest_rejects_an_unversioned_prompt() -> None:
    document, _ = _document()
    payload = _vision_manifest(document).model_dump(mode="json")
    payload["prompt_sha256"] = "0" * 64

    with pytest.raises(ValueError, match="prompt does not match"):
        VisionSearchTagManifest.model_validate(payload)


def test_cli_searches_a_local_docx_without_vision_upload(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    document_path = tmp_path / "ТЗ_ремонт_кондиционеров.docx"
    document_path.write_bytes(_docx_bytes("Выполнение работ"))

    exit_code = main(
        [
            "search-document",
            str(document_path),
            "--company-id",
            "tenant-a",
            "--query",
            "кондиционирование",
        ]
    )

    assert exit_code == 0
    output = json.loads(capsys.readouterr().out)
    assert output["search_status"] == "PASS"
    assert output["document_id"].startswith("doc-")
    assert [hit["source_kind"] for hit in output["hits"]] == ["FILENAME"]


def test_cli_returns_fail_without_turning_no_match_into_an_error(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    document_path = tmp_path / "specification.docx"
    document_path.write_bytes(_docx_bytes("Выполнение работ"))

    exit_code = main(
        [
            "search-document",
            str(document_path),
            "--company-id",
            "tenant-a",
            "--query",
            "General Climate",
        ]
    )

    assert exit_code == 4
    output = json.loads(capsys.readouterr().out)
    assert output["search_status"] == "FAIL"
    assert output["hits"] == []
