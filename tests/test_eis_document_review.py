"""Bounded DOCX review retains native facts without claiming requirement completeness."""

from __future__ import annotations

import hashlib
from io import BytesIO

import pytest
from docx import Document
from pydantic import ValidationError

from procurement.documents.extraction import (
    DocumentRejected,
    DocumentRejectionCode,
    ExtractedBlock,
    ExtractionReason,
    ExtractionResult,
    ExtractionStatus,
)
from procurement.domain.provenance import Extractor, Provenance
from procurement.sources import eis_document_review as review


def docx_bytes(*paragraphs: str) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def docx_with_header(body: str, header: str) -> bytes:
    document = Document()
    document.sections[0].header.paragraphs[0].text = header
    document.add_paragraph(body)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def extraction(blocks: tuple[ExtractedBlock, ...], *, reasons=()) -> ExtractionResult:
    return ExtractionResult(
        status=ExtractionStatus.UNKNOWN if reasons else ExtractionStatus.PASS,
        blocks=blocks,
        reasons=reasons,
        source_units_total=len(blocks),
        source_units_extracted=len(blocks),
        characters_extracted=sum(len(block.text) for block in blocks),
    )


def blocks(digest: str, texts: tuple[str, ...]) -> tuple[ExtractedBlock, ...]:
    return tuple(
        ExtractedBlock(
            text=text,
            provenance=Provenance(
                document_id=digest,
                locator=f"word/document.xml :: paragraph[{index}]",
                extractor=Extractor.NATIVE_DOCX,
            ),
        )
        for index, text in enumerate(texts, 1)
    )


def test_native_docx_analysis_and_page_preserve_digest_provenance_and_delivery_context():
    content = docx_bytes("Введение", "Срок поставки 10 дней", "Адрес склада", "Иное")
    digest = hashlib.sha256(content).hexdigest()
    analysis = review.analyze_review_document(content, "ТЗ.docx")
    assert analysis.review_version == "eis-native-review-v1"
    assert analysis.sha256 == digest and analysis.byte_count == len(content)
    assert analysis.extraction_status == "PASS"
    assert (
        analysis.extraction_input_byte_limit,
        analysis.extraction_expanded_byte_limit,
        analysis.extraction_source_unit_limit,
        analysis.extraction_character_limit,
    ) == (4 * 1024 * 1024, 32 * 1024 * 1024, 20_000, 1_000_000)
    assert analysis.source_units_total == analysis.source_units_extracted == 4
    assert analysis.blocks_extracted_total == analysis.retained_blocks == 4
    assert analysis.characters_extracted == analysis.retained_characters
    assert analysis.text_scope == "SUPPORTED_DOCX_BODY_PARAGRAPHS_AND_TABLES"
    assert analysis.retained_extracted_text_complete
    assert analysis.native_text_completeness == "UNKNOWN"
    assert all(block.provenance.document_id == digest for block in analysis.blocks)
    assert (
        analysis.document_role
        == analysis.position_association
        == analysis.approved_requirements
        == analysis.full_tspec_completeness
        == analysis.participation
        == "UNKNOWN"
    )

    page = review.review_document_page(analysis, 0, 3)
    assert [block.block_index for block in page.blocks] == [0, 1, 2]
    assert page.delivery_lead_indices == (0, 1, 2)
    assert [page.blocks[index].text for index in page.delivery_lead_indices] == [
        "Введение",
        "Срок поставки 10 дней",
        "Адрес склада",
    ]
    assert page.next_offset == 3 and page.structured_delivery == "UNKNOWN"
    assert page.sha256 == digest and page.filename == "ТЗ.docx"
    assert len(page.model_dump_json().encode()) <= review.MAX_PAGE_RESPONSE_BYTES


def test_docx_header_is_outside_declared_text_scope_and_never_complete_native_text():
    content = docx_with_header("Основной текст", "Срок доставки в заголовке")
    analysis = review.analyze_review_document(content, "spec.docx")
    assert [block.text for block in analysis.blocks] == ["Основной текст"]
    assert analysis.text_scope == "SUPPORTED_DOCX_BODY_PARAGRAPHS_AND_TABLES"
    assert analysis.retained_extracted_text_complete
    assert analysis.native_text_completeness == "UNKNOWN"


def test_retention_counts_native_totals_before_whole_block_prefix(monkeypatch):
    content = b"synthetic-docx"
    digest = hashlib.sha256(content).hexdigest()
    native_blocks = blocks(digest, tuple("x" for _ in range(1_001)))
    monkeypatch.setattr(
        review, "extract_document", lambda *_args, **_kwargs: extraction(native_blocks)
    )
    analysis = review.analyze_review_document(content, "spec.docx")
    assert analysis.blocks_extracted_total == analysis.source_units_extracted == 1_001
    assert analysis.retained_blocks == 1_000 and analysis.retained_characters == 1_000
    assert analysis.retention_truncated
    assert analysis.retention_reasons == ("RETAINED_BLOCK_LIMIT",)
    assert analysis.native_text_completeness == "UNKNOWN"
    assert not analysis.retained_extracted_text_complete

    native_blocks = blocks(digest, ("x" * 100_001, "later"))
    monkeypatch.setattr(
        review, "extract_document", lambda *_args, **_kwargs: extraction(native_blocks)
    )
    analysis = review.analyze_review_document(content, "spec.docx")
    assert not analysis.blocks and analysis.retained_characters == 0
    assert analysis.blocks_extracted_total == 2
    assert analysis.retention_reasons == ("RETAINED_CHARACTER_LIMIT",)
    payload = analysis.model_dump(mode="python")
    payload["characters_extracted"] = analysis.retained_characters
    with pytest.raises(ValidationError, match="RETENTION_CHARACTER_COUNT_INVALID"):
        review.ReviewDocumentAnalysis.model_validate(payload)


def test_ignored_native_reason_and_empty_result_never_claim_complete(monkeypatch):
    content = b"synthetic-docx"
    digest = hashlib.sha256(content).hexdigest()
    native_blocks = blocks(digest, ("text",))
    monkeypatch.setattr(
        review,
        "extract_document",
        lambda *_args, **_kwargs: extraction(
            native_blocks, reasons=(ExtractionReason.EXTERNAL_LINKS_IGNORED,)
        ),
    )
    analysis = review.analyze_review_document(content, "spec.docx")
    assert analysis.extraction_status == "UNKNOWN"
    assert analysis.extraction_reasons == ("EXTERNAL_LINKS_IGNORED",)
    assert analysis.native_text_completeness == "UNKNOWN"
    assert analysis.retained_extracted_text_complete

    monkeypatch.setattr(
        review,
        "extract_document",
        lambda *_args, **_kwargs: ExtractionResult(
            status=ExtractionStatus.UNKNOWN,
            reasons=(ExtractionReason.OCR_REQUIRED,),
            source_units_total=0,
            source_units_extracted=0,
            characters_extracted=0,
        ),
    )
    empty = review.analyze_review_document(content, "spec.docx")
    assert empty.blocks == () and empty.native_text_completeness == "UNKNOWN"


def test_rejection_is_unknown_without_invented_totals(monkeypatch):
    def reject(*_args, **_kwargs):
        raise DocumentRejected(DocumentRejectionCode.MALFORMED_CONTAINER)

    monkeypatch.setattr(review, "extract_document", reject)
    content = b"broken"
    analysis = review.analyze_review_document(content, "spec.docx")
    assert analysis.extraction_status == "UNKNOWN"
    assert analysis.rejection_code == "MALFORMED_CONTAINER"
    assert analysis.source_units_total is analysis.source_units_extracted is None
    assert analysis.characters_extracted is analysis.blocks_extracted_total is None
    assert analysis.blocks == () and analysis.native_text_completeness == "UNKNOWN"


def test_input_format_filename_and_oversize_are_bounded():
    with pytest.raises(ValueError, match="FORMAT_UNSUPPORTED"):
        review.analyze_review_document(b"x", "spec.pdf")
    with pytest.raises(ValueError, match="FILENAME_INVALID"):
        review.analyze_review_document(b"x", "../spec.docx")
    with pytest.raises(ValueError, match="CONTENT_INVALID"):
        review.analyze_review_document(b"", "spec.docx")
    with pytest.raises(ValueError, match="CONTENT_INVALID"):
        review.analyze_review_document(bytearray(b"x"), "spec.docx")  # type: ignore[arg-type]
    oversized = b"x" * (review.MAX_REVIEW_INPUT_BYTES + 1)
    with pytest.raises(ValueError, match="DOCUMENT_BYTES_LIMIT"):
        review.analyze_review_document(oversized, "spec.docx")


def test_page_inputs_text_cap_and_cached_model_consistency(monkeypatch):
    content = b"synthetic-docx"
    digest = hashlib.sha256(content).hexdigest()
    native_blocks = blocks(digest, ("поставка " + "x" * 30_000, "context"))
    monkeypatch.setattr(
        review, "extract_document", lambda *_args, **_kwargs: extraction(native_blocks)
    )
    analysis = review.analyze_review_document(content, "spec.docx")
    page = review.review_document_page(analysis, 0, 2)
    assert page.page_text_characters == review.MAX_PAGE_TEXT_CHARACTERS
    assert page.page_text_truncated and page.page_reasons == ("PAGE_TEXT_LIMIT",)
    assert len(page.blocks) == 1 and page.blocks[0].text_truncated
    assert page.next_offset == 0
    assert page.next_character_offset == review.MAX_PAGE_TEXT_CHARACTERS
    assert page.delivery_lead_indices == (0,)
    assert len(page.model_dump_json().encode()) <= review.MAX_PAGE_RESPONSE_BYTES

    pages = [page]
    while pages[-1].next_offset is not None:
        pages.append(
            review.review_document_page(
                analysis,
                pages[-1].next_offset,
                2,
                character_offset=pages[-1].next_character_offset,
            )
        )
    reconstructed: dict[int, str] = {}
    for current in pages:
        for block in current.blocks:
            reconstructed[block.block_index] = reconstructed.get(block.block_index, "") + block.text
    assert tuple(reconstructed[index] for index in sorted(reconstructed)) == tuple(
        block.text for block in analysis.blocks
    )
    assert pages[-1].blocks[-1].text == "context"

    for offset, limit in ((True, 1), (-1, 1), (1_000, 1), (0, True), (0, 0), (0, 26)):
        with pytest.raises(ValueError):
            review.review_document_page(analysis, offset, limit)  # type: ignore[arg-type]
    for character_offset in (True, -1, len(analysis.blocks[0].text)):
        with pytest.raises(ValueError, match="CHARACTER_OFFSET_INVALID"):
            review.review_document_page(
                analysis,
                0,
                1,
                character_offset=character_offset,  # type: ignore[arg-type]
            )

    payload = analysis.model_dump(mode="python")
    payload["retained_blocks"] = 1
    with pytest.raises(ValidationError, match="RETAINED_BLOCK_COUNT_INVALID"):
        review.ReviewDocumentAnalysis.model_validate(payload)
    payload = analysis.model_dump(mode="python")
    payload["blocks"][0]["provenance"]["document_id"] = "f" * 64
    with pytest.raises(ValidationError, match="PROVENANCE_MISMATCH"):
        review.ReviewDocumentAnalysis.model_validate(payload)
    page_payload = page.model_dump(mode="python")
    page_payload["blocks"][0]["block_index"] = 1
    with pytest.raises(ValidationError, match="PAGE_INDEX_INVALID"):
        review.ReviewDocumentPage.model_validate(page_payload)


def test_strict_counters_and_unknown_policy_are_not_coerced():
    content = docx_bytes("text")
    analysis = review.analyze_review_document(content, "spec.docx")
    payload = analysis.model_dump(mode="python")
    payload["retained_blocks"] = True
    with pytest.raises(ValidationError):
        review.ReviewDocumentAnalysis.model_validate(payload)
    payload = analysis.model_dump(mode="python")
    payload["full_tspec_completeness"] = "COMPLETE"
    with pytest.raises(ValidationError):
        review.ReviewDocumentAnalysis.model_validate(payload)


def test_cached_counter_and_completeness_corruption_is_rejected(monkeypatch):
    content = b"synthetic-docx"
    digest = hashlib.sha256(content).hexdigest()
    native_blocks = blocks(digest, ("text",))
    monkeypatch.setattr(
        review, "extract_document", lambda *_args, **_kwargs: extraction(native_blocks)
    )
    analysis = review.analyze_review_document(content, "spec.docx")

    payload = analysis.model_dump(mode="python")
    payload["source_units_total"] = 0
    with pytest.raises(ValidationError, match="SOURCE_UNIT_COUNT_INVALID"):
        review.ReviewDocumentAnalysis.model_validate(payload)

    payload = analysis.model_dump(mode="python")
    payload["characters_extracted"] = 0
    with pytest.raises(ValidationError, match="RETENTION_CHARACTER_COUNT_INVALID"):
        review.ReviewDocumentAnalysis.model_validate(payload)

    payload = analysis.model_dump(mode="python")
    payload.update(
        blocks=(),
        source_units_total=0,
        source_units_extracted=0,
        characters_extracted=0,
        blocks_extracted_total=0,
        retained_blocks=0,
        retained_characters=0,
        retained_extracted_text_complete=True,
        native_text_completeness="UNKNOWN",
    )
    with pytest.raises(ValidationError, match="RETAINED_COMPLETENESS_INVALID"):
        review.ReviewDocumentAnalysis.model_validate(payload)

    rejected = analysis.model_copy(
        update={
            "extraction_status": "UNKNOWN",
            "extraction_reasons": ("DOCUMENT_REJECTED",),
            "rejection_code": DocumentRejectionCode.PARSE_FAILED,
            "source_units_total": None,
            "source_units_extracted": None,
            "characters_extracted": None,
            "blocks_extracted_total": None,
            "blocks": (),
            "retained_blocks": 0,
            "retained_characters": 1,
            "native_text_completeness": "UNKNOWN",
        }
    )
    with pytest.raises(ValidationError, match="REJECTION_HAS_FACTS"):
        review.ReviewDocumentAnalysis.model_validate(rejected.model_dump(mode="python"))
