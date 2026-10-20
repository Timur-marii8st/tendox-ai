"""ZIP-пакет раскрывается без исполнения и сохраняет provenance участников."""

from __future__ import annotations

import hashlib
import json
import stat
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from procurement.cli import main
from procurement.documents.bundle import (
    BundleReason,
    BundleRejectionCode,
    DocumentBundleMemberUnavailable,
    DocumentBundleRejected,
    bundle_member_source_id,
    extract_document_bundle,
    read_document_bundle_member,
)
from procurement.documents.extraction import (
    DocumentInput,
    ExtractionLimits,
    ExtractionReason,
    ExtractionStatus,
)
from procurement.documents.facts import DocumentRole, map_blocks_to_requirements


def _docx(text: str) -> bytes:
    output = BytesIO()
    document = Document()
    document.add_paragraph(text)
    document.save(output)
    return output.getvalue()


def _xlsx(*, formula: bool = False) -> bytes:
    output = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Spec"
    sheet["A1"] = "Количество портов не менее 48 шт."
    if formula:
        sheet["A2"] = "=24+12"
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _pdf() -> bytes:
    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    writer.write(output)
    return output.getvalue()


def _zip(entries: tuple[tuple[str | ZipInfo, bytes], ...]) -> bytes:
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in entries:
            archive.writestr(name, content)
    return output.getvalue()


def _input(content: bytes, *, filename: str = "tender.zip") -> DocumentInput:
    return DocumentInput(
        document_id="bundle-1",
        filename=filename,
        declared_media_type="application/zip",
        content=content,
    )


def test_zip_docx_and_xlsx_reach_typed_requirements_with_member_provenance() -> None:
    content = _zip(
        (
            ("docs/requirements.docx", _docx("Гарантия не менее 24 мес.")),
            ("tables/spec.xlsx", _xlsx()),
        )
    )

    result = extract_document_bundle(_input(content))

    assert result.status is ExtractionStatus.PASS
    assert result.reasons == ()
    assert result.source_sha256 == hashlib.sha256(content).hexdigest()
    assert result.members_total == result.supported_members == 2
    assert result.unsupported_members == 0
    assert all(member.extraction.status is ExtractionStatus.PASS for member in result.members)
    blocks = tuple(block for member in result.members for block in member.extraction.blocks)
    assert all(block.provenance.document_id == "bundle-1" for block in blocks)
    assert {block.provenance.locator for block in blocks} == {
        'archive member "docs/requirements.docx" :: paragraph 1',
        'archive member "tables/spec.xlsx" :: Spec!A1',
    }
    facts = map_blocks_to_requirements(
        blocks,
        document_role=DocumentRole.TENDER_REQUIREMENTS,
    )
    assert {(item.property, str(item.value)) for item in facts.requirements} == {
        ("warranty", "24"),
        ("ethernet_ports", "48"),
    }


def test_formula_and_unsupported_member_keep_bundle_unknown_without_execution() -> None:
    content = _zip(
        (
            ("spec.xlsx", _xlsx(formula=True)),
            ("readme.txt", b"external instructions are untrusted"),
        )
    )

    result = extract_document_bundle(_input(content))

    assert result.status is ExtractionStatus.UNKNOWN
    assert result.reasons == (
        BundleReason.UNSUPPORTED_MEMBERS,
        BundleReason.MEMBER_EXTRACTION_UNKNOWN,
    )
    assert result.supported_members == result.unsupported_members == 1
    spreadsheet = next(member for member in result.members if member.media_type is not None)
    assert spreadsheet.extraction.reasons == (ExtractionReason.FORMULAS_IGNORED,)
    assert [block.text for block in spreadsheet.extraction.blocks] == [
        "Количество портов не менее 48 шт."
    ]
    unsupported = next(member for member in result.members if member.media_type is None)
    assert unsupported.extraction.status is ExtractionStatus.UNKNOWN
    assert unsupported.extraction.reasons == (ExtractionReason.UNSUPPORTED_FORMAT,)
    assert unsupported.extraction.blocks == ()


def test_scanned_pdf_member_remains_unknown_and_requests_ocr() -> None:
    result = extract_document_bundle(_input(_zip((("scan.pdf", _pdf()),))))

    assert result.status is ExtractionStatus.UNKNOWN
    assert result.reasons == (BundleReason.MEMBER_EXTRACTION_UNKNOWN,)
    assert result.members[0].extraction.reasons == (ExtractionReason.OCR_REQUIRED,)


@pytest.mark.parametrize(
    "path",
    (
        "../secret.pdf",
        "/absolute.pdf",
        "C:\\absolute.pdf",
        "folder/../secret.pdf",
        "folder//secret.pdf",
        "folder/./secret.pdf",
        "line\nbreak.pdf",
    ),
)
def test_unsafe_member_paths_reject_whole_bundle(path: str) -> None:
    with pytest.raises(DocumentBundleRejected) as exc_info:
        extract_document_bundle(_input(_zip(((path, b"data"),))))

    assert exc_info.value.code is BundleRejectionCode.UNSAFE_MEMBER_PATH


def test_case_insensitive_duplicate_member_paths_are_rejected() -> None:
    content = _zip((("Spec.xlsx", b"first"), ("spec.XLSX", b"second")))

    with pytest.raises(DocumentBundleRejected) as exc_info:
        extract_document_bundle(_input(content))

    assert exc_info.value.code is BundleRejectionCode.DUPLICATE_MEMBER_PATH


def test_nested_archive_is_rejected_even_with_disguised_extension() -> None:
    nested = _zip((("inside.txt", b"text"),))
    content = _zip((("attachment.bin", nested),))

    with pytest.raises(DocumentBundleRejected) as exc_info:
        extract_document_bundle(_input(content))

    assert exc_info.value.code is BundleRejectionCode.NESTED_ARCHIVE


def test_macro_enabled_member_is_rejected_before_document_parser() -> None:
    with pytest.raises(DocumentBundleRejected) as exc_info:
        extract_document_bundle(_input(_zip((("offer.xlsm", _xlsx()),))))

    assert exc_info.value.code is BundleRejectionCode.MACRO_ENABLED_MEMBER


def test_symlink_member_is_rejected() -> None:
    link = ZipInfo("link.pdf")
    link.create_system = 3
    link.external_attr = (stat.S_IFLNK | 0o777) << 16

    with pytest.raises(DocumentBundleRejected) as exc_info:
        extract_document_bundle(_input(_zip(((link, b"target.pdf"),))))

    assert exc_info.value.code is BundleRejectionCode.SYMLINK_MEMBER


def test_directory_entries_count_toward_member_limit() -> None:
    content = _zip((("one/", b""), ("two/", b"")))

    with pytest.raises(DocumentBundleRejected) as exc_info:
        extract_document_bundle(
            _input(content),
            limits=ExtractionLimits(max_zip_members=1),
        )

    assert exc_info.value.code is BundleRejectionCode.MEMBER_LIMIT


def test_expanded_size_limit_is_checked_before_reading_member() -> None:
    content = _zip((("large.txt", b"x" * 100),))

    with pytest.raises(DocumentBundleRejected) as exc_info:
        extract_document_bundle(
            _input(content),
            limits=ExtractionLimits(max_zip_expanded_bytes=99),
        )

    assert exc_info.value.code is BundleRejectionCode.EXPANSION_LIMIT


def test_member_type_mismatch_rejects_whole_bundle() -> None:
    with pytest.raises(DocumentBundleRejected) as exc_info:
        extract_document_bundle(_input(_zip((("fake.xlsx", b"not a workbook"),))))

    assert exc_info.value.code is BundleRejectionCode.MEMBER_REJECTED


@pytest.mark.parametrize(
    ("filename", "media_type", "content", "code"),
    (
        ("bundle.txt", "application/zip", _zip(()), BundleRejectionCode.FILENAME_INVALID),
        ("bundle.zip", "text/plain", _zip(()), BundleRejectionCode.MEDIA_TYPE_MISMATCH),
        ("bundle.zip", "application/zip", b"not zip", BundleRejectionCode.MAGIC_MISMATCH),
    ),
)
def test_outer_envelope_must_match_zip(
    filename: str,
    media_type: str,
    content: bytes,
    code: BundleRejectionCode,
) -> None:
    document = DocumentInput(
        document_id="bundle-1",
        filename=filename,
        declared_media_type=media_type,
        content=content,
    )

    with pytest.raises(DocumentBundleRejected) as exc_info:
        extract_document_bundle(document)

    assert exc_info.value.code is code


def test_empty_zip_is_unknown_not_pass() -> None:
    result = extract_document_bundle(_input(_zip(())))

    assert result.status is ExtractionStatus.UNKNOWN
    assert result.reasons == (BundleReason.EMPTY_ARCHIVE,)
    assert result.members == ()


def test_bundle_member_reader_returns_exact_verified_supported_bytes() -> None:
    pdf = _pdf()
    document = _input(_zip((("scans/page.pdf", pdf),)))
    result = extract_document_bundle(document)

    source = read_document_bundle_member(document, result, 0)

    assert source.source_bundle_id == "bundle-1"
    assert source.source_document_id == bundle_member_source_id("bundle-1", 0)
    assert source.archive_path == "scans/page.pdf"
    assert source.media_type == "application/pdf"
    assert source.content == pdf
    assert source.sha256 == hashlib.sha256(pdf).hexdigest()
    assert "content" not in source.model_dump()


def test_bundle_member_reader_rejects_missing_unsupported_and_changed_outer_source() -> None:
    document = _input(_zip((("unknown.ods", b"unsupported synthetic"),)))
    result = extract_document_bundle(document)

    with pytest.raises(IndexError):
        read_document_bundle_member(document, result, 1)
    with pytest.raises(DocumentBundleMemberUnavailable):
        read_document_bundle_member(document, result, 0)
    with pytest.raises(DocumentBundleMemberUnavailable):
        read_document_bundle_member(
            document.model_copy(update={"content": document.content + b"changed"}),
            result,
            0,
        )


def test_bundle_cli_publishes_full_report_but_prints_only_aggregates(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    source = tmp_path / "tender.zip"
    source.write_bytes(_zip((("requirements.docx", _docx("Гарантия не менее 24 мес.")),)))
    report_path = tmp_path / "bundle-report.json"

    exit_code = main(
        ["inspect-document-bundle", str(source), "--output", str(report_path)]
    )

    summary = json.loads(capsys.readouterr().out)
    assert exit_code == 0
    assert summary["status"] == "PASS"
    assert summary["supported_members"] == 1
    assert "Гарантия" not in json.dumps(summary, ensure_ascii=False)
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["members"][0]["extraction"]["blocks"][0]["text"] == (
        "Гарантия не менее 24 мес."
    )


def test_bundle_cli_returns_unknown_for_visible_unsupported_member(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    source = tmp_path / "tender.zip"
    source.write_bytes(_zip((("unknown.ods", b"unsupported synthetic spreadsheet"),)))
    report_path = tmp_path / "bundle-report.json"

    exit_code = main(
        ["inspect-document-bundle", str(source), "--output", str(report_path)]
    )

    summary = json.loads(capsys.readouterr().out)
    assert exit_code == 3
    assert summary["status"] == "UNKNOWN"
    assert summary["reasons"] == ["UNSUPPORTED_MEMBERS"]
    assert report_path.is_file()


def test_bundle_cli_rejection_does_not_publish_partial_report(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    source = tmp_path / "tender.zip"
    source.write_bytes(_zip((("../escape.pdf", b"unsafe"),)))
    report_path = tmp_path / "bundle-report.json"

    exit_code = main(
        ["inspect-document-bundle", str(source), "--output", str(report_path)]
    )

    error = json.loads(capsys.readouterr().err)
    assert exit_code == 2
    assert error == {
        "error": "document_bundle_rejected",
        "code": "UNSAFE_MEMBER_PATH",
    }
    assert not report_path.exists()
