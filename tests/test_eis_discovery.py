"""EIS import must never turn a matching title into a participation decision."""

from datetime import UTC, datetime
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document

from procurement.sources.eis_discovery import (
    EisExportError,
    discover_from_export,
    inspect_official_attachment,
)


def _notice(
    number: str = "0123456789012345678",
    title: str = "Поставка серверного оборудования",
    end: str = "2030-01-02T12:00:00+03:00",
    status: str = "published",
) -> bytes:
    return f"""<?xml version="1.0" encoding="utf-8"?>
<epNotificationEF2020 xmlns="http://zakupki.gov.ru/oos/types/1">
  <commonInfo><purchaseNumber>{number}</purchaseNumber>
    <purchaseObjectInfo>{title}</purchaseObjectInfo><status>{status}</status>
    <currency>RUB</currency><maxPrice>2500000.00</maxPrice>
  </commonInfo>
  <notificationInfo><procedureInfo><collecting><endDT>{end}</endDT>
  </collecting></procedureInfo></notificationInfo>
  <attachments><attachment><fileName>ТЗ.pdf</fileName>
    <url>https://zakupki.gov.ru/epz/order/notice/download/abc</url>
  </attachment></attachments>
</epNotificationEF2020>""".encode()


def _zip(*members: tuple[str, bytes]) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        for name, data in members:
            archive.writestr(name, data)
    return buffer.getvalue()


def test_official_export_envelope_preserves_notice_version_and_publication() -> None:
    notice = _notice().split(b"?>", 1)[1]
    notice = notice.replace(
        b"<commonInfo>",
        b"<versionNumber>2</versionNumber><commonInfo>"
        b"<publishDTInEIS>2026-09-25T17:03:33+03:00</publishDTInEIS>",
        1,
    )
    wrapped = b'<export xmlns="http://zakupki.gov.ru/oos/export/1">' + notice + b"</export>"
    result = discover_from_export(
        _zip(("notice.xml", wrapped)), "серверное", now=datetime(2026, 9, 25, tzinfo=UTC)
    )
    assert result.inspected_notices == 1
    lead = result.matched[0]
    assert lead.notice_type == "epNotificationEF2020"
    assert lead.version_number == 2
    assert lead.published_at is not None
    assert lead.published_at.isoformat() == "2026-09-25T17:03:33+03:00"


def test_export_envelope_with_two_documents_is_rejected() -> None:
    notice = _notice().split(b"?>", 1)[1]
    wrapped = (
        b'<export xmlns="http://zakupki.gov.ru/oos/export/1">' + notice + notice + b"</export>"
    )
    with pytest.raises(EisExportError, match="NO_VALID_NOTICES"):
        discover_from_export(_zip(("notice.xml", wrapped)), "серверное")


def test_structured_official_price_requires_ruble_code_in_same_block() -> None:
    notice = _notice().replace(b"<currency>RUB</currency><maxPrice>2500000.00</maxPrice>", b"")
    notice = notice.replace(
        b"<notificationInfo>",
        b"<notificationInfo><contractConditionsInfo><maxPriceInfo>"
        b"<maxPrice>182420.00</maxPrice><currency><code>RUB</code></currency>"
        b"</maxPriceInfo></contractConditionsInfo>",
        1,
    )
    result = discover_from_export(_zip(("notice.xml", notice)), "серверное")
    lead = result.matched[0]
    assert str(lead.nmck) == "182420.00"
    assert lead.nmck_currency == "RUB"
    assert not any("НМЦК в рублях не подтверждена" in reason for reason in lead.explanation)


def test_official_attachment_info_requires_exact_content_uid() -> None:
    uid = "A" * 32
    source = _notice().replace(
        b"<attachments><attachment><fileName>\xd0\xa2\xd0\x97.pdf</fileName>",
        f"<attachmentsInfo><attachmentInfo><publishedContentId>{uid}</publishedContentId>"
        "<fileName>Спецификация.docx</fileName><fileSize>123</fileSize>".encode(),
    )
    source = source.replace(
        b"<url>https://zakupki.gov.ru/epz/order/notice/download/abc</url>",
        f"<url>https://zakupki.gov.ru/44fz/filestore/public/1.0/download/priz/file.html?uid={uid}</url>".encode(),
    ).replace(b"</attachment></attachments>", b"</attachmentInfo></attachmentsInfo>")
    result = discover_from_export(source, "серверное")
    assert len(result.matched[0].attachments) == 1
    assert result.matched[0].attachments[0].locator == "attachmentInfo[1]/url"
    wrong = source.replace(f"uid={uid}".encode(), f"uid={'B' * 32}".encode())
    assert discover_from_export(wrong, "серверное").matched[0].attachments == ()


def test_eis_export_finds_competing_notices_with_provenance_and_unknown_economics() -> None:
    export = _zip(
        ("a.xml", _notice()),
        ("b.xml", _notice(number="0123456789012345679", title="Поставка серверного шкафа")),
        ("c.xml", _notice(number="0123456789012345680", title="Поставка бумаги")),
    )
    result = discover_from_export(export, "серверное", now=datetime(2026, 9, 25, tzinfo=UTC))
    assert result.inspected_notices == 3
    assert len(result.matched) == 2
    lead = result.matched[0]
    assert lead.purchase_number == "0123456789012345678"
    assert lead.nmck == 2500000
    assert lead.economic_status == "UNKNOWN"
    assert lead.application_status == "UNKNOWN"
    assert lead.xml_sha256 and lead.export_sha256 == result.export_sha256
    assert lead.attachments[0].url.startswith("https://zakupki.gov.ru/epz/")
    assert "прибыль" in " ".join(lead.explanation).lower()


def test_expired_or_naive_deadline_never_becomes_open() -> None:
    expired = discover_from_export(
        _notice(end="2020-01-01T00:00:00+03:00"),
        "серверного",
        now=datetime(2026, 9, 25, tzinfo=UTC),
    )
    assert expired.matched[0].application_status == "FAIL"
    naive = discover_from_export(_notice(end="2030-01-01T00:00:00"), "серверного")
    assert naive.matched[0].application_status == "UNKNOWN"
    assert naive.matched[0].deadline is None


def test_supplier_request_matches_paper_inflection() -> None:
    result = discover_from_export(_notice(title="Поставка бумаги А4"), "хочу бумагу")
    assert len(result.matched) == 1
    assert result.matched[0].relevance_points == 1
    assert result.matched[0].match_quality == "STRONG"


@pytest.mark.parametrize(
    "document_type", ["epNotificationEF2020", "epNotificationEZK2020", "epNotificationEOK2020"]
)
def test_export_matches_one_lot_item_with_exact_xml_locator(document_type: str) -> None:
    notice = _notice(title="Поставка оборудования").replace(
        b"<notificationInfo>",
        "<notificationInfo><purchaseObjectsInfo><notDrugPurchaseObjectsInfo>"
        "<purchaseObject><name>Шкаф металлический</name></purchaseObject>"
        "<purchaseObject><name>Серверное оборудование</name></purchaseObject>"
        "</notDrugPurchaseObjectsInfo></purchaseObjectsInfo>".encode(),
        1,
    )
    notice = notice.replace(b"epNotificationEF2020", document_type.encode())
    result = discover_from_export(notice, "серверное оборудование")
    assert len(result.matched) == 1
    lead = result.matched[0]
    assert lead.match_quality == "STRONG"
    assert lead.match_locator.endswith("purchaseObject[2]/name")
    assert lead.economic_status == "UNKNOWN"


def test_query_words_from_different_lot_items_are_not_combined() -> None:
    notice = _notice(title="Поставка товаров").replace(
        b"<notificationInfo>",
        "<notificationInfo><purchaseObjectsInfo><notDrugPurchaseObjectsInfo>"
        "<purchaseObject><name>Серверное оборудование</name></purchaseObject>"
        "<purchaseObject><name>Бумага офисная</name></purchaseObject>"
        "</notDrugPurchaseObjectsInfo></purchaseObjectsInfo>".encode(),
        1,
    )
    result = discover_from_export(notice, "серверная бумага")
    assert len(result.matched) == 1
    assert result.matched[0].match_quality == "PARTIAL"
    assert result.matched[0].relevance_points == 1


def test_rejects_untrusted_xml_and_unsupported_export() -> None:
    with pytest.raises(EisExportError, match="NO_VALID_NOTICES"):
        discover_from_export(b"<!DOCTYPE a><epNotificationEF2020/>", "серверного")
    with pytest.raises(EisExportError, match="NO_VALID_NOTICES"):
        discover_from_export(b"<! DOCTYPE a><epNotificationEF2020/>", "серверного")
    with pytest.raises(EisExportError, match="NO_VALID_NOTICES"):
        discover_from_export(
            b"<notAnEisNotice><purchaseNumber>0123456789012345678</purchaseNumber></notAnEisNotice>",
            "серверного",
        )
    with pytest.raises(EisExportError, match="NO_XML_NOTICES"):
        discover_from_export(_zip(("readme.txt", b"not xml")), "серверного")


def test_attachment_download_has_official_host_gate() -> None:
    for url in (
        "http://zakupki.gov.ru/epz/a",
        "https://evil.test/epz/a",
        "https://zakupki.gov.ru.evil.test/epz/a",
        "https://zakupki.gov.ru@evil.test/epz/a",
        "https://zakupki.gov.ru:8443/epz/a",
    ):
        with pytest.raises(EisExportError, match="ATTACHMENT_URL_INVALID"):
            inspect_official_attachment(url, "ТЗ.pdf")


def test_official_attachment_download_extracts_exact_bytes() -> None:
    document = Document()
    document.add_paragraph("Поставка серверного оборудования")
    buffer = BytesIO()
    document.save(buffer)
    content = buffer.getvalue()

    url = "https://zakupki.gov.ru/44fz/filestore/public/1.0/download/priz/file.html?uid=fixture"

    def fetcher(requested: str) -> bytes:
        assert requested == url
        return content

    result = inspect_official_attachment(url, "ТЗ.docx", fetcher=fetcher)
    assert result.byte_count == len(content)
    assert result.extraction_status == "PASS"
    assert "серверного" in result.text_preview
    assert len(result.sha256) == 64


def test_official_zip_attachment_extracts_member_text() -> None:
    document = Document()
    document.add_paragraph("Бумага формата А4")
    buffer = BytesIO()
    document.save(buffer)
    archive = _zip(("техническое_задание.docx", buffer.getvalue()))
    result = inspect_official_attachment(
        "https://zakupki.gov.ru/44fz/filestore/public/1.0/download/priz/file.html?uid=fixture",
        "документы.zip",
        fetcher=lambda _: archive,
    )
    assert result.extraction_status == "PASS"
    assert result.bundle_members == ("техническое_задание.docx",)
    assert "Бумага" in result.text_preview
