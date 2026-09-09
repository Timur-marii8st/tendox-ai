"""Детерминированное и безопасное извлечение нативного текста.

Этот модуль принимает уже неизменяемый байтовый снимок документа. Он не
интерпретирует содержание как команды, не вычисляет формулы и не обращается к
внешним ссылкам. Неудача извлечения никогда не становится фактом: пустой или
сканированный документ возвращает ``UNKNOWN``, а повреждённый или запрещённый
файл отклоняется до разбора.
"""

from __future__ import annotations

from dataclasses import dataclass
from enum import StrEnum
from io import BytesIO
from pathlib import PurePath, PureWindowsPath
from typing import TYPE_CHECKING
from zipfile import BadZipFile, ZipFile

from pydantic import BaseModel, ConfigDict, Field, model_validator

from procurement.domain.provenance import Extractor, Provenance

if TYPE_CHECKING:
    from collections.abc import Iterable


MAX_UPLOAD_BYTES = 20 * 1024 * 1024
"""Единый лимит загружаемого документа для API и извлекателя: 20 МиБ."""


class ExtractionStatus(StrEnum):
    """Состояние извлечения, не являющееся решением о соответствии."""

    PASS = "PASS"
    FAIL = "FAIL"
    UNKNOWN = "UNKNOWN"


class ExtractionReason(StrEnum):
    """Причина неполноты или ограничение результата извлечения."""

    OCR_REQUIRED = "OCR_REQUIRED"
    UNSUPPORTED_FORMAT = "UNSUPPORTED_FORMAT"
    EXTERNAL_LINKS_IGNORED = "EXTERNAL_LINKS_IGNORED"
    FORMULAS_IGNORED = "FORMULAS_IGNORED"


class DocumentRejectionCode(StrEnum):
    """Детерминированные причины отказа до создания фактов из документа."""

    INPUT_TOO_LARGE = "INPUT_TOO_LARGE"
    FILENAME_INVALID = "FILENAME_INVALID"
    MEDIA_TYPE_MISMATCH = "MEDIA_TYPE_MISMATCH"
    MAGIC_MISMATCH = "MAGIC_MISMATCH"
    MACRO_ENABLED_FORMAT = "MACRO_ENABLED_FORMAT"
    ENCRYPTED_DOCUMENT = "ENCRYPTED_DOCUMENT"
    MALFORMED_CONTAINER = "MALFORMED_CONTAINER"
    ZIP_MEMBER_LIMIT = "ZIP_MEMBER_LIMIT"
    ZIP_EXPANSION_LIMIT = "ZIP_EXPANSION_LIMIT"
    SOURCE_UNIT_LIMIT = "SOURCE_UNIT_LIMIT"
    EXTRACTED_TEXT_LIMIT = "EXTRACTED_TEXT_LIMIT"
    PARSE_FAILED = "PARSE_FAILED"


class DocumentRejected(ValueError):
    """Документ нельзя безопасно использовать; детали не содержат его текста."""

    def __init__(self, code: DocumentRejectionCode) -> None:
        self.code = code
        super().__init__(f"document rejected: {code.value}")


class DocumentInput(BaseModel):
    """Неизменяемый вход нативного извлекателя.

    ``document_id`` задаётся хранилищем и используется в provenance; имя и MIME
    не считаются доказательством формата без проверки сигнатуры содержимого.
    """

    model_config = ConfigDict(frozen=True)

    document_id: str = Field(min_length=1)
    filename: str = Field(min_length=1)
    declared_media_type: str = Field(min_length=1)
    content: bytes = Field(min_length=1)

    @model_validator(mode="after")
    def _filename_is_basename(self) -> DocumentInput:
        if "\x00" in self.filename:
            raise ValueError("filename must not contain NUL")
        if PurePath(self.filename).name != self.filename:
            raise ValueError("filename must be a basename")
        if PureWindowsPath(self.filename).name != self.filename:
            raise ValueError("filename must be a basename")
        return self


class ExtractionLimits(BaseModel):
    """Ограничения ресурсов до передачи байтов сторонним парсерам."""

    model_config = ConfigDict(frozen=True)

    max_input_bytes: int = Field(default=MAX_UPLOAD_BYTES, ge=1)
    max_zip_members: int = Field(default=1_000, ge=1)
    max_zip_expanded_bytes: int = Field(default=100 * 1024 * 1024, ge=1)
    max_source_units: int = Field(default=200_000, ge=1)
    max_pdf_pages: int = Field(default=2_000, ge=1)
    max_pdf_stream_bytes: int = Field(default=20 * 1024 * 1024, ge=1)
    max_extracted_characters: int = Field(default=5_000_000, ge=1)


DEFAULT_EXTRACTION_LIMITS = ExtractionLimits()


class ExtractedBlock(BaseModel):
    """Нативный фрагмент текста с точным происхождением."""

    model_config = ConfigDict(frozen=True)

    text: str = Field(min_length=1)
    provenance: Provenance


class ExtractionResult(BaseModel):
    """Результат извлечения без семантических выводов о документе."""

    model_config = ConfigDict(frozen=True)

    status: ExtractionStatus
    blocks: tuple[ExtractedBlock, ...] = Field(default=())
    reasons: tuple[ExtractionReason, ...] = Field(default=())
    source_units_total: int = Field(ge=0)
    source_units_extracted: int = Field(ge=0)
    characters_extracted: int = Field(ge=0)

    @model_validator(mode="after")
    def _result_is_honest(self) -> ExtractionResult:
        if self.source_units_extracted > self.source_units_total:
            raise ValueError("extracted source units exceed total")
        if self.source_units_extracted != len(self.blocks):
            raise ValueError("each extracted source unit must have one provenance block")
        if self.characters_extracted != sum(len(block.text) for block in self.blocks):
            raise ValueError("characters_extracted must equal the block text length")
        if self.status is ExtractionStatus.PASS and not self.blocks:
            raise ValueError("PASS requires useful native text")
        if self.status is ExtractionStatus.PASS and self.characters_extracted == 0:
            raise ValueError("PASS requires useful native text")
        return self


class _DocumentKind(StrEnum):
    PDF = "pdf"
    DOCX = "docx"
    XLSX = "xlsx"


@dataclass(frozen=True)
class _ZipInspection:
    names: frozenset[str]
    has_external_links: bool
    has_formulas: bool


_MEDIA_TYPES: dict[_DocumentKind, frozenset[str]] = {
    _DocumentKind.PDF: frozenset({"application/pdf"}),
    _DocumentKind.DOCX: frozenset(
        {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
    ),
    _DocumentKind.XLSX: frozenset(
        {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    ),
}
_EXTENSIONS: dict[str, _DocumentKind] = {
    ".pdf": _DocumentKind.PDF,
    ".docx": _DocumentKind.DOCX,
    ".xlsx": _DocumentKind.XLSX,
}
_MACRO_EXTENSIONS = frozenset({".docm", ".xlsm", ".xltm"})
_ZIP_MAGIC = b"PK\x03\x04"
_PDF_MAGIC = b"%PDF-"


def extract_document(
    document: DocumentInput,
    *,
    limits: ExtractionLimits | None = None,
) -> ExtractionResult:
    """Извлечь нативный текст из PDF, DOCX или XLSX без побочных действий.

    Пустой PDF часто означает скан. В этом случае результат остаётся
    ``UNKNOWN`` с ``OCR_REQUIRED`` и не создаёт текстовых фактов.

    >>> result = extract_document(DocumentInput(
    ...     document_id="scan.pdf", filename="scan.pdf",
    ...     declared_media_type="application/pdf", content=b"%PDF-1.4\\n",
    ... ))
    Traceback (most recent call last):
    ...
    procurement.documents.extraction.DocumentRejected: document rejected: PARSE_FAILED
    """
    active_limits = limits or DEFAULT_EXTRACTION_LIMITS
    _validate_input_size(document, active_limits)
    kind = _classify(document)
    if kind is None:
        return _unknown_unsupported()
    if kind is _DocumentKind.PDF:
        return _extract_pdf(document, active_limits)

    inspection = _inspect_zip(document.content, kind, active_limits)
    if kind is _DocumentKind.DOCX:
        return _extract_docx(document, inspection, active_limits)
    return _extract_xlsx(document, inspection, active_limits)


def _validate_input_size(document: DocumentInput, limits: ExtractionLimits) -> None:
    if len(document.content) > limits.max_input_bytes:
        raise DocumentRejected(DocumentRejectionCode.INPUT_TOO_LARGE)


def _classify(document: DocumentInput) -> _DocumentKind | None:
    suffix = PurePath(document.filename).suffix.lower()
    if suffix in _MACRO_EXTENSIONS:
        raise DocumentRejected(DocumentRejectionCode.MACRO_ENABLED_FORMAT)
    kind = _EXTENSIONS.get(suffix)
    if kind is None:
        return None

    if document.declared_media_type.strip().lower() not in _MEDIA_TYPES[kind]:
        raise DocumentRejected(DocumentRejectionCode.MEDIA_TYPE_MISMATCH)
    if kind is _DocumentKind.PDF and not document.content.startswith(_PDF_MAGIC):
        raise DocumentRejected(DocumentRejectionCode.MAGIC_MISMATCH)
    if kind in (_DocumentKind.DOCX, _DocumentKind.XLSX) and not document.content.startswith(
        _ZIP_MAGIC
    ):
        raise DocumentRejected(DocumentRejectionCode.MAGIC_MISMATCH)
    return kind


def _inspect_zip(content: bytes, kind: _DocumentKind, limits: ExtractionLimits) -> _ZipInspection:
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > limits.max_zip_members:
                raise DocumentRejected(DocumentRejectionCode.ZIP_MEMBER_LIMIT)
            if sum(info.file_size for info in infos) > limits.max_zip_expanded_bytes:
                raise DocumentRejected(DocumentRejectionCode.ZIP_EXPANSION_LIMIT)

            names = frozenset(info.filename for info in infos)
            required_member = (
                "word/document.xml" if kind is _DocumentKind.DOCX else "xl/workbook.xml"
            )
            if required_member not in names:
                raise DocumentRejected(DocumentRejectionCode.MALFORMED_CONTAINER)
            if any("vbaProject.bin" in name for name in names):
                raise DocumentRejected(DocumentRejectionCode.MACRO_ENABLED_FORMAT)

            has_external_links = any(name.startswith("xl/externalLinks/") for name in names)
            has_external_links = has_external_links or any(
                b'TargetMode="External"' in archive.read(name)
                for name in names
                if name.startswith("word/_rels/") and name.endswith(".rels")
            )
            has_formulas = False
            if kind is _DocumentKind.XLSX:
                has_formulas = any(
                    b"<f" in archive.read(name)
                    for name in names
                    if name.startswith("xl/worksheets/") and name.endswith(".xml")
                )
            return _ZipInspection(
                names=names,
                has_external_links=has_external_links,
                has_formulas=has_formulas,
            )
    except DocumentRejected:
        raise
    except (BadZipFile, OSError, RuntimeError) as error:
        raise DocumentRejected(DocumentRejectionCode.MALFORMED_CONTAINER) from error


def _extract_pdf(document: DocumentInput, limits: ExtractionLimits) -> ExtractionResult:
    try:
        from pypdf import PdfReader, apply_configuration

        with apply_configuration(
            maximum_declared_stream_length=limits.max_pdf_stream_bytes,
            array_based_stream_maximum_output_length=limits.max_pdf_stream_bytes,
            jbig2_maximum_output_length=limits.max_pdf_stream_bytes,
            lzw_maximum_output_length=limits.max_pdf_stream_bytes,
            run_length_maximum_output_length=limits.max_pdf_stream_bytes,
            zlib_maximum_output_length=limits.max_pdf_stream_bytes,
            image_maximum_buffer_size=limits.max_pdf_stream_bytes,
        ):
            reader = PdfReader(BytesIO(document.content), strict=True)
            if reader.is_encrypted:
                raise DocumentRejected(DocumentRejectionCode.ENCRYPTED_DOCUMENT)
            if len(reader.pages) > limits.max_pdf_pages:
                raise DocumentRejected(DocumentRejectionCode.SOURCE_UNIT_LIMIT)

            extracted: list[ExtractedBlock] = []
            characters = 0
            for number, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                if not text.strip():
                    continue
                characters += len(text.strip())
                if characters > limits.max_extracted_characters:
                    raise DocumentRejected(DocumentRejectionCode.EXTRACTED_TEXT_LIMIT)
                extracted.append(
                    _block(document.document_id, f"page {number}", Extractor.NATIVE_PDF, text)
                )
            blocks = tuple(extracted)
    except DocumentRejected:
        raise
    except Exception as error:
        raise DocumentRejected(DocumentRejectionCode.PARSE_FAILED) from error
    return _result_or_ocr_required(blocks, len(reader.pages))


def _extract_docx(
    document: DocumentInput,
    inspection: _ZipInspection,
    limits: ExtractionLimits,
) -> ExtractionResult:
    try:
        from docx import Document

        parsed = Document(BytesIO(document.content))
        units: list[tuple[str, str]] = []
        units.extend(
            (f"paragraph {number}", paragraph.text)
            for number, paragraph in enumerate(parsed.paragraphs, start=1)
        )
        seen_cells: set[int] = set()
        for table_number, table in enumerate(parsed.tables, start=1):
            for row_number, row in enumerate(table.rows, start=1):
                for cell_number, cell in enumerate(row.cells, start=1):
                    if id(cell._tc) in seen_cells:
                        continue
                    seen_cells.add(id(cell._tc))
                    units.append(
                        (f"table {table_number} row {row_number} cell {cell_number}", cell.text)
                    )
        _check_source_units(len(units), limits)
    except DocumentRejected:
        raise
    except Exception as error:
        raise DocumentRejected(DocumentRejectionCode.PARSE_FAILED) from error

    blocks = _blocks_from_units(
        document.document_id,
        Extractor.NATIVE_DOCX,
        units,
        max_characters=limits.max_extracted_characters,
    )
    warnings = _link_warning(inspection)
    return _result_or_unknown(blocks, len(units), warnings)


def _extract_xlsx(
    document: DocumentInput,
    inspection: _ZipInspection,
    limits: ExtractionLimits,
) -> ExtractionResult:
    try:
        from openpyxl import load_workbook  # type: ignore[import-untyped]

        workbook = load_workbook(
            BytesIO(document.content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        formula_cells: set[tuple[str, str]] = set()
        if inspection.has_formulas:
            formula_workbook = load_workbook(
                BytesIO(document.content),
                read_only=True,
                data_only=False,
                keep_links=False,
            )
            for formula_sheet in formula_workbook.worksheets:
                for row in formula_sheet.iter_rows():
                    for cell in row:
                        if cell.data_type == "f":
                            formula_cells.add((formula_sheet.title, cell.coordinate))
            formula_workbook.close()
        units: list[tuple[str, str]] = []
        source_units_total = 0
        for worksheet in workbook.worksheets:
            capacity = worksheet.max_row * worksheet.max_column
            if capacity > limits.max_source_units - source_units_total:
                raise DocumentRejected(DocumentRejectionCode.SOURCE_UNIT_LIMIT)
            source_units_total += capacity
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    is_formula = (worksheet.title, cell.coordinate) in formula_cells
                    if not is_formula:
                        units.append((f"{worksheet.title}!{cell.coordinate}", str(cell.value)))
        workbook.close()
    except DocumentRejected:
        raise
    except Exception as error:
        raise DocumentRejected(DocumentRejectionCode.PARSE_FAILED) from error

    reasons = list(_link_warning(inspection))
    if inspection.has_formulas:
        reasons.append(ExtractionReason.FORMULAS_IGNORED)
    blocks = _blocks_from_units(
        document.document_id,
        Extractor.NATIVE_XLSX,
        units,
        max_characters=limits.max_extracted_characters,
    )
    return _result_or_unknown(blocks, source_units_total, tuple(reasons))


def _check_source_units(total: int, limits: ExtractionLimits) -> None:
    if total > limits.max_source_units:
        raise DocumentRejected(DocumentRejectionCode.SOURCE_UNIT_LIMIT)


def _blocks_from_units(
    document_id: str,
    extractor: Extractor,
    units: Iterable[tuple[str, str]],
    max_characters: int | None = None,
) -> tuple[ExtractedBlock, ...]:
    blocks: list[ExtractedBlock] = []
    characters = 0
    for locator, text in units:
        if not text.strip():
            continue
        characters += len(text.strip())
        if max_characters is not None and characters > max_characters:
            raise DocumentRejected(DocumentRejectionCode.EXTRACTED_TEXT_LIMIT)
        blocks.append(_block(document_id, locator, extractor, text))
    return tuple(blocks)


def _block(
    document_id: str,
    locator: str,
    extractor: Extractor,
    text: str,
) -> ExtractedBlock:
    cleaned = text.strip()
    return ExtractedBlock(
        text=cleaned,
        provenance=Provenance(
            document_id=document_id,
            locator=locator,
            extractor=extractor,
            confidence=1.0,
        ),
    )


def _result_or_ocr_required(
    blocks: tuple[ExtractedBlock, ...], source_units_total: int
) -> ExtractionResult:
    if not blocks:
        return ExtractionResult(
            status=ExtractionStatus.UNKNOWN,
            reasons=(ExtractionReason.OCR_REQUIRED,),
            source_units_total=source_units_total,
            source_units_extracted=0,
            characters_extracted=0,
        )
    return _pass_result(blocks, source_units_total)


def _result_or_unknown(
    blocks: tuple[ExtractedBlock, ...],
    source_units_total: int,
    reasons: tuple[ExtractionReason, ...],
) -> ExtractionResult:
    if not blocks:
        return ExtractionResult(
            status=ExtractionStatus.UNKNOWN,
            reasons=reasons,
            source_units_total=source_units_total,
            source_units_extracted=0,
            characters_extracted=0,
        )
    return ExtractionResult(
        status=ExtractionStatus.UNKNOWN if reasons else ExtractionStatus.PASS,
        blocks=blocks,
        reasons=reasons,
        source_units_total=source_units_total,
        source_units_extracted=len(blocks),
        characters_extracted=sum(len(block.text) for block in blocks),
    )


def _pass_result(blocks: tuple[ExtractedBlock, ...], source_units_total: int) -> ExtractionResult:
    return _result_or_unknown(blocks, source_units_total, ())


def _unknown_unsupported() -> ExtractionResult:
    return ExtractionResult(
        status=ExtractionStatus.UNKNOWN,
        reasons=(ExtractionReason.UNSUPPORTED_FORMAT,),
        source_units_total=0,
        source_units_extracted=0,
        characters_extracted=0,
    )


def _link_warning(inspection: _ZipInspection) -> tuple[ExtractionReason, ...]:
    if inspection.has_external_links:
        return (ExtractionReason.EXTERNAL_LINKS_IGNORED,)
    return ()
